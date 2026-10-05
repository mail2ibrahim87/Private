"""Generates Unit_Converter.xlsx (formula-only, no macros).

Usage:  python3 tools/build_unit_converter.py Unit_Converter.xlsx
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

# (unit label, system, factor-to-SI-base [number or formula string], offset)
# SI value = (x + offset) * factor ;  x = SI value / factor - offset
CATS = [
 ("Length", "metre (m)", "Meter (m)", "Foot (ft)", [
  ("Kilometre (km)","SI / Metric",1000),("Meter (m)","SI / Metric",1),("Centimetre (cm)","SI / Metric",0.01),
  ("Millimetre (mm)","SI / Metric",0.001),("Micrometre (µm)","SI / Metric",1e-6),("Nanometre (nm)","SI / Metric",1e-9),
  ("Mile (mi)","Imperial / US",1609.344),("Yard (yd)","Imperial / US",0.9144),("Foot (ft)","Imperial / US",0.3048),
  ("Inch (in)","Imperial / US",0.0254),("Thou / mil","Imperial / US",0.0000254),("Furlong","Imperial / US",201.168),
  ("Chain","Imperial / US",20.1168),("Fathom","Imperial / US",1.8288),("Nautical mile (nmi)","Other",1852),
  ("Ångström (Å)","Other",1e-10),("Astronomical unit (au)","Other",149597870700),("Light-year (ly)","Other",9460730472580800)]),
 ("Area", "square metre (m²)", "Square meter (m²)", "Square foot (ft²)", [
  ("Square kilometre (km²)","SI / Metric",1e6),("Hectare (ha)","SI / Metric",1e4),("Are (a)","SI / Metric",100),
  ("Square meter (m²)","SI / Metric",1),("Square centimetre (cm²)","SI / Metric",1e-4),("Square millimetre (mm²)","SI / Metric",1e-6),
  ("Square mile (mi²)","Imperial / US",2589988.110336),("Acre (ac)","Imperial / US",4046.8564224),
  ("Square yard (yd²)","Imperial / US",0.83612736),("Square foot (ft²)","Imperial / US",0.09290304),
  ("Square inch (in²)","Imperial / US",0.00064516)]),
 ("Volume", "cubic metre (m³)", "Litre (L)", "US gallon (gal)", [
  ("Cubic meter (m³)","SI / Metric",1),("Litre (L)","SI / Metric",0.001),("Millilitre (mL)","SI / Metric",1e-6),
  ("Cubic centimetre (cm³)","SI / Metric",1e-6),
  ("US gallon (gal)","US customary",0.003785411784),("US quart (qt)","US customary",0.000946352946),
  ("US pint (pt)","US customary",0.000473176473),("US cup","US customary",0.0002365882365),
  ("US fluid ounce (fl oz)","US customary",0.0000295735295625),("US tablespoon (tbsp)","US customary",0.00001478676478125),
  ("US teaspoon (tsp)","US customary",0.00000492892159375),
  ("Imperial gallon (gal)","Imperial (UK)",0.00454609),("Imperial quart (qt)","Imperial (UK)",0.0011365225),
  ("Imperial pint (pt)","Imperial (UK)",0.00056826125),("Imperial fluid ounce (fl oz)","Imperial (UK)",0.0000284130625),
  ("Cubic yard (yd³)","Imperial / US",0.764554857984),("Cubic foot (ft³)","Imperial / US",0.028316846592),
  ("Cubic inch (in³)","Imperial / US",0.000016387064),("Oil barrel (bbl)","Other",0.158987294928)]),
 ("Mass", "kilogram (kg)", "Kilogram (kg)", "Pound (lb)", [
  ("Tonne (t)","SI / Metric",1000),("Kilogram (kg)","SI / Metric",1),("Gram (g)","SI / Metric",0.001),
  ("Milligram (mg)","SI / Metric",1e-6),("Microgram (µg)","SI / Metric",1e-9),
  ("Long ton (UK)","Imperial / US",1016.0469088),("Short ton (US)","Imperial / US",907.18474),
  ("Stone (st)","Imperial / US",6.35029318),("Pound (lb)","Imperial / US",0.45359237),
  ("Ounce (oz)","Imperial / US",0.028349523125),("Grain (gr)","Imperial / US",0.00006479891),
  ("Troy ounce (oz t)","Other",0.0311034768),("Carat (ct)","Other",0.0002)]),
 ("Temperature", "kelvin (K)", "Degree Celsius (°C)", "Degree Fahrenheit (°F)", [
  ("Degree Celsius (°C)","SI / Metric",1,273.15),("Kelvin (K)","SI / Metric",1,0),
  ("Degree Fahrenheit (°F)","Imperial / US","=5/9",459.67),("Rankine (°R)","Imperial / US","=5/9",0)]),
 ("Speed", "metre per second (m/s)", "Kilometre per hour (km/h)", "Mile per hour (mph)", [
  ("Metre per second (m/s)","SI / Metric",1),("Kilometre per hour (km/h)","SI / Metric","=1/3.6"),
  ("Mile per hour (mph)","Imperial / US",0.44704),("Foot per second (ft/s)","Imperial / US",0.3048),
  ("Knot (kn)","Other","=1852/3600")]),
 ("Time", "second (s)", "Hour (h)", "Minute (min)", [
  ("Nanosecond (ns)","SI / Metric",1e-9),("Microsecond (µs)","SI / Metric",1e-6),("Millisecond (ms)","SI / Metric",0.001),
  ("Second (s)","SI / Metric",1),("Minute (min)","Other",60),("Hour (h)","Other",3600),("Day (d)","Other",86400),
  ("Week (wk)","Other",604800),("Month (avg, 30.436875 d)","Other",2629746),("Year (avg, 365.2425 d)","Other",31556952),
  ("Decade","Other",315569520),("Century","Other",3155695200)]),
 ("Pressure", "pascal (Pa)", "Bar (bar)", "Pound per square inch (psi)", [
  ("Pascal (Pa)","SI / Metric",1),("Kilopascal (kPa)","SI / Metric",1000),("Megapascal (MPa)","SI / Metric",1e6),
  ("Bar (bar)","SI / Metric",1e5),("Millibar (mbar)","SI / Metric",100),
  ("Pound per square inch (psi)","Imperial / US",6894.757293168361),("Kilopound per sq inch (ksi)","Imperial / US",6894757.293168361),
  ("Inch of mercury (inHg)","Imperial / US",3386.389),("Standard atmosphere (atm)","Other",101325),
  ("Torr","Other","=101325/760"),("Millimetre of mercury (mmHg)","Other",133.322387415),
  ("Kilogram-force per cm² (kgf/cm²)","Other",98066.5)]),
 ("Energy", "joule (J)", "Kilojoule (kJ)", "Kilocalorie (kcal)", [
  ("Joule (J)","SI / Metric",1),("Kilojoule (kJ)","SI / Metric",1000),("Megajoule (MJ)","SI / Metric",1e6),
  ("Watt-hour (Wh)","SI / Metric",3600),("Kilowatt-hour (kWh)","SI / Metric",3.6e6),
  ("Calorie (cal)","Other",4.184),("Kilocalorie (kcal)","Other",4184),
  ("British thermal unit (BTU)","Imperial / US",1055.05585262),("Therm (US)","Imperial / US",105480400),
  ("Foot-pound force (ft·lbf)","Imperial / US",1.3558179483314004),("Electronvolt (eV)","Other",1.602176634e-19),
  ("Erg","Other",1e-7)]),
 ("Power", "watt (W)", "Kilowatt (kW)", "Horsepower (hp)", [
  ("Watt (W)","SI / Metric",1),("Kilowatt (kW)","SI / Metric",1000),("Megawatt (MW)","SI / Metric",1e6),
  ("Horsepower (hp)","Imperial / US",745.6998715822702),("Metric horsepower (PS)","Other",735.49875),
  ("BTU per hour (BTU/h)","Imperial / US",0.29307107017222),("Foot-pound force per second (ft·lbf/s)","Imperial / US",1.3558179483314004),
  ("Ton of refrigeration (TR)","Imperial / US",3516.8528420667),("Kilocalorie per hour (kcal/h)","Other",1.163)]),
 ("Force", "newton (N)", "Newton (N)", "Pound-force (lbf)", [
  ("Newton (N)","SI / Metric",1),("Kilonewton (kN)","SI / Metric",1000),("Dyne (dyn)","SI / Metric",1e-5),
  ("Pound-force (lbf)","Imperial / US",4.4482216152605),("Kip (kip)","Imperial / US",4448.2216152605),
  ("Poundal (pdl)","Imperial / US",0.138254954376),("Kilogram-force (kgf)","Other",9.80665),
  ("Tonne-force (tf)","Other",9806.65)]),
 ("Torque", "newton metre (N·m)", "Newton metre (N·m)", "Pound-force foot (lbf·ft)", [
  ("Newton metre (N·m)","SI / Metric",1),("Kilonewton metre (kN·m)","SI / Metric",1000),
  ("Pound-force foot (lbf·ft)","Imperial / US",1.3558179483314004),("Pound-force inch (lbf·in)","Imperial / US",0.1129848290276167),
  ("Ounce-force inch (ozf·in)","Imperial / US",0.00706155181422604),("Kilogram-force metre (kgf·m)","Other",9.80665)]),
 ("Density", "kilogram per cubic metre (kg/m³)", "Kilogram per cubic metre (kg/m³)", "Pound per cubic foot (lb/ft³)", [
  ("Kilogram per cubic metre (kg/m³)","SI / Metric",1),("Gram per cubic centimetre (g/cm³)","SI / Metric",1000),
  ("Kilogram per litre (kg/L)","SI / Metric",1000),("Gram per litre (g/L)","SI / Metric",1),
  ("Pound per cubic foot (lb/ft³)","Imperial / US",16.018463373960138),("Pound per cubic inch (lb/in³)","Imperial / US",27679.904710203125),
  ("Pound per US gallon (lb/gal)","Imperial / US",119.82642731689663),("Ounce per cubic inch (oz/in³)","Imperial / US",1729.994044387695)]),
 ("Flow Rate", "cubic metre per second (m³/s)", "Litre per minute (L/min)", "US gallon per minute (gpm)", [
  ("Cubic metre per second (m³/s)","SI / Metric",1),("Cubic metre per hour (m³/h)","SI / Metric","=1/3600"),
  ("Litre per second (L/s)","SI / Metric",0.001),("Litre per minute (L/min)","SI / Metric","=0.001/60"),
  ("US gallon per minute (gpm)","US customary","=0.003785411784/60"),("Imperial gallon per minute","Imperial (UK)","=0.00454609/60"),
  ("Cubic foot per minute (cfm)","Imperial / US","=0.028316846592/60"),("Cubic foot per second (cfs)","Imperial / US",0.028316846592),
  ("Oil barrel per day (bbl/d)","Other","=0.158987294928/86400")]),
 ("Angle", "radian (rad)", "Degree (°)", "Radian (rad)", [
  ("Radian (rad)","SI / Metric",1),("Milliradian (mrad)","SI / Metric",0.001),("Degree (°)","Other","=PI()/180"),
  ("Arcminute (′)","Other","=PI()/10800"),("Arcsecond (″)","Other","=PI()/648000"),("Gradian (gon)","Other","=PI()/200"),
  ("Revolution (rev)","Other","=2*PI()")]),
 ("Frequency", "hertz (Hz)", "Hertz (Hz)", "Revolutions per minute (rpm)", [
  ("Hertz (Hz)","SI / Metric",1),("Kilohertz (kHz)","SI / Metric",1000),("Megahertz (MHz)","SI / Metric",1e6),
  ("Gigahertz (GHz)","SI / Metric",1e9),("Revolutions per minute (rpm)","Other","=1/60"),
  ("Radian per second (rad/s)","Other","=1/(2*PI())")]),
 ("Data Storage", "byte (B)", "Megabyte (MB)", "Mebibyte (MiB)", [
  ("Bit (b)","Other",0.125),("Byte (B)","Other",1),("Kilobit (kb)","Decimal (SI)",125),("Megabit (Mb)","Decimal (SI)",125000),("Gigabit (Gb)","Decimal (SI)",125000000),("Kilobyte (kB)","Decimal (SI)",1e3),("Megabyte (MB)","Decimal (SI)",1e6),
  ("Gigabyte (GB)","Decimal (SI)",1e9),("Terabyte (TB)","Decimal (SI)",1e12),("Kibibyte (KiB)","Binary (IEC)",1024),
  ("Mebibyte (MiB)","Binary (IEC)",1048576),("Gibibyte (GiB)","Binary (IEC)",1073741824),("Tebibyte (TiB)","Binary (IEC)",1099511627776)]),
]


# ---------------------------------------------------------------------------
# Mixed-unit calculators.  Every input has its own unit drop-down, so values
# are typed exactly as measured (e.g. diameter in inches, length in metres).
#   input : (key, label, category or None, default unit, default value, flags)
#           flags: "opt" = blank counts as 0, "delta" = temperature difference
#   result: (label, category or None, default unit, SI expression, flags)
#           expressions use {key} for the input's value in SI base units.
# ---------------------------------------------------------------------------
PIPE_AREA = "PI()/4*{D}^2"
CALCS = {
 "Length": [
  ("Add lengths given in different units", "Total = L1 + L2 + L3",
   [("a", "Length 1", "Length", "Foot (ft)", 5, ""), ("b", "Length 2", "Length", "Meter (m)", 2, ""),
    ("c", "Length 3 (optional)", "Length", "Inch (in)", 6, "opt")],
   [("Total length", "Length", "Meter (m)", "{a}+{b}+{c}", ""), ("Total length", "Length", "Foot (ft)", "{a}+{b}+{c}", "")]),
  ("Pipe length from volume and diameter  (reverse of pipe volume)", "L = V ÷ (π/4 × D²)",
   [("V", "Volume inside pipe", "Volume", "Litre (L)", 50, ""), ("D", "Inside diameter", "Length", "Inch (in)", 4, "")],
   [("Pipe length", "Length", "Meter (m)", "{V}/(" + PIPE_AREA + ")", ""), ("Pipe length", "Length", "Foot (ft)", "{V}/(" + PIPE_AREA + ")", "")]),
  ("Pipe diameter from volume and length  (reverse of pipe volume)", "D = √(4 × V ÷ (π × L))",
   [("V", "Volume inside pipe", "Volume", "US gallon (gal)", 100, ""), ("L", "Pipe length", "Length", "Meter (m)", 12, "")],
   [("Inside diameter", "Length", "Inch (in)", "SQRT(4*{V}/(PI()*{L}))", ""), ("Inside diameter", "Length", "Millimetre (mm)", "SQRT(4*{V}/(PI()*{L}))", "")]),
  ("Distance travelled", "Distance = speed × time",
   [("v", "Speed", "Speed", "Kilometre per hour (km/h)", 60, ""), ("t", "Time", "Time", "Minute (min)", 45, "")],
   [("Distance", "Length", "Mile (mi)", "{v}*{t}", ""), ("Distance", "Length", "Kilometre (km)", "{v}*{t}", "")]),
 ],
 "Area": [
  ("Rectangle area", "A = length × width",
   [("L", "Length", "Length", "Foot (ft)", 12, ""), ("W", "Width", "Length", "Meter (m)", 3, "")],
   [("Area", "Area", "Square meter (m²)", "{L}*{W}", ""), ("Area", "Area", "Square foot (ft²)", "{L}*{W}", "")]),
  ("Circle / pipe cross-section area", "A = π/4 × D²",
   [("D", "Diameter", "Length", "Inch (in)", 6, "")],
   [("Area", "Area", "Square centimetre (cm²)", PIPE_AREA, ""), ("Area", "Area", "Square inch (in²)", PIPE_AREA, "")]),
  ("Pipe outside surface area (painting / insulation)", "A = π × D × L",
   [("D", "Outside diameter", "Length", "Inch (in)", 4, ""), ("L", "Pipe length", "Length", "Meter (m)", 6, "")],
   [("Surface area", "Area", "Square meter (m²)", "PI()*{D}*{L}", ""), ("Surface area", "Area", "Square foot (ft²)", "PI()*{D}*{L}", "")]),
 ],
 "Volume": [
  ("Pipe / cylinder volume", "V = π/4 × D² × L",
   [("D", "Inside diameter", "Length", "Inch (in)", 4, ""), ("L", "Pipe length", "Length", "Meter (m)", 6, "")],
   [("Volume", "Volume", "Litre (L)", PIPE_AREA + "*{L}", ""), ("Volume", "Volume", "US gallon (gal)", PIPE_AREA + "*{L}", ""),
    ("Volume", "Volume", "Cubic meter (m³)", PIPE_AREA + "*{L}", "")]),
  ("Pipe wall (steel) volume", "V = π/4 × (OD² − ID²) × L",
   [("OD", "Outside diameter", "Length", "Inch (in)", 4.5, ""), ("ID", "Inside diameter", "Length", "Inch (in)", 4.026, ""),
    ("L", "Pipe length", "Length", "Meter (m)", 6, "")],
   [("Wall volume", "Volume", "Cubic meter (m³)", "PI()/4*({OD}^2-{ID}^2)*{L}", ""), ("Wall volume", "Volume", "Cubic inch (in³)", "PI()/4*({OD}^2-{ID}^2)*{L}", "")]),
  ("Box / rectangular tank volume", "V = L × W × H",
   [("L", "Length", "Length", "Foot (ft)", 10, ""), ("W", "Width", "Length", "Meter (m)", 2, ""), ("H", "Height", "Length", "Inch (in)", 36, "")],
   [("Volume", "Volume", "Cubic meter (m³)", "{L}*{W}*{H}", ""), ("Volume", "Volume", "US gallon (gal)", "{L}*{W}*{H}", "")]),
  ("Volume delivered by a flow", "V = flow rate × time",
   [("Q", "Flow rate", "Flow Rate", "US gallon per minute (gpm)", 50, ""), ("t", "Time", "Time", "Hour (h)", 2, "")],
   [("Volume", "Volume", "Litre (L)", "{Q}*{t}", ""), ("Volume", "Volume", "Cubic meter (m³)", "{Q}*{t}", "")]),
 ],
 "Mass": [
  ("Mass from density and volume", "m = ρ × V",
   [("rho", "Density", "Density", "Pound per cubic foot (lb/ft³)", 62.4, ""), ("V", "Volume", "Volume", "US gallon (gal)", 100, "")],
   [("Mass", "Mass", "Kilogram (kg)", "{rho}*{V}", ""), ("Mass", "Mass", "Pound (lb)", "{rho}*{V}", "")]),
  ("Pipe weight (empty)", "m = ρ × π/4 × (OD² − ID²) × L   (steel ρ ≈ 7850 kg/m³)",
   [("OD", "Outside diameter", "Length", "Inch (in)", 4.5, ""), ("ID", "Inside diameter", "Length", "Inch (in)", 4.026, ""),
    ("L", "Pipe length", "Length", "Meter (m)", 6, ""), ("rho", "Material density", "Density", "Kilogram per cubic metre (kg/m³)", 7850, "")],
   [("Pipe weight", "Mass", "Kilogram (kg)", "{rho}*PI()/4*({OD}^2-{ID}^2)*{L}", ""), ("Pipe weight", "Mass", "Pound (lb)", "{rho}*PI()/4*({OD}^2-{ID}^2)*{L}", "")]),
  ("Weight of liquid filling a pipe", "m = ρ × π/4 × D² × L   (water ρ ≈ 1000 kg/m³)",
   [("D", "Inside diameter", "Length", "Inch (in)", 4, ""), ("L", "Pipe length", "Length", "Meter (m)", 6, ""),
    ("rho", "Liquid density", "Density", "Kilogram per cubic metre (kg/m³)", 1000, "")],
   [("Liquid weight", "Mass", "Kilogram (kg)", "{rho}*" + PIPE_AREA + "*{L}", ""), ("Liquid weight", "Mass", "Pound (lb)", "{rho}*" + PIPE_AREA + "*{L}", "")]),
  ("Add masses given in different units", "Total = m1 + m2 + m3",
   [("a", "Mass 1", "Mass", "Pound (lb)", 10, ""), ("b", "Mass 2", "Mass", "Kilogram (kg)", 3, ""), ("c", "Mass 3 (optional)", "Mass", "Ounce (oz)", 8, "opt")],
   [("Total mass", "Mass", "Kilogram (kg)", "{a}+{b}+{c}", ""), ("Total mass", "Mass", "Pound (lb)", "{a}+{b}+{c}", "")]),
 ],
 "Temperature": [
  ("Difference and average of two temperatures", "ΔT = T2 − T1 ;  average = (T1 + T2) ÷ 2",
   [("a", "Temperature T1", "Temperature", "Degree Celsius (°C)", 25, ""), ("b", "Temperature T2", "Temperature", "Degree Fahrenheit (°F)", 100, "")],
   [("Difference ΔT (T2 − T1)", "Temperature", "Degree Celsius (°C)", "{b}-{a}", "delta"),
    ("Difference ΔT (T2 − T1)", "Temperature", "Degree Fahrenheit (°F)", "{b}-{a}", "delta"),
    ("Average temperature", "Temperature", "Degree Fahrenheit (°F)", "({a}+{b})/2", "")]),
  ("Temperature after a rise or fall", "T final = T start + ΔT   (ΔT is a temperature difference)",
   [("T", "Starting temperature", "Temperature", "Degree Fahrenheit (°F)", 68, ""), ("dT", "Change ΔT (difference)", "Temperature", "Degree Celsius (°C)", 10, "delta")],
   [("Final temperature", "Temperature", "Degree Celsius (°C)", "{T}+{dT}", ""), ("Final temperature", "Temperature", "Degree Fahrenheit (°F)", "{T}+{dT}", "")]),
 ],
 "Speed": [
  ("Average speed", "v = distance ÷ time",
   [("d", "Distance", "Length", "Mile (mi)", 26.2, ""), ("t", "Time", "Time", "Hour (h)", 3.5, "")],
   [("Speed", "Speed", "Kilometre per hour (km/h)", "{d}/{t}", ""), ("Speed", "Speed", "Metre per second (m/s)", "{d}/{t}", "")]),
  ("Fluid velocity in a pipe", "v = Q ÷ (π/4 × D²)",
   [("Q", "Flow rate", "Flow Rate", "US gallon per minute (gpm)", 100, ""), ("D", "Inside diameter", "Length", "Inch (in)", 3, "")],
   [("Velocity", "Speed", "Metre per second (m/s)", "{Q}/(" + PIPE_AREA + ")", ""), ("Velocity", "Speed", "Foot per second (ft/s)", "{Q}/(" + PIPE_AREA + ")", "")]),
 ],
 "Time": [
  ("Travel time", "t = distance ÷ speed",
   [("d", "Distance", "Length", "Kilometre (km)", 250, ""), ("v", "Speed", "Speed", "Mile per hour (mph)", 55, "")],
   [("Time", "Time", "Hour (h)", "{d}/{v}", ""), ("Time", "Time", "Minute (min)", "{d}/{v}", "")]),
  ("Time to fill a tank", "t = volume ÷ flow rate",
   [("V", "Volume", "Volume", "US gallon (gal)", 500, ""), ("Q", "Flow rate", "Flow Rate", "Litre per minute (L/min)", 40, "")],
   [("Fill time", "Time", "Minute (min)", "{V}/{Q}", ""), ("Fill time", "Time", "Hour (h)", "{V}/{Q}", "")]),
  ("Time to fill a pipe", "t = (π/4 × D² × L) ÷ Q",
   [("D", "Inside diameter", "Length", "Inch (in)", 6, ""), ("L", "Pipe length", "Length", "Meter (m)", 100, ""),
    ("Q", "Flow rate", "Flow Rate", "Litre per second (L/s)", 5, "")],
   [("Fill time", "Time", "Minute (min)", PIPE_AREA + "*{L}/{Q}", ""), ("Fill time", "Time", "Second (s)", PIPE_AREA + "*{L}/{Q}", "")]),
 ],
 "Pressure": [
  ("Pressure from force and area", "P = F ÷ A",
   [("F", "Force", "Force", "Pound-force (lbf)", 500, ""), ("A", "Area", "Area", "Square centimetre (cm²)", 20, "")],
   [("Pressure", "Pressure", "Bar (bar)", "{F}/{A}", ""), ("Pressure", "Pressure", "Pound per square inch (psi)", "{F}/{A}", "")]),
  ("Pressure on a circular piston / pipe end", "P = F ÷ (π/4 × D²)",
   [("F", "Force", "Force", "Kilonewton (kN)", 10, ""), ("D", "Diameter", "Length", "Inch (in)", 2, "")],
   [("Pressure", "Pressure", "Bar (bar)", "{F}/(" + PIPE_AREA + ")", ""), ("Pressure", "Pressure", "Pound per square inch (psi)", "{F}/(" + PIPE_AREA + ")", "")]),
  ("Hydrostatic pressure of a liquid column", "P = ρ × g × h   (g = 9.80665 m/s²)",
   [("rho", "Liquid density", "Density", "Kilogram per cubic metre (kg/m³)", 1000, ""), ("h", "Liquid height", "Length", "Foot (ft)", 33, "")],
   [("Pressure", "Pressure", "Bar (bar)", "{rho}*9.80665*{h}", ""), ("Pressure", "Pressure", "Pound per square inch (psi)", "{rho}*9.80665*{h}", "")]),
 ],
 "Energy": [
  ("Energy used", "E = power × time",
   [("P", "Power", "Power", "Kilowatt (kW)", 2, ""), ("t", "Time", "Time", "Hour (h)", 3, "")],
   [("Energy", "Energy", "Kilowatt-hour (kWh)", "{P}*{t}", ""), ("Energy", "Energy", "British thermal unit (BTU)", "{P}*{t}", "")]),
  ("Work done", "W = force × distance",
   [("F", "Force", "Force", "Pound-force (lbf)", 100, ""), ("d", "Distance", "Length", "Meter (m)", 5, "")],
   [("Work", "Energy", "Joule (J)", "{F}*{d}", ""), ("Work", "Energy", "Foot-pound force (ft·lbf)", "{F}*{d}", "")]),
  ("Potential energy", "E = m × g × h   (g = 9.80665 m/s²)",
   [("m", "Mass", "Mass", "Pound (lb)", 150, ""), ("h", "Height", "Length", "Foot (ft)", 10, "")],
   [("Energy", "Energy", "Joule (J)", "{m}*9.80665*{h}", ""), ("Energy", "Energy", "Kilocalorie (kcal)", "{m}*9.80665*{h}", "")]),
 ],
 "Power": [
  ("Power from energy and time", "P = E ÷ t",
   [("E", "Energy", "Energy", "British thermal unit (BTU)", 12000, ""), ("t", "Time", "Time", "Hour (h)", 1, "")],
   [("Power", "Power", "Kilowatt (kW)", "{E}/{t}", ""), ("Power", "Power", "Horsepower (hp)", "{E}/{t}", "")]),
  ("Shaft power from torque and speed", "P = 2π × N × T",
   [("T", "Torque", "Torque", "Pound-force foot (lbf·ft)", 200, ""), ("N", "Rotational speed", "Frequency", "Revolutions per minute (rpm)", 1800, "")],
   [("Power", "Power", "Kilowatt (kW)", "2*PI()*{N}*{T}", ""), ("Power", "Power", "Horsepower (hp)", "2*PI()*{N}*{T}", "")]),
  ("Pump hydraulic power", "P = flow rate × pressure",
   [("Q", "Flow rate", "Flow Rate", "US gallon per minute (gpm)", 100, ""), ("p", "Pressure (head)", "Pressure", "Pound per square inch (psi)", 50, "")],
   [("Power", "Power", "Kilowatt (kW)", "{Q}*{p}", ""), ("Power", "Power", "Horsepower (hp)", "{Q}*{p}", "")]),
 ],
 "Force": [
  ("Force from pressure and area", "F = P × A",
   [("p", "Pressure", "Pressure", "Pound per square inch (psi)", 100, ""), ("A", "Area", "Area", "Square inch (in²)", 12, "")],
   [("Force", "Force", "Kilonewton (kN)", "{p}*{A}", ""), ("Force", "Force", "Pound-force (lbf)", "{p}*{A}", "")]),
  ("Force on a circular piston / pipe end", "F = P × π/4 × D²",
   [("p", "Pressure", "Pressure", "Bar (bar)", 10, ""), ("D", "Diameter", "Length", "Inch (in)", 3, "")],
   [("Force", "Force", "Kilonewton (kN)", "{p}*" + PIPE_AREA, ""), ("Force", "Force", "Pound-force (lbf)", "{p}*" + PIPE_AREA, "")]),
  ("Weight (force) of a mass", "F = m × g   (g = 9.80665 m/s²)",
   [("m", "Mass", "Mass", "Pound (lb)", 100, "")],
   [("Weight", "Force", "Newton (N)", "{m}*9.80665", ""), ("Weight", "Force", "Kilogram-force (kgf)", "{m}*9.80665", "")]),
 ],
 "Torque": [
  ("Torque from force and lever arm", "T = F × r",
   [("F", "Force", "Force", "Pound-force (lbf)", 50, ""), ("r", "Lever arm length", "Length", "Meter (m)", 0.5, "")],
   [("Torque", "Torque", "Newton metre (N·m)", "{F}*{r}", ""), ("Torque", "Torque", "Pound-force foot (lbf·ft)", "{F}*{r}", "")]),
  ("Torque from power and speed", "T = P ÷ (2π × N)",
   [("P", "Power", "Power", "Horsepower (hp)", 10, ""), ("N", "Rotational speed", "Frequency", "Revolutions per minute (rpm)", 1750, "")],
   [("Torque", "Torque", "Newton metre (N·m)", "{P}/(2*PI()*{N})", ""), ("Torque", "Torque", "Pound-force foot (lbf·ft)", "{P}/(2*PI()*{N})", "")]),
 ],
 "Density": [
  ("Density from mass and volume", "ρ = m ÷ V",
   [("m", "Mass", "Mass", "Pound (lb)", 10, ""), ("V", "Volume", "Volume", "Litre (L)", 5, "")],
   [("Density", "Density", "Kilogram per cubic metre (kg/m³)", "{m}/{V}", ""), ("Density", "Density", "Pound per cubic foot (lb/ft³)", "{m}/{V}", "")]),
  ("Density from specific gravity", "ρ = SG × 1000 kg/m³   (water reference)",
   [("SG", "Specific gravity (no unit)", None, None, 0.85, "")],
   [("Density", "Density", "Kilogram per cubic metre (kg/m³)", "{SG}*1000", ""), ("Density", "Density", "Pound per US gallon (lb/gal)", "{SG}*1000", "")]),
 ],
 "Flow Rate": [
  ("Flow rate from volume and time", "Q = V ÷ t",
   [("V", "Volume", "Volume", "US gallon (gal)", 1000, ""), ("t", "Time", "Time", "Minute (min)", 30, "")],
   [("Flow rate", "Flow Rate", "Cubic metre per hour (m³/h)", "{V}/{t}", ""), ("Flow rate", "Flow Rate", "Litre per second (L/s)", "{V}/{t}", "")]),
  ("Flow rate in a pipe from velocity", "Q = v × π/4 × D²",
   [("v", "Velocity", "Speed", "Foot per second (ft/s)", 6, ""), ("D", "Inside diameter", "Length", "Millimetre (mm)", 100, "")],
   [("Flow rate", "Flow Rate", "Litre per second (L/s)", "{v}*" + PIPE_AREA, ""), ("Flow rate", "Flow Rate", "US gallon per minute (gpm)", "{v}*" + PIPE_AREA, "")]),
 ],
 "Angle": [
  ("Slope angle from rise and run", "angle = ATAN(rise ÷ run) ;  slope % = 100 × rise ÷ run",
   [("rise", "Rise (vertical)", "Length", "Inch (in)", 12, ""), ("run", "Run (horizontal)", "Length", "Meter (m)", 2, "")],
   [("Angle", "Angle", "Degree (°)", "ATAN({rise}/{run})", ""), ("Slope (%)", None, None, "100*{rise}/{run}", "")]),
  ("Arc length", "s = radius × angle (in radians)",
   [("r", "Radius", "Length", "Foot (ft)", 10, ""), ("a", "Angle", "Angle", "Degree (°)", 90, "")],
   [("Arc length", "Length", "Meter (m)", "{r}*{a}", ""), ("Arc length", "Length", "Foot (ft)", "{r}*{a}", "")]),
 ],
 "Frequency": [
  ("Frequency from period", "f = 1 ÷ T",
   [("T", "Period (time of one cycle)", "Time", "Millisecond (ms)", 20, "")],
   [("Frequency", "Frequency", "Hertz (Hz)", "1/{T}", ""), ("Frequency", "Frequency", "Revolutions per minute (rpm)", "1/{T}", "")]),
  ("Rotational speed from surface speed", "N = v ÷ (π × D)",
   [("v", "Surface (cutting) speed", "Speed", "Metre per second (m/s)", 2, ""), ("D", "Diameter", "Length", "Inch (in)", 6, "")],
   [("Speed", "Frequency", "Revolutions per minute (rpm)", "{v}/(PI()*{D})", ""), ("Speed", "Frequency", "Hertz (Hz)", "{v}/(PI()*{D})", "")]),
 ],
 "Data Storage": [
  ("Download / transfer time", "t = file size ÷ connection speed   (speed = amount per second, e.g. Megabit = Mbps)",
   [("S", "File size", "Data Storage", "Gigabyte (GB)", 4, ""), ("R", "Connection speed (per second)", "Data Storage", "Megabit (Mb)", 100, "")],
   [("Transfer time", "Time", "Minute (min)", "{S}/{R}", ""), ("Transfer time", "Time", "Second (s)", "{S}/{R}", "")]),
  ("Add file sizes given in different units", "Total = S1 + S2 + S3",
   [("a", "Size 1", "Data Storage", "Megabyte (MB)", 750, ""), ("b", "Size 2", "Data Storage", "Gibibyte (GiB)", 2, ""),
    ("c", "Size 3 (optional)", "Data Storage", "Kilobyte (kB)", 500, "opt")],
   [("Total size", "Data Storage", "Gigabyte (GB)", "{a}+{b}+{c}", ""), ("Total size", "Data Storage", "Gibibyte (GiB)", "{a}+{b}+{c}", "")]),
 ],
}

F = "Arial"
NAVY = "1F3864"
title_font = Font(name=F, size=18, bold=True, color=NAVY)
sub_font = Font(name=F, size=10, italic=True, color="595959")
sect_font = Font(name=F, size=12, bold=True, color=NAVY)
lbl_font = Font(name=F, size=11, bold=True, color="404040")
input_font = Font(name=F, size=16, bold=True, color="0000FF")
result_font = Font(name=F, size=16, bold=True, color="000000")
unit_font = Font(name=F, size=12, bold=True, color="000000")
hdr_font = Font(name=F, size=11, bold=True, color="FFFFFF")
body = Font(name=F, size=11)
grey = Font(name=F, size=9, color="808080")
blue_in = Font(name=F, size=11, bold=True, color="0000FF")
bold = Font(name=F, size=11, bold=True)
YEL = PatternFill("solid", fgColor="FFF2CC")
YEL_STRONG = PatternFill("solid", fgColor="FFFF00")
GRN = PatternFill("solid", fgColor="E2EFDA")
HDR = PatternFill("solid", fgColor=NAVY)
CALC_HDR = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
med = Side(style="medium", color=NAVY)
B_THIN = Border(left=thin, right=thin, top=thin, bottom=thin)
B_MED = Border(left=med, right=med, top=med, bottom=med)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)
L = Alignment(horizontal="left", vertical="center", wrap_text=True)
LNW = Alignment(horizontal="left", vertical="center")
R = Alignment(horizontal="right", vertical="center")
NUMFMT = "General"


def sig(e):
    """Round to 12 significant digits (hides floating-point noise)."""
    return f"IF(({e})=0,0,ROUND({e},11-INT(LOG10(ABS({e})))))"


def nm(cat):
    return cat.replace(" ", "_")


def base_symbol(base):
    return base.split("(")[-1].rstrip(")")


CAT_BY_NAME = {c[0]: c for c in CATS}
UNITS_BY_CAT = {c[0]: [u[0] for u in c[4]] for c in CATS}
BASE_BY_CAT = {c[0]: base_symbol(c[1]) for c in CATS}
for sheet, calcs in CALCS.items():  # catch typos in default units
    for title, _, inputs, results in calcs:
        for _, _, cat, unit, *_ in inputs:
            assert cat is None or unit in UNITS_BY_CAT[cat], (sheet, title, unit)
        for _, cat, unit, *_ in results:
            assert cat is None or unit in UNITS_BY_CAT[cat], (sheet, title, unit)

CALC_START = 12


def calc_rows(name):
    return sum(2 + len(i) + len(r) + 1 for _, _, i, r in CALCS.get(name, []))


TABLE_TOP = {c[0]: CALC_START + 2 + calc_rows(c[0]) + 4 for c in CATS}  # first unit row

wb = Workbook()
home = wb.active
home.title = "Home"

for name, base, _, _, units in CATS:
    top = TABLE_TOP[name]
    rng = lambda col: f"'{name}'!${col}${top}:${col}${top + len(units) - 1}"
    wb.defined_names[f"U_{nm(name)}"] = DefinedName(f"U_{nm(name)}", attr_text=rng("B"))
    wb.defined_names[f"F_{nm(name)}"] = DefinedName(f"F_{nm(name)}", attr_text=rng("I"))
    wb.defined_names[f"O_{nm(name)}"] = DefinedName(f"O_{nm(name)}", attr_text=rng("J"))


def fac(cat, cell):
    return f"INDEX(F_{nm(cat)},MATCH({cell},U_{nm(cat)},0))"


def off(cat, cell):
    return f"INDEX(O_{nm(cat)},MATCH({cell},U_{nm(cat)},0))"


def build(name, base, d_from, d_to, units):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 38, "C": 18, "D": 24, "E": 10, "F": 22, "G": 34, "H": 3, "I": 20, "J": 12}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    n = len(units)
    TOP = TABLE_TOP[name]
    first, last = TOP, TOP + n - 1
    U, FAC, OFF = f"U_{nm(name)}", f"F_{nm(name)}", f"O_{nm(name)}"
    INP = f"$D${first}:$D${last}"
    dvs = {}

    def unit_dv(cat):
        if cat not in dvs:
            dv = DataValidation(type="list", formula1=f"=U_{nm(cat)}", allow_blank=False,
                                showErrorMessage=True, errorTitle="Unit", error="Pick a unit from the list.")
            ws.add_data_validation(dv)
            dvs[cat] = dv
        return dvs[cat]

    dv_num = DataValidation(type="decimal", operator="between", formula1="-1E+307", formula2="1E+307",
                            allow_blank=True, showErrorMessage=True, errorTitle="Number only", error="Please type a number.")
    ws.add_data_validation(dv_num)

    ws["B1"] = f"{name} Converter"; ws["B1"].font = title_font
    ws["B2"] = (f"SI base unit for this sheet: {base}.  Blue cells on yellow = type here; drop-downs = choose unit.  "
                "← Back to the Home sheet")
    ws["B2"].font = sub_font
    ws["B2"].hyperlink = "#Home!A1"

    # ---- 1. Quick converter (Google style) ----
    ws["B4"] = "1.  QUICK CONVERTER  (works like Google: type a value, pick the two units)"
    ws["B4"].font = sect_font
    ws["B5"] = "Value"; ws["B5"].font = lbl_font; ws["B5"].alignment = R
    ws["C5"] = 1
    ws["D5"] = d_from
    ws["F5"] = "="  # placeholder, replaced below
    raw = f"((C5+{off(name, 'D5')})*{fac(name, 'D5')})/{fac(name, 'G5')}-{off(name, 'G5')}"
    ws["E5"] = "="; ws["E5"].font = Font(name=F, size=18, bold=True, color="404040"); ws["E5"].alignment = C
    ws["F5"] = f'=IF(ISNUMBER(C5),{sig(raw)},"")'
    ws["G5"] = d_to
    for c in ("C5", "D5", "E5", "F5", "G5"):
        ws.merge_cells(f"{c}:{c[0]}6")
    ws["C5"].font = input_font; ws["C5"].fill = YEL_STRONG; ws["C5"].alignment = C; ws["C5"].number_format = NUMFMT
    ws["D5"].font = unit_font; ws["D5"].fill = YEL; ws["D5"].alignment = C
    ws["F5"].font = result_font; ws["F5"].fill = GRN; ws["F5"].alignment = C; ws["F5"].number_format = NUMFMT
    ws["G5"].font = unit_font; ws["G5"].fill = YEL; ws["G5"].alignment = C
    for col in "CDFG":
        for r in (5, 6):
            ws[f"{col}{r}"].border = B_MED
    for c, t in (("C7", "↑ type value"), ("D7", "↑ FROM unit (drop-down)"), ("F7", "↑ result"), ("G7", "↑ TO unit (drop-down)")):
        ws[c] = t; ws[c].font = grey; ws[c].alignment = C
    ws["B8"] = "Formula:"; ws["B8"].font = lbl_font; ws["B8"].alignment = R
    ws["C8"] = '=IF(ISNUMBER(F5),C5&" "&D5&" = "&TEXT(F5,"General")&" "&G5,"Enter a number in the yellow Value box")'
    ws["C8"].font = Font(name=F, size=11, italic=True)
    ws["B9"] = "Reverse:"; ws["B9"].font = lbl_font; ws["B9"].alignment = R
    rev = f"((1+{off(name, 'G5')})*{fac(name, 'G5')})/{fac(name, 'D5')}-{off(name, 'D5')}"
    ws["C9"] = f'="1 "&G5&" = "&TEXT({sig(rev)},"General")&" "&D5'
    ws["C9"].font = Font(name=F, size=11, italic=True)
    ws["B10"] = "Tip: to convert the other way (e.g. foot → meter instead of meter → foot) just swap the two drop-downs."
    ws["B10"].font = grey
    unit_dv(name).add("D5"); unit_dv(name).add("G5"); dv_num.add("C5")

    # ---- 2. Mixed-unit calculators ----
    r = CALC_START
    ws[f"B{r}"] = "2.  MIXED-UNIT CALCULATORS  (enter each value in the unit you have – pick its unit from the drop-down next to it)"
    ws[f"B{r}"].font = sect_font
    r += 1
    for col, t in zip("BCDF", ["Quantity", "Value", "Unit (drop-down)", "Same value in SI base"]):
        c = ws[f"{col}{r}"]; c.value = t; c.font = hdr_font; c.fill = HDR; c.alignment = C; c.border = B_THIN
    ws.merge_cells(f"D{r}:E{r}")
    r += 1
    for title, formula, inputs, results in CALCS[name]:
        ws[f"B{r}"] = title
        for col in "BCDEF":
            ws[f"{col}{r}"].fill = CALC_HDR
        ws[f"B{r}"].font = Font(name=F, size=11, bold=True, color=NAVY)
        ws[f"B{r}"].alignment = LNW
        r += 1
        ws[f"B{r}"] = "Formula:  " + formula; ws[f"B{r}"].font = Font(name=F, size=10, italic=True, color="595959")
        r += 1
        refs, required = {}, []
        for key, label, cat, unit, default, flags in inputs:
            ws[f"B{r}"] = label; ws[f"B{r}"].font = body; ws[f"B{r}"].alignment = R
            vc = ws[f"C{r}"]; vc.value = default; vc.font = blue_in; vc.fill = YEL; vc.alignment = C; vc.number_format = NUMFMT
            dv_num.add(f"C{r}")
            ws.merge_cells(f"D{r}:E{r}")
            blank = "0" if "opt" in flags else '""'
            if cat:
                ws[f"D{r}"] = unit; ws[f"D{r}"].font = bold; ws[f"D{r}"].fill = YEL; ws[f"D{r}"].alignment = LNW
                unit_dv(cat).add(f"D{r}")
                si = f"C{r}*{fac(cat, f'D{r}')}" if "delta" in flags else f"(C{r}+{off(cat, f'D{r}')})*{fac(cat, f'D{r}')}"
                ws[f"F{r}"] = f"=IF(ISNUMBER(C{r}),{si},{blank})"
                ws[f"G{r}"] = BASE_BY_CAT[cat] + (" (difference)" if "delta" in flags else "")
            else:
                ws[f"D{r}"] = "(no unit)"; ws[f"D{r}"].font = grey
                ws[f"F{r}"] = f"=IF(ISNUMBER(C{r}),C{r},{blank})"
            ws[f"F{r}"].font = grey; ws[f"F{r}"].number_format = NUMFMT; ws[f"F{r}"].alignment = C
            ws[f"G{r}"].font = grey
            for col in "BCDEF":
                ws[f"{col}{r}"].border = B_THIN
            refs[key] = f"F{r}"
            if "opt" not in flags:
                required.append(f'F{r}=""')
            r += 1
        for label, cat, unit, expr, flags in results:
            si_expr = expr.format(**refs)
            ws[f"B{r}"] = "⇒ " + label; ws[f"B{r}"].font = bold; ws[f"B{r}"].alignment = R
            cond = f"OR({','.join(required)})" if len(required) > 1 else required[0]
            ws[f"F{r}"] = f'=IF({cond},"",IFERROR({si_expr},""))'
            ws[f"F{r}"].font = grey; ws[f"F{r}"].number_format = NUMFMT; ws[f"F{r}"].alignment = C
            ws.merge_cells(f"D{r}:E{r}")
            if cat:
                ws[f"D{r}"] = unit; ws[f"D{r}"].font = bold; ws[f"D{r}"].fill = YEL; ws[f"D{r}"].alignment = LNW
                unit_dv(cat).add(f"D{r}")
                out = f"F{r}/{fac(cat, f'D{r}')}" + ("" if "delta" in flags else f"-{off(cat, f'D{r}')}")
                ws[f"G{r}"] = BASE_BY_CAT[cat] + (" (difference)" if "delta" in flags else ""); ws[f"G{r}"].font = grey
            else:
                ws[f"D{r}"] = "(no unit)"; ws[f"D{r}"].font = grey
                out = f"F{r}"
            ws[f"C{r}"] = f'=IF(F{r}="","",{sig(out)})'
            ws[f"C{r}"].font = Font(name=F, size=12, bold=True); ws[f"C{r}"].fill = GRN
            ws[f"C{r}"].alignment = C; ws[f"C{r}"].number_format = NUMFMT
            for col in "BCDEF":
                ws[f"{col}{r}"].border = B_THIN
            r += 1
        r += 1

    # ---- 3. All-units table ----
    ws[f"B{TOP-4}"] = "3.  ALL UNITS AT ONCE  (type a value in ANY ONE yellow cell – every other unit is shown instantly)"
    ws[f"B{TOP-4}"].font = sect_font
    msg = ws[f"B{TOP-3}"]
    msg.value = (f'=IF(COUNT({INP})=0,"Type a number next to any unit below (e.g. next to the unit you know).",'
                 f'IF(COUNT({INP})>1,"⚠ More than one value entered – clear all but one yellow cell.",'
                 f'"Showing everything equal to "&INDEX({INP},MATCH(TRUE,INDEX(ISNUMBER({INP}),0),0))&" "'
                 f'&INDEX({U},MATCH(TRUE,INDEX(ISNUMBER({INP}),0),0))))')
    msg.font = Font(name=F, size=11, bold=True, color="C00000")
    ws.merge_cells(f"B{TOP-3}:G{TOP-3}")
    basecell = f"$J${TOP-2}"
    ws[f"I{TOP-2}"] = "Value in SI base:"; ws[f"I{TOP-2}"].font = grey; ws[f"I{TOP-2}"].alignment = R
    bc = ws[basecell.replace("$", "")]
    bc.value = f'=IF(COUNT({INP})<>1,"",SUMPRODUCT(--ISNUMBER({INP}),({INP}+{OFF})*{FAC}))'
    bc.font = grey; bc.number_format = NUMFMT
    hdrs = {"B": "Unit", "C": "System", "D": "Enter value here\n(only ONE row)", "E": "", "F": "Equivalent value",
            "G": "Unit", "I": f"1 unit = ? {base_symbol(base)}\n(factor)", "J": "Offset"}
    for col, t in hdrs.items():
        cell = ws[f"{col}{TOP-1}"]; cell.value = t; cell.font = hdr_font; cell.fill = HDR; cell.alignment = C; cell.border = B_THIN
    ws.row_dimensions[TOP-1].height = 32
    for i, u in enumerate(units):
        r = TOP + i
        label, system, factor = u[0], u[1], u[2]
        offset = u[3] if len(u) > 3 else 0
        ws[f"B{r}"] = label; ws[f"B{r}"].font = bold
        ws[f"C{r}"] = system; ws[f"C{r}"].font = Font(name=F, size=10, color="595959")
        ws[f"D{r}"].fill = YEL; ws[f"D{r}"].font = blue_in; ws[f"D{r}"].number_format = NUMFMT; ws[f"D{r}"].alignment = C
        ws[f"E{r}"] = "→"; ws[f"E{r}"].font = Font(name=F, size=11, color="808080"); ws[f"E{r}"].alignment = C
        ws[f"F{r}"] = f'=IF({basecell}="","",' + sig(f"{basecell}/I{r}-J{r}") + ')'
        ws[f"F{r}"].font = Font(name=F, size=11, bold=True); ws[f"F{r}"].fill = GRN
        ws[f"F{r}"].number_format = NUMFMT; ws[f"F{r}"].alignment = C
        ws[f"G{r}"] = f"=B{r}"; ws[f"G{r}"].font = body
        ws[f"I{r}"] = factor; ws[f"I{r}"].font = grey; ws[f"I{r}"].number_format = NUMFMT
        ws[f"J{r}"] = offset; ws[f"J{r}"].font = grey; ws[f"J{r}"].number_format = NUMFMT
        for col in "BCDEFGIJ":
            ws[f"{col}{r}"].border = B_THIN
        ws.row_dimensions[r].height = 20
    dv_num.add(f"D{first}:D{last}")
    ws.conditional_formatting.add(f"B{first}:G{last}",
        FormulaRule(formula=[f"ISNUMBER($D{first})"], font=Font(name=F, bold=True, color="C00000")))

    nr = last + 2
    notes = [
        "How it works: every value is converted to the SI base unit, calculated, then converted to the unit you pick:  "
        "SI = (value + offset) × factor;  result = SI ÷ factor − offset.",
        "Offsets are only used for temperature (°C and °F are not zero-based). All factors are exact by international "
        "definition (NIST SP 811 / BIPM SI Brochure) unless noted.",
        "Grey columns I:J are the conversion constants used by every sheet (named ranges U_, F_, O_ + sheet name).",
    ]
    extra = {
        "Volume": "US cup = 236.5882365 mL (customary). US and Imperial gallons/pints/fl oz differ – pick the right one.",
        "Time": "Month and year are Gregorian averages (365.2425 days per year).",
        "Energy": "BTU = International Table BTU; calorie = thermochemical calorie (4.184 J).",
        "Pressure": "inHg at 0 °C (3386.389 Pa); mmHg = 133.322387415 Pa. Pressures are as entered (gauge or absolute).",
        "Data Storage": "Decimal units (kB, MB, GB…) use powers of 1000; binary units (KiB, MiB, GiB…) use powers of 1024. "
                        "1 byte = 8 bits; internet speeds are usually quoted in Megabit (Mb) per second.",
        "Temperature": "A temperature DIFFERENCE (ΔT) ignores the offset: a change of 10 °C = a change of 18 °F, "
                       "but a temperature of 10 °C = 50 °F.",
    }
    if name in extra:
        notes.append(extra[name])
    for j, t in enumerate(notes):
        c = ws[f"B{nr + j}"]; c.value = t; c.font = grey
    ws.freeze_panes = "A4"


for cat in CATS:
    build(*cat)

# ---- Home sheet ----
home.sheet_view.showGridLines = False
for k, v in {"A": 2, "B": 6, "C": 22, "D": 36, "E": 60, "F": 60}.items():
    home.column_dimensions[k].width = v
home["B1"] = "Unit Converter – SI ⇄ Imperial / US / Other"; home["B1"].font = Font(name=F, size=20, bold=True, color=NAVY)
home["B2"] = "One sheet per quantity. Works like the Google unit converter – no macros needed."; home["B2"].font = sub_font
how = [
    "HOW TO USE  –  every sheet has three sections",
    "1.  QUICK CONVERTER: type a number in the bright-yellow Value box, choose the FROM unit and the TO unit from the drop-downs – the green box shows the answer.",
    "     To go the other way (e.g. foot → meter instead of meter → foot) simply swap the two drop-downs.",
    "2.  MIXED-UNIT CALCULATORS: enter each value exactly as you have it, each with its own unit – e.g. pipe volume with the diameter in inches and the length",
    "     in metres. Pick any unit for the answer too. Reverse calculators are included (e.g. Length sheet: pipe length or diameter from a known volume).",
    "3.  ALL UNITS AT ONCE: type a number next to ANY unit in the yellow column (e.g. 10 next to Foot) – every other unit updates instantly.",
    "     Delete it and type in another row (e.g. 3 next to Meter) to go the other way. Keep only ONE value in that column.",
    "Colour key:  blue text on yellow = your input / unit choice  ·  green = calculated result  ·  grey = SI values and conversion constants.",
]
for i, t in enumerate(how):
    c = home[f"B{4+i}"]; c.value = t
    c.font = Font(name=F, size=12, bold=True, color=NAVY) if i == 0 else Font(name=F, size=11)
r0 = 4 + len(how) + 1
for col, t in zip("BCDEF", ["#", "Sheet (click)", "SI base unit", "Mixed-unit calculators", "Units included"]):
    c = home[f"{col}{r0}"]; c.value = t; c.font = hdr_font; c.fill = HDR; c.alignment = C; c.border = B_THIN
for i, (name, base, _, _, units) in enumerate(CATS):
    r = r0 + 1 + i
    home[f"B{r}"] = i + 1
    home[f"C{r}"] = name
    home[f"C{r}"].hyperlink = f"#'{name}'!A1"
    home[f"D{r}"] = base
    home[f"E{r}"] = "\n".join("• " + c[0] for c in CALCS[name])
    home[f"F{r}"] = ", ".join(u[0] for u in units)
    for col in "BCDEF":
        c = home[f"{col}{r}"]; c.border = B_THIN
        c.font = Font(name=F, size=10)
        c.alignment = L if col in "DEF" else C
    home[f"C{r}"].font = Font(name=F, size=11, bold=True, color="0563C1", underline="single")
    home.row_dimensions[r].height = max(30, 15 * len(CALCS[name]) + 6, 15 * (len(home[f"F{r}"].value) // 70 + 1) + 6)

import sys
out = sys.argv[1] if len(sys.argv) > 1 else "Unit_Converter.xlsx"
wb.save(out)
print("saved", out)

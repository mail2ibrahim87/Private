"""Generates Cleaning_Decon_Volume_Template.xlsx – volume & chemical estimate template for
chemical-cleaning / decontamination proposals (exchangers, vessels/columns/tanks, piping; 5 trains).

Usage:  python3 tools/build_cleaning_template.py Cleaning_Decon_Volume_Template.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

N_HX, N_VES, N_PIPE, N_TRAINS, N_STEPS = 20, 20, 10, 5, 8
TR = ["D", "E", "F", "G", "H"]          # train columns on the calculation sheet
HELP = ["L", "M", "N", "O", "P"]        # hidden helper columns (equipment type per train)

# ---------------------------------------------------------------- styles
FN = "Arial"
NAVY, BLUE_HDR = "1F3864", "D9E1F2"
f_title = Font(name=FN, size=16, bold=True, color=NAVY)
f_sub = Font(name=FN, size=10, italic=True, color="595959")
f_sect = Font(name=FN, size=12, bold=True, color="FFFFFF")
f_blk = Font(name=FN, size=11, bold=True, color=NAVY)
f_hdr = Font(name=FN, size=10, bold=True, color="FFFFFF")
f_body = Font(name=FN, size=10)
f_bold = Font(name=FN, size=10, bold=True)
f_in = Font(name=FN, size=10, color="0000FF")
f_grey = Font(name=FN, size=9, color="808080")
f_res = Font(name=FN, size=10, bold=True)
f_red = Font(name=FN, size=11, bold=True, color="C00000")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_RES = PatternFill("solid", fgColor="E2EFDA")
FILL_KEY = PatternFill("solid", fgColor="C6E0B4")
FILL_SECT = PatternFill("solid", fgColor=NAVY)
FILL_BLK = PatternFill("solid", fgColor=BLUE_HDR)
FILL_SUBH = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
AL_LN = Alignment(horizontal="left", vertical="center")
AL_R = Alignment(horizontal="right", vertical="center")
NF_M3 = '#,##0.000;-#,##0.000;"-"'
NF_M2 = '#,##0.0000;-#,##0.0000;"-"'
NF_L = '#,##0;-#,##0;"-"'
NF_MM = '#,##0.00;-#,##0.00;"-"'
NF_PCT = '0%;-0%;"-"'
NF_PCT2 = '0.0%'

# When run via runpy with init_globals={"TARGET_WB": wb, "SHEET_NAMES": {...}} the sheets are added to
# an existing workbook (used by build_pricing.py); otherwise a standalone template is created.
SHEET_NAMES = globals().get("SHEET_NAMES") or {"set": "Settings", "vc": "Volume Calculation", "sum": "Summary", "lk": "Lookup"}
STANDALONE = globals().get("TARGET_WB") is None
q = lambda n: f"'{n}'" if " " in n else n
SS, SVC, SSUM, SL = (q(SHEET_NAMES[k]) for k in ("set", "vc", "sum", "lk"))
if STANDALONE:
    wb = Workbook()
    ws_set = wb.active
    ws_set.title = SHEET_NAMES["set"]
else:
    wb = TARGET_WB
    ws_set = wb.create_sheet(SHEET_NAMES["set"])
ws = wb.create_sheet(SHEET_NAMES["vc"])
ws_sum = wb.create_sheet(SHEET_NAMES["sum"])
ws_lk = wb.create_sheet(SHEET_NAMES["lk"])
ITEMS = []  # (kind, item no, tag row, considered-volume row)


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def style(c, font=f_body, fill=None, nf=None, al=None, border=True):
    c.font = font
    if fill: c.fill = fill
    if nf: c.number_format = nf
    if al: c.alignment = al
    if border: c.border = BORDER


# ================================================================ LOOKUP sheet
LEN_U = [("mm", 0.001), ("cm", 0.01), ("m", 1), ("in", 0.0254), ("ft", 0.3048)]
VOL_U = [("m³", 1), ("L", 0.001), ("US gal", 0.003785411784), ("Imp gal", 0.00454609),
         ("bbl", 0.158987294928), ("ft³", 0.028316846592)]
WT_U = [("kg", 1), ("t", 1000), ("lb", 0.45359237)]
HEADS = [("None / Open", "0", "No head volume"), ("Flat", "0", "Flat cover / blind – no extra volume"),
         ("2:1 Ellipsoidal", "=PI()/24", "V = π/24 × D³"), ("Hemispherical", "=PI()/12", "V = π/12 × D³"),
         ("Torispherical (F&D)", "0.0809", "ASME flanged & dished, V ≈ 0.0809 × D³"),
         ("Conical", "0", "V = π/12 × D² × cone height (enter cone height)")]
LAYOUTS = [("Triangular (30°)", "=SQRT(3)/2"), ("Rotated triangular (60°)", "=SQRT(3)/2"),
           ("Square (90°)", 1), ("Rotated square (45°)", 1)]
BASIS = ["Dimensions", "Water weight"]
JOBS = ["Chemical Cleaning", "Decontamination"]

# ASME B36.10M / B36.19M – OD (mm) and wall thickness (mm)
SCHEDS = ["Sch 5S", "Sch 10S", "Sch 10", "Sch 20", "Sch 30", "Sch 40", "STD", "Sch 40S", "Sch 60",
          "Sch 80", "XS", "Sch 80S", "Sch 100", "Sch 120", "Sch 140", "Sch 160", "XXS"]
PIPES = [  # NPS label, OD, {schedule: wall}
 ('1/2" (DN 15)', 21.3, {"Sch 5S": 1.65, "Sch 10S": 2.11, "Sch 40": 2.77, "STD": 2.77, "Sch 40S": 2.77, "Sch 80": 3.73, "XS": 3.73, "Sch 80S": 3.73, "Sch 160": 4.78, "XXS": 7.47}),
 ('3/4" (DN 20)', 26.7, {"Sch 5S": 1.65, "Sch 10S": 2.11, "Sch 40": 2.87, "STD": 2.87, "Sch 40S": 2.87, "Sch 80": 3.91, "XS": 3.91, "Sch 80S": 3.91, "Sch 160": 5.56, "XXS": 7.82}),
 ('1" (DN 25)', 33.4, {"Sch 5S": 1.65, "Sch 10S": 2.77, "Sch 40": 3.38, "STD": 3.38, "Sch 40S": 3.38, "Sch 80": 4.55, "XS": 4.55, "Sch 80S": 4.55, "Sch 160": 6.35, "XXS": 9.09}),
 ('1-1/4" (DN 32)', 42.2, {"Sch 5S": 1.65, "Sch 10S": 2.77, "Sch 40": 3.56, "STD": 3.56, "Sch 40S": 3.56, "Sch 80": 4.85, "XS": 4.85, "Sch 80S": 4.85, "Sch 160": 6.35, "XXS": 9.70}),
 ('1-1/2" (DN 40)', 48.3, {"Sch 5S": 1.65, "Sch 10S": 2.77, "Sch 40": 3.68, "STD": 3.68, "Sch 40S": 3.68, "Sch 80": 5.08, "XS": 5.08, "Sch 80S": 5.08, "Sch 160": 7.14, "XXS": 10.15}),
 ('2" (DN 50)', 60.3, {"Sch 5S": 1.65, "Sch 10S": 2.77, "Sch 40": 3.91, "STD": 3.91, "Sch 40S": 3.91, "Sch 80": 5.54, "XS": 5.54, "Sch 80S": 5.54, "Sch 160": 8.74, "XXS": 11.07}),
 ('2-1/2" (DN 65)', 73.0, {"Sch 5S": 2.11, "Sch 10S": 3.05, "Sch 40": 5.16, "STD": 5.16, "Sch 40S": 5.16, "Sch 80": 7.01, "XS": 7.01, "Sch 80S": 7.01, "Sch 160": 9.53, "XXS": 14.02}),
 ('3" (DN 80)', 88.9, {"Sch 5S": 2.11, "Sch 10S": 3.05, "Sch 40": 5.49, "STD": 5.49, "Sch 40S": 5.49, "Sch 80": 7.62, "XS": 7.62, "Sch 80S": 7.62, "Sch 160": 11.13, "XXS": 15.24}),
 ('4" (DN 100)', 114.3, {"Sch 5S": 2.11, "Sch 10S": 3.05, "Sch 40": 6.02, "STD": 6.02, "Sch 40S": 6.02, "Sch 80": 8.56, "XS": 8.56, "Sch 80S": 8.56, "Sch 120": 11.13, "Sch 160": 13.49, "XXS": 17.12}),
 ('5" (DN 125)', 141.3, {"Sch 5S": 2.77, "Sch 10S": 3.40, "Sch 40": 6.55, "STD": 6.55, "Sch 40S": 6.55, "Sch 80": 9.53, "XS": 9.53, "Sch 80S": 9.53, "Sch 120": 12.70, "Sch 160": 15.88, "XXS": 19.05}),
 ('6" (DN 150)', 168.3, {"Sch 5S": 2.77, "Sch 10S": 3.40, "Sch 40": 7.11, "STD": 7.11, "Sch 40S": 7.11, "Sch 80": 10.97, "XS": 10.97, "Sch 80S": 10.97, "Sch 120": 14.27, "Sch 160": 18.26, "XXS": 21.95}),
 ('8" (DN 200)', 219.1, {"Sch 5S": 2.77, "Sch 10S": 3.76, "Sch 20": 6.35, "Sch 30": 7.04, "Sch 40": 8.18, "STD": 8.18, "Sch 40S": 8.18, "Sch 60": 10.31, "Sch 80": 12.70, "XS": 12.70, "Sch 80S": 12.70, "Sch 100": 15.09, "Sch 120": 18.26, "Sch 140": 20.62, "Sch 160": 23.01, "XXS": 22.23}),
 ('10" (DN 250)', 273.0, {"Sch 5S": 3.40, "Sch 10S": 4.19, "Sch 20": 6.35, "Sch 30": 7.80, "Sch 40": 9.27, "STD": 9.27, "Sch 40S": 9.27, "Sch 60": 12.70, "XS": 12.70, "Sch 80S": 12.70, "Sch 80": 15.09, "Sch 100": 18.26, "Sch 120": 21.44, "Sch 140": 25.40, "Sch 160": 28.58, "XXS": 25.40}),
 ('12" (DN 300)', 323.8, {"Sch 5S": 3.96, "Sch 10S": 4.57, "Sch 20": 6.35, "Sch 30": 8.38, "STD": 9.53, "Sch 40S": 9.53, "Sch 40": 10.31, "XS": 12.70, "Sch 80S": 12.70, "Sch 60": 14.27, "Sch 80": 17.48, "Sch 100": 21.44, "Sch 120": 25.40, "Sch 140": 28.58, "Sch 160": 33.32, "XXS": 25.40}),
 ('14" (DN 350)', 355.6, {"Sch 5S": 3.96, "Sch 10S": 4.78, "Sch 10": 6.35, "Sch 20": 7.92, "Sch 30": 9.53, "STD": 9.53, "Sch 40": 11.13, "XS": 12.70, "Sch 60": 15.09, "Sch 80": 19.05, "Sch 100": 23.83, "Sch 120": 27.79, "Sch 140": 31.75, "Sch 160": 35.71}),
 ('16" (DN 400)', 406.4, {"Sch 5S": 4.19, "Sch 10S": 4.78, "Sch 10": 6.35, "Sch 20": 7.92, "Sch 30": 9.53, "STD": 9.53, "Sch 40": 12.70, "XS": 12.70, "Sch 60": 16.66, "Sch 80": 21.44, "Sch 100": 26.19, "Sch 120": 30.96, "Sch 140": 36.53, "Sch 160": 40.49}),
 ('18" (DN 450)', 457.0, {"Sch 5S": 4.19, "Sch 10S": 4.78, "Sch 10": 6.35, "Sch 20": 7.92, "STD": 9.53, "Sch 30": 11.13, "XS": 12.70, "Sch 40": 14.27, "Sch 60": 19.05, "Sch 80": 23.83, "Sch 100": 29.36, "Sch 120": 34.93, "Sch 140": 39.67, "Sch 160": 45.24}),
 ('20" (DN 500)', 508.0, {"Sch 5S": 4.78, "Sch 10S": 5.54, "Sch 10": 6.35, "Sch 20": 9.53, "STD": 9.53, "Sch 30": 12.70, "XS": 12.70, "Sch 40": 15.09, "Sch 60": 20.62, "Sch 80": 26.19, "Sch 100": 32.54, "Sch 120": 38.10, "Sch 140": 44.45, "Sch 160": 50.01}),
 ('24" (DN 600)', 610.0, {"Sch 5S": 5.54, "Sch 10S": 6.35, "Sch 10": 6.35, "Sch 20": 9.53, "STD": 9.53, "XS": 12.70, "Sch 30": 14.27, "Sch 40": 17.48, "Sch 60": 24.61, "Sch 80": 30.96, "Sch 100": 38.89, "Sch 120": 46.02, "Sch 140": 52.37, "Sch 160": 59.54}),
 ('30" (DN 750)', 762.0, {"Sch 5S": 6.35, "Sch 10S": 7.92, "Sch 10": 7.92, "STD": 9.53, "Sch 20": 12.70, "XS": 12.70, "Sch 30": 15.88}),
 ('36" (DN 900)', 914.0, {"Sch 10": 7.92, "STD": 9.53, "Sch 20": 12.70, "XS": 12.70, "Sch 30": 15.88, "Sch 40": 19.05}),
]

lk = ws_lk
lk["A1"] = "Lookup tables used by the template (edit with care)"; lk["A1"].font = f_blk


def lk_table(col, title, rows, hdrs, nm_prefix, numcols):
    c0 = col
    lk.cell(row=3, column=c0, value=title).font = f_bold
    for j, h in enumerate(hdrs):
        style(lk.cell(row=4, column=c0 + j, value=h), f_hdr, FILL_SECT, al=AL_C)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            style(lk.cell(row=5 + i, column=c0 + j, value=v))
    first, last = 5, 4 + len(rows)
    for j, n in enumerate(nm_prefix):
        if n:
            L = get_column_letter(c0 + j)
            name(n, f"{SL}!${L}${first}:${L}${last}")
    return last


lk_table(1, "Length units", LEN_U, ["Unit", "× to m"], ["LenU", "LenF"], 2)
lk_table(4, "Volume units", VOL_U, ["Unit", "× to m³"], ["VolU", "VolF"], 2)
lk_table(7, "Weight units", WT_U, ["Unit", "× to kg"], ["WtU", "WtF"], 2)
lk_table(10, "Head types", HEADS, ["Head type", "k (V = k·D³)", "Formula"], ["HeadU", "HeadK", None], 3)
lk_table(14, "Tube layout", LAYOUTS, ["Layout", "Area factor (× pitch²)"], ["LayoutU", "LayoutF"], 2)
lk_table(17, "Volume basis", [(b,) for b in BASIS], ["Basis"], ["BasisU"], 1)
lk_table(19, "Job type", [(j,) for j in JOBS], ["Job type"], ["JobU"], 1)
for w, cols in ((6, "AB"), (2, "C"), (10, "DE"), (2, "F"), (10, "GH"), (2, "I"), (22, "J"), (14, "K"), (40, "L"), (2, "M"), (24, "N"), (14, "O"), (2, "P"), (14, "Q"), (2, "R"), (18, "S")):
    for c in cols:
        lk.column_dimensions[c].width = w
lk.column_dimensions["A"].width = 8; lk.column_dimensions["B"].width = 8

# pipe table
pr = 14
lk.cell(row=pr, column=1, value="Pipe wall thickness (mm) – ASME B36.10M / B36.19M.  Verify against the project piping class; "
        "blank = schedule not standard for that size.").font = f_bold
style(lk.cell(row=pr + 1, column=1, value="NPS"), f_hdr, FILL_SECT, al=AL_C)
style(lk.cell(row=pr + 1, column=2, value="OD (mm)"), f_hdr, FILL_SECT, al=AL_C)
for j, s in enumerate(SCHEDS):
    style(lk.cell(row=pr + 1, column=3 + j, value=s), f_hdr, FILL_SECT, al=AL_C)
for i, (nps, od, wt) in enumerate(PIPES):
    r = pr + 2 + i
    style(lk.cell(row=r, column=1, value=nps), f_bold)
    style(lk.cell(row=r, column=2, value=od))
    for j, s in enumerate(SCHEDS):
        style(lk.cell(row=r, column=3 + j, value=wt.get(s)), f_body)
p_first, p_last = pr + 2, pr + 1 + len(PIPES)
lastcol = get_column_letter(2 + len(SCHEDS))
name("NPSU", f"{SL}!$A${p_first}:$A${p_last}")
name("PipeOD", f"{SL}!$B${p_first}:$B${p_last}")
name("SchU", f"{SL}!$C${pr+1}:${lastcol}${pr+1}")
name("PipeWT", f"{SL}!$C${p_first}:${lastcol}${p_last}")
lk.column_dimensions["A"].width = 16

# ================================================================ SETTINGS sheet
s = ws_set
s.sheet_view.showGridLines = False
for c, w in {"A": 2, "B": 34, "C": 26, "D": 44, "E": 14, "F": 16, "G": 56}.items():
    s.column_dimensions[c].width = w
s["B1"] = "Chemical Cleaning / Decontamination – Volume Estimate Template"; s["B1"].font = f_title
s["B2"] = (f"Fill in the yellow cells on this sheet first, then the '{SHEET_NAMES['vc']}' sheet. "
           "The 'Summary' sheet is the proposal output."); s["B2"].font = f_sub

r = 4
s[f"B{r}"] = "PROJECT INFORMATION"; s[f"B{r}"].font = f_blk
info = ["Client", "Plant / Unit", "Proposal / Project No.", "Location", "Prepared by", "Date", "Revision"]
for i, t in enumerate(info):
    rr = r + 1 + i
    style(s[f"B{rr}"], f_bold, al=AL_LN); s[f"B{rr}"].value = t
    style(s[f"C{rr}"], f_in, FILL_IN, al=AL_LN)
    s.merge_cells(f"C{rr}:D{rr}")
name("ProjClient", f"{SS}!$C${r+1}"); name("ProjPlant", f"{SS}!$C${r+2}"); name("ProjNo", f"{SS}!$C${r+3}")
s[f"C{r+6}"].number_format = "dd-mmm-yyyy"

r = 13
s[f"B{r}"] = "TRAIN / UNIT NAMES  (shown as column headings)"; s[f"B{r}"].font = f_blk
for i in range(N_TRAINS):
    rr = r + 1 + i
    style(s[f"B{rr}"], f_bold, al=AL_LN); s[f"B{rr}"].value = f"Train {i+1} name"
    style(s[f"C{rr}"], f_in, FILL_IN, al=AL_LN); s[f"C{rr}"].value = f"Train {i+1}"
    name(f"TrainName{i+1}", f"{SS}!$C${rr}")

r = 20
s[f"B{r}"] = "JOB TYPE  (drives the whole calculation)"; s[f"B{r}"].font = f_blk
style(s[f"B{r+1}"], f_bold, al=AL_LN); s[f"B{r+1}"].value = "Job type"
style(s[f"C{r+1}"], Font(name=FN, size=12, bold=True, color="0000FF"), PatternFill("solid", fgColor="FFFF00"), al=AL_C)
s[f"C{r+1}"].value = "Chemical Cleaning"
name("JobType", f"{SS}!$C${r+1}")
dv = DataValidation(type="list", formula1="=JobU", allow_blank=False); s.add_data_validation(dv); dv.add(f"C{r+1}")

r = 23
s[f"B{r}"] = "GENERAL PARAMETERS"; s[f"B{r}"].font = f_blk
gp = [("Water density for water-weight → volume", 1000, "kg/m³", "WaterRho", "Volume = water weight ÷ this density"),
      ("Tube pitch ratio (pitch ÷ tube OD)", 1.25, "–", "PitchRatio", "Used to estimate tube OD when only the pitch is given (TEMA typical 1.25)"),
      ("Drum size", 200, "L", "DrumL", "Chemical containers – used for the drum count on the Summary"),
      ("IBC / tote size", 1000, "L", "IBCL", "Chemical containers – used for the IBC count on the Summary")]
for i, (t, v, u, n, note) in enumerate(gp):
    rr = r + 1 + i
    style(s[f"B{rr}"], f_bold, al=AL_LN); s[f"B{rr}"].value = t
    style(s[f"C{rr}"], f_in, FILL_IN, al=AL_C); s[f"C{rr}"].value = v
    style(s[f"D{rr}"], f_body, al=AL_LN); s[f"D{rr}"].value = u
    s[f"G{rr}"] = note; s[f"G{rr}"].font = f_grey
    name(n, f"{SS}!$C${rr}")

r = 29
s[f"B{r}"] = "METHODOLOGY – % of equipment volume considered and chemical dosage (edit per project)"; s[f"B{r}"].font = f_blk
hd = ["Equipment", "Chemical cleaning – fill %", "Decontamination – method", "Decon – fill %", "Decon – chemical % vol", "Remarks"]
for j, h in enumerate(hd):
    style(s.cell(row=r + 1, column=2 + j, value=h), f_hdr, FILL_SECT, al=AL_C)
s.row_dimensions[r + 1].height = 30
METH = [
    ("Exchanger", 1, "Vapour phase – complete volume (same as column)", 1, 0.01, "CC: full volume, tube side + shell side"),
    ("Piping", 1, "Vapour phase – complete volume (same as column)", 1, 0.01, "CC: full volume"),
    ("Column", "per item ▼", "Vapour phase – column filled for complete volume", 1, 0.01, "CC: Full fill or 20% + gamma jet – chosen per item"),
    ("Vessel – Vertical", "per item ▼", "Boil-out – water fill, steam injected at bottom (rumbling)", 0.3, 0.02, "CC: Full fill or 20% + gamma jet – chosen per item"),
    ("Vessel – Horizontal", "per item ▼", "Boil-out – water fill, steam injected at bottom (rumbling)", 0.3, 0.02, "CC: Full fill or 20% + gamma jet – chosen per item"),
    ("Tank", "per item ▼", "Gamma-jet circulation (heating up to 80 °C if required)", 0.3, 0.02, "⚠ TO CONFIRM: decon fill % and chemical % for tanks (set equal to boil-out)"),
]
for i, row in enumerate(METH):
    rr = r + 2 + i
    for j, v in enumerate(row):
        c = s.cell(row=rr, column=2 + j, value=v)
        if j == 0:
            style(c, f_bold, al=AL_LN)
        elif j == 5:
            style(c, Font(name=FN, size=9, color="C00000" if "TO CONFIRM" in v else "595959"), al=AL_L)
        elif j == 1 and isinstance(v, str):
            style(c, f_grey, al=AL_C)
        else:
            style(c, f_in, FILL_IN, al=AL_C if j != 2 else AL_L, nf=NF_PCT2 if j in (1, 3, 4) else None)
    s.row_dimensions[rr].height = 28
m1, m2 = r + 2, r + 1 + len(METH)
name("DecTypeU", f"{SS}!$B${m1}:$B${m2}")
name("DecMeth", f"{SS}!$D${m1}:$D${m2}")
name("DecFill", f"{SS}!$E${m1}:$E${m2}")
name("DecChem", f"{SS}!$F${m1}:$F${m2}")
name("EqTypeU", f"{SS}!$B${m1+2}:$B${m2}")          # Column … Tank (vessel-type equipment)
name("CC_HX", f"{SS}!$C${m1}")
name("CC_PIPE", f"{SS}!$C${m1+1}")

r = m2 + 2
s[f"B{r}"] = "CHEMICAL CLEANING – method options for vessels / columns / tanks (chosen per item)"; s[f"B{r}"].font = f_blk
for j, h in enumerate(["Method", "% of volume considered", "Remarks"]):
    style(s.cell(row=r + 1, column=2 + j, value=h), f_hdr, FILL_SECT, al=AL_C)
CCM = [("Full fill", 1, "Equipment completely filled with cleaning solution"),
       ("20% + gamma jet circulation", 0.2, "Large equipment: 20% of volume circulated through gamma jetting nozzle")]
for i, (a, b, c_) in enumerate(CCM):
    rr = r + 2 + i
    style(s.cell(row=rr, column=2, value=a), f_bold, al=AL_LN)
    style(s.cell(row=rr, column=3, value=b), f_in, FILL_IN, nf=NF_PCT2, al=AL_C)
    style(s.cell(row=rr, column=4, value=c_), Font(name=FN, size=9, color="595959"), al=AL_L)
name("CCMethU", f"{SS}!$B${r+2}:$B${r+1+len(CCM)}")
name("CCMethPct", f"{SS}!$C${r+2}:$C${r+1+len(CCM)}")

r = r + 2 + len(CCM) + 1
s[f"B{r}"] = "CHEMICAL CLEANING – chemical steps (applied to the total considered volume)"; s[f"B{r}"].font = f_blk
for j, h in enumerate(["Step / chemical", "Concentration % vol", "No. of fills / batches", "", "", "Remarks"]):
    if h:
        style(s.cell(row=r + 1, column=2 + j, value=h), f_hdr, FILL_SECT, al=AL_C)
st1 = r + 2
for i in range(N_STEPS):
    rr = st1 + i
    style(s.cell(row=rr, column=2), f_in, FILL_IN, al=AL_LN)
    style(s.cell(row=rr, column=3), f_in, FILL_IN, nf=NF_PCT2, al=AL_C)
    style(s.cell(row=rr, column=4), f_in, FILL_IN, al=AL_C)
    style(s.cell(row=rr, column=7), Font(name=FN, size=9, color="595959"), FILL_IN, al=AL_L)
s.cell(row=st1, column=2, value="EXAMPLE – Alkaline degreasing")
s.cell(row=st1, column=3, value=0.02); s.cell(row=st1, column=4, value=1)
s.cell(row=st1, column=7, value="Example row – replace with the project's chemical programme")
name("StepName", f"{SS}!$B${st1}:$B${st1+N_STEPS-1}")
name("StepConc", f"{SS}!$C${st1}:$C${st1+N_STEPS-1}")
name("StepFills", f"{SS}!$D${st1}:$D${st1+N_STEPS-1}")
STEP_ROWS = list(range(st1, st1 + N_STEPS))
r = st1 + N_STEPS + 1
notes = [
    "LEGEND:  yellow = input (blue text) · green = calculated · drop-downs are provided for units, head types, layout, basis and methods.",
    f"All equipment volumes are calculated in m³. Each input row on '{SHEET_NAMES['vc']}' has its own unit drop-down (mm, cm, m, in, ft / kg, t, lb / m³, L, gal …).",
    "Head volumes exclude the straight flange; torispherical = ASME F&D approximation. Internals (trays, packing, demisters) are not deducted.",
]
for i, t in enumerate(notes):
    s[f"B{r+i}"] = t; s[f"B{r+i}"].font = f_grey

# ================================================================ VOLUME CALCULATION sheet
ws.sheet_view.showGridLines = False
for c, w in {"A": 5, "B": 50, "C": 9, "D": 17, "E": 17, "F": 17, "G": 17, "H": 17, "I": 17, "J": 58, "K": 9}.items():
    ws.column_dimensions[c].width = w
for c in HELP:
    ws.column_dimensions[c].width = 14
    ws.column_dimensions[c].hidden = True
ws.column_dimensions["K"].hidden = True

ws["B1"] = "VOLUME CALCULATION – Exchangers, Vessels / Columns / Tanks, Piping"; ws["B1"].font = f_title
ws["B2"] = '="Client: "&ProjClient&"   ·   Plant: "&ProjPlant&"   ·   Proposal: "&ProjNo'; ws["B2"].font = f_sub
ws["B3"] = '="JOB TYPE: "&UPPER(JobType)&"   (change on the Settings sheet)"'; ws["B3"].font = f_red
ws["B4"] = ("Enter data per train in the yellow cells (blank = not applicable). Pick the unit of each row in column C. "
            "Green rows are calculated in m³.  Tip: copy a Train column to the next train when trains are identical.")
ws["B4"].font = f_grey
HR = 6
heads = ["#", "Parameter", "Unit"] + [f"=TrainName{i+1}" for i in range(N_TRAINS)] + ["Total (all trains)", "Notes / source"]
for j, h in enumerate(heads):
    style(ws.cell(row=HR, column=1 + j, value=h), f_hdr, FILL_SECT, al=AL_C)
ws.row_dimensions[HR].height = 24
ws.freeze_panes = "D7"

dvs = {}


def dv_list(formula):
    if formula not in dvs:
        d = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                           errorTitle="Invalid entry", error="Pick a value from the drop-down list.")
        ws.add_data_validation(d)
        dvs[formula] = d
    return dvs[formula]


dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="Number", error="Enter a positive number.")
ws.add_data_validation(dv_num)

UNIT_LIST = {"len": "=LenU", "vol": "=VolU", "wt": "=WtU"}
UNIT_NAME = {"len": "Len", "vol": "Vol", "wt": "Wt"}
row = HR + 1


def fac(kind, r):
    return f"INDEX({UNIT_NAME[kind]}F,MATCH($C${r},{UNIT_NAME[kind]}U,0))"


def head_vol(tcell, D, h="0"):
    return (f'IF({tcell}="",0,IF({tcell}="Conical",PI()/12*({D})^2*({h}),'
            f'INDEX(HeadK,MATCH({tcell},HeadU,0))*({D})^3))')


def section(title):
    global row
    row += 1
    ws.merge_cells(f"A{row}:J{row}")
    c = ws[f"A{row}"]; c.value = title; c.font = f_sect; c.fill = FILL_SECT; c.alignment = AL_LN
    ws.row_dimensions[row].height = 22
    row += 1


def block_header(num, title):
    global row
    for col in "ABCDEFGHIJ":
        ws[f"{col}{row}"].fill = FILL_BLK; ws[f"{col}{row}"].border = BORDER
    ws[f"A{row}"] = num; ws[f"A{row}"].font = f_blk; ws[f"A{row}"].alignment = AL_C
    ws[f"B{row}"] = title; ws[f"B{row}"].font = f_blk
    row += 1


def subhead(t):
    global row
    ws[f"B{row}"] = t; ws[f"B{row}"].font = Font(name=FN, size=9, bold=True, color="595959")
    for col in "BCDEFGHIJ":
        ws[f"{col}{row}"].fill = FILL_SUBH
    row += 1


def inp(label, kind=None, default_unit=None, lst=None, default=None, note=None, text=False, example=None):
    """Input row. kind: len/vol/wt (unit drop-down) | None. lst: drop-down list formula for train cells."""
    global row
    r = row
    style(ws[f"B{r}"], f_body, al=AL_LN); ws[f"B{r}"].value = label
    cu = ws[f"C{r}"]
    if kind:
        style(cu, f_in, FILL_IN, al=AL_C); cu.value = default_unit
        dv_list(UNIT_LIST[kind]).add(f"C{r}")
    else:
        style(cu, f_grey, al=AL_C); cu.value = "nos" if (not lst and not text) else "–"
    for i, col in enumerate(TR):
        c = ws[f"{col}{r}"]
        style(c, f_in, FILL_IN, al=AL_C, nf="General")
        if lst:
            dv_list(lst).add(f"{col}{r}")
            if default is not None:
                c.value = default
        elif not text:
            dv_num.add(f"{col}{r}")
        if example is not None and i == 0:
            c.value = example
    style(ws[f"I{r}"])
    if note:
        ws[f"J{r}"] = note
    style(ws[f"J{r}"], f_grey, al=AL_L)
    row += 1
    return r


def calc(label, fn, nf=NF_M3, key=None, total=True, strong=False, note=None, unit="m³", helper_from=None):
    """Calculated row; fn(col) returns the formula (without '=') for that train column."""
    global row
    r = row
    style(ws[f"B{r}"], f_bold if strong else f_body, al=AL_LN); ws[f"B{r}"].value = label
    style(ws[f"C{r}"], f_grey, al=AL_C); ws[f"C{r}"].value = unit
    for col in TR:
        c = ws[f"{col}{r}"]
        c.value = "=" + fn(col)
        style(c, f_res if strong else f_body, FILL_KEY if strong else FILL_RES, nf=nf, al=AL_C)
    if total:
        ws[f"I{r}"] = f'=IF(COUNT(D{r}:H{r})=0,"",SUM(D{r}:H{r}))'
    style(ws[f"I{r}"], f_res, FILL_KEY if strong else FILL_RES, nf=nf, al=AL_C)
    if note:
        ws[f"J{r}"] = note
    style(ws[f"J{r}"], f_grey, al=AL_L)
    if key:
        ws[f"K{r}"] = key
        if helper_from:
            for col, hc in zip(TR, HELP):
                ws[f"{hc}{r}"] = f"={col}{helper_from}"
    row += 1
    return r


def num(col, r):
    return f"ISNUMBER({col}{r})"


TAG_ROWS = []
# ---------------------------------------------------------------- A. Exchangers
section(f"A.  HEAT EXCHANGERS  –  tube side & shell side  ({N_HX} items)")
for k in range(1, N_HX + 1):
    ex = (k == 1)
    block_header(k, f"HEAT EXCHANGER {k}")
    R = {}
    R["tag"] = inp("Tag No.", text=True, example="EXAMPLE E-101" if ex else None,
                   note="Example data in Train 1 of item 1 – overwrite or delete" if ex else None)
    TAG_ROWS.append(R["tag"])
    R["svc"] = inp("Service / description", text=True)
    R["tema"] = inp("TEMA type (e.g. AES, BEU) – reference", text=True)
    subhead("TUBE SIDE – from data sheet")
    R["n"] = inp("Number of tubes (U-tube: no. of tube holes)", example=500 if ex else None)
    R["od"] = inp("Tube OD", "len", "mm", example=19.05 if ex else None, note="Leave blank if not given – estimated from tube pitch ÷ pitch ratio")
    R["id"] = inp("Tube ID", "len", "mm", example=15.75 if ex else None)
    R["lt"] = inp("Tube length (straight length per tube)", "len", "mm", example=6096 if ex else None)
    R["p"] = inp("Tube pitch", "len", "mm", example=23.81 if ex else None, note="Used when tube OD and/or shell ID are not given")
    R["lay"] = inp("Tube layout", lst="=LayoutU", default="Triangular (30°)")
    R["fid"] = inp("Front channel / bonnet ID", "len", "mm", example=800 if ex else None)
    R["fl"] = inp("Front channel cylindrical length", "len", "mm", example=600 if ex else None)
    R["fh"] = inp("Front head / cover type", lst="=HeadU", example="Flat" if ex else None)
    R["rid"] = inp("Rear channel / bonnet ID (blank for U-tube)", "len", "mm", example=800 if ex else None)
    R["rl"] = inp("Rear channel cylindrical length", "len", "mm", example=300 if ex else None)
    R["rh"] = inp("Rear head / cover type", lst="=HeadU", example="2:1 Ellipsoidal" if ex else None)
    R["cg"] = inp("Channel / bonnet volume – if given on data sheet", "vol", "m³", note="If filled, overrides the channel dimensions above")
    R["tw"] = inp("Tube side water weight – if given", "wt", "kg")
    subhead("SHELL SIDE – from data sheet")
    R["sid"] = inp("Shell ID", "len", "mm", example=800 if ex else None, note="If blank: area estimated as no. of tubes × pitch² × layout factor")
    R["sl"] = inp("Shell length (tubesheet to tubesheet / cover)", "len", "mm", example=6000 if ex else None, note="If blank, tube length is used")
    R["sh"] = inp("Shell cover / head type (e.g. U-tube or floating head cover)", lst="=HeadU")
    R["sw"] = inp("Shell side water weight – if given", "wt", "kg")
    subhead("VOLUME BASIS – choose per item")
    R["tb"] = inp("Tube side basis", lst="=BasisU", default="Dimensions")
    R["sb"] = inp("Shell side basis", lst="=BasisU", default="Dimensions")
    subhead("CALCULATION")
    g = lambda key, col: f"{col}{R[key]}"
    si = lambda key, kind, col: f"{col}{R[key]}*{fac(kind, R[key])}"
    R["odu"] = calc("Tube OD used", lambda c: f'IF({num(c,R["od"])},{si("od","len",c)}*1000,IF({num(c,R["p"])},{si("p","len",c)}/PitchRatio*1000,""))',
                    nf=NF_MM, total=False, unit="mm", note="Given OD, or pitch ÷ pitch ratio (Settings) when OD is blank")
    R["tubes"] = calc("Tubes internal volume  (N × π/4 × ID² × L)",
                      lambda c: f'IF(AND({num(c,R["n"])},{num(c,R["id"])},{num(c,R["lt"])}),{g("n",c)}*PI()/4*({si("id","len",c)})^2*{si("lt","len",c)},"")')

    def chan(c):
        front = (f'IF({num(c,R["fid"])},PI()/4*({si("fid","len",c)})^2*N({g("fl",c)})*{fac("len",R["fl"])}+'
                 + head_vol(g("fh", c), si("fid", "len", c)) + ",0)")
        rear = (f'IF({num(c,R["rid"])},PI()/4*({si("rid","len",c)})^2*N({g("rl",c)})*{fac("len",R["rl"])}+'
                + head_vol(g("rh", c), si("rid", "len", c)) + ",0)")
        return f'IF({num(c,R["cg"])},{si("cg","vol",c)},{front}+{rear})'
    R["chan"] = calc("Channel / bonnet volume (front + rear)", chan, note="Given volume, else cylinder + head for front and rear")
    R["tsd"] = calc("Tube side volume – from dimensions", lambda c: f'IF({g("tubes",c)}="","",{g("tubes",c)}+{g("chan",c)})')
    R["tsw"] = calc("Tube side volume – from water weight", lambda c: f'IF({num(c,R["tw"])},{si("tw","wt",c)}/WaterRho,"")')
    R["TS"] = calc("TUBE SIDE VOLUME (selected basis)", lambda c: f'IF({g("tb",c)}="Water weight",{g("tsw",c)},{g("tsd",c)})',
                   key="HX_TS", strong=True)
    R["area"] = calc("Shell cross-section area used",
                     lambda c: (f'IF({num(c,R["sid"])},PI()/4*({si("sid","len",c)})^2,IF(AND({num(c,R["n"])},{num(c,R["p"])},{g("lay",c)}<>""),'
                                f'{g("n",c)}*({si("p","len",c)})^2*INDEX(LayoutF,MATCH({g("lay",c)},LayoutU,0)),""))'),
                     nf=NF_M2, total=False, unit="m²", note="From shell ID; else no. of tubes × pitch² × layout factor (bundle area estimate)")

    def gross(c):
        L = f'IF({num(c,R["sl"])},{si("sl","len",c)},{si("lt","len",c)})'
        return (f'IF(OR({g("area",c)}="",AND(NOT({num(c,R["sl"])}),NOT({num(c,R["lt"])}))),"",'
                f'{g("area",c)}*{L}+' + head_vol(g("sh", c), f'SQRT(4*{g("area",c)}/PI())') + ")")
    R["gross"] = calc("Shell gross volume (incl. shell cover)", gross)
    R["disp"] = calc("Tube bundle displacement  (N × π/4 × OD² × L)",
                     lambda c: f'IF(AND({num(c,R["n"])},{num(c,R["odu"])},{num(c,R["lt"])}),{g("n",c)}*PI()/4*({g("odu",c)}/1000)^2*{si("lt","len",c)},"")')
    R["ssd"] = calc("Shell side volume – from dimensions", lambda c: f'IF(OR({g("gross",c)}="",{g("disp",c)}=""),"",{g("gross",c)}-{g("disp",c)})',
                    note="Negative value = check shell / tube data")
    R["ssw"] = calc("Shell side volume – from water weight", lambda c: f'IF({num(c,R["sw"])},{si("sw","wt",c)}/WaterRho,"")')
    R["SS"] = calc("SHELL SIDE VOLUME (selected basis)", lambda c: f'IF({g("sb",c)}="Water weight",{g("ssw",c)},{g("ssd",c)})',
                   key="HX_SS", strong=True)
    R["TOT"] = calc("EXCHANGER TOTAL VOLUME (tube side + shell side)",
                    lambda c: f'IF(AND({g("TS",c)}="",{g("SS",c)}=""),"",N({g("TS",c)})+N({g("SS",c)}))', key="HX_TOT", strong=True)
    R["pct"] = calc("% of volume considered",
                    lambda c: f'IF({g("TOT",c)}="","",IF(JobType="Chemical Cleaning",CC_HX,INDEX(DecFill,MATCH("Exchanger",DecTypeU,0))))',
                    nf=NF_PCT, total=False, unit="%", note="From Settings – methodology table")
    R["CONS"] = calc("CONSIDERED VOLUME for job", lambda c: f'IF(OR({g("TOT",c)}="",{g("pct",c)}=""),"",{g("TOT",c)}*{g("pct",c)})',
                     key="HX_CONS", strong=True)
    R["CHEM"] = calc("Decontamination chemical required",
                     lambda c: f'IF(OR(JobType<>"Decontamination",{g("CONS",c)}=""),"",{g("CONS",c)}*INDEX(DecChem,MATCH("Exchanger",DecTypeU,0))*1000)',
                     nf=NF_L, key="HX_CHEM", unit="L", note="Decon only. Chemical-cleaning chemicals are calculated per step on the Summary")
    ITEMS.append(("Exchanger", k, R["tag"], R["CONS"]))
    # flag negative shell side
    ws.conditional_formatting.add(f"D{R['ssd']}:H{R['ssd']}",
                                  FormulaRule(formula=[f'AND(ISNUMBER(D{R["ssd"]}),D{R["ssd"]}<0)'], font=Font(color="C00000", bold=True)))
    ws.conditional_formatting.add(f"D{R['TS']}:H{R['TS']}",
                                  FormulaRule(formula=[f'AND(D{R["tb"]}="Water weight",NOT(ISNUMBER(D{R["tw"]})))'], fill=PatternFill("solid", fgColor="F4B084")))
    ws.conditional_formatting.add(f"D{R['SS']}:H{R['SS']}",
                                  FormulaRule(formula=[f'AND(D{R["sb"]}="Water weight",NOT(ISNUMBER(D{R["sw"]})))'], fill=PatternFill("solid", fgColor="F4B084")))
    row += 1

# ---------------------------------------------------------------- B. Vessels
section(f"B.  VESSELS / COLUMNS / TANKS  ({N_VES} items)")
for k in range(1, N_VES + 1):
    ex = (k == 1)
    block_header(k, f"VESSEL / COLUMN / TANK {k}")
    R = {}
    R["tag"] = inp("Tag No.", text=True, example="EXAMPLE V-101" if ex else None,
                   note="Example data in Train 1 of item 1 – overwrite or delete" if ex else None)
    TAG_ROWS.append(R["tag"])
    R["svc"] = inp("Service / description", text=True)
    R["type"] = inp("Equipment type", lst="=EqTypeU", example="Vessel – Vertical" if ex else None,
                    note="Decides the decontamination method (Settings)")
    subhead("DIMENSIONS – from data sheet")
    R["D"] = inp("Inside diameter", "len", "mm", example=2000 if ex else None)
    R["L"] = inp("Shell length T/T (straight side)", "len", "mm", example=6000 if ex else None)
    R["h1"] = inp("Head 1 type (top / left)", lst="=HeadU", example="2:1 Ellipsoidal" if ex else None)
    R["c1"] = inp("Head 1 cone height (conical head only)", "len", "mm")
    R["h2"] = inp("Head 2 type (bottom / right)", lst="=HeadU", example="2:1 Ellipsoidal" if ex else None)
    R["c2"] = inp("Head 2 cone height (conical head only)", "len", "mm")
    R["w"] = inp("Water weight – if given", "wt", "kg")
    subhead("SELECTIONS – choose per item")
    R["b"] = inp("Volume basis", lst="=BasisU", default="Dimensions")
    R["m"] = inp("Chemical cleaning method", lst="=CCMethU", default="Full fill",
                 note="Large equipment: '20% + gamma jet circulation'")
    subhead("CALCULATION")
    g = lambda key, col: f"{col}{R[key]}"
    si = lambda key, kind, col: f"{col}{R[key]}*{fac(kind, R[key])}"
    R["sh"] = calc("Shell volume  (π/4 × D² × L)", lambda c: f'IF({num(c,R["D"])},PI()/4*({si("D","len",c)})^2*N({g("L",c)})*{fac("len",R["L"])},"")')
    R["hv1"] = calc("Head 1 volume", lambda c: f'IF({num(c,R["D"])},' + head_vol(g("h1", c), si("D", "len", c), f'N({g("c1",c)})*{fac("len",R["c1"])}') + ',"")')
    R["hv2"] = calc("Head 2 volume", lambda c: f'IF({num(c,R["D"])},' + head_vol(g("h2", c), si("D", "len", c), f'N({g("c2",c)})*{fac("len",R["c2"])}') + ',"")')
    R["vd"] = calc("Volume – from dimensions", lambda c: f'IF({num(c,R["D"])},{g("sh",c)}+{g("hv1",c)}+{g("hv2",c)},"")')
    R["vw"] = calc("Volume – from water weight", lambda c: f'IF({num(c,R["w"])},{si("w","wt",c)}/WaterRho,"")')
    R["SEL"] = calc("EQUIPMENT VOLUME (selected basis)", lambda c: f'IF({g("b",c)}="Water weight",{g("vw",c)},{g("vd",c)})',
                    key="V_SEL", strong=True, helper_from=R["type"])
    R["meth"] = calc("Method applied",
                     lambda c: (f'IF({g("SEL",c)}="","",IF(JobType="Chemical Cleaning",{g("m",c)},'
                                f'IFERROR(INDEX(DecMeth,MATCH({g("type",c)},DecTypeU,0)),"⚠ select equipment type")))'),
                     nf="@", total=False, unit="–")
    for col in TR:
        ws[f"{col}{R['meth']}"].alignment = AL_C
        ws[f"{col}{R['meth']}"].font = Font(name=FN, size=8)
    ws.row_dimensions[R["meth"]].height = 36
    R["pct"] = calc("% of volume considered",
                    lambda c: (f'IF({g("SEL",c)}="","",IF(JobType="Chemical Cleaning",IFERROR(INDEX(CCMethPct,MATCH({g("m",c)},CCMethU,0)),""),'
                               f'IFERROR(INDEX(DecFill,MATCH({g("type",c)},DecTypeU,0)),"")))'),
                    nf=NF_PCT, total=False, unit="%", note="From Settings – methodology / CC method tables")
    R["CONS"] = calc("CONSIDERED VOLUME for job", lambda c: f'IF(OR({g("SEL",c)}="",{g("pct",c)}=""),"",{g("SEL",c)}*{g("pct",c)})',
                     key="V_CONS", strong=True, helper_from=R["type"])
    R["CHEM"] = calc("Decontamination chemical required",
                     lambda c: (f'IF(OR(JobType<>"Decontamination",{g("CONS",c)}=""),"",'
                                f'IFERROR({g("CONS",c)}*INDEX(DecChem,MATCH({g("type",c)},DecTypeU,0))*1000,""))'),
                     nf=NF_L, key="V_CHEM", unit="L", helper_from=R["type"],
                     note="Decon only: column 1% (vapour), vessels 2% (boil-out), tanks per Settings")
    ITEMS.append(("Vessel", k, R["tag"], R["CONS"]))
    ws.conditional_formatting.add(f"D{R['SEL']}:H{R['SEL']}",
                                  FormulaRule(formula=[f'AND(D{R["b"]}="Water weight",NOT(ISNUMBER(D{R["w"]})))'], fill=PatternFill("solid", fgColor="F4B084")))
    row += 1

# ---------------------------------------------------------------- C. Piping
section(f"C.  PIPING FROM ISOMETRICS  ({N_PIPE} lines)")
for k in range(1, N_PIPE + 1):
    ex = (k == 1)
    block_header(k, f"PIPING LINE {k}")
    R = {}
    R["tag"] = inp("Line No. (from isometric)", text=True, example='EXAMPLE 6"-P-1001' if ex else None,
                   note="Example data in Train 1 of line 1 – overwrite or delete" if ex else None)
    TAG_ROWS.append(R["tag"])
    R["nps"] = inp("Nominal pipe size (NPS)", lst="=NPSU", example='6" (DN 150)' if ex else None)
    R["sch"] = inp("Schedule / wall", lst="=SchU", example="Sch 40" if ex else None)
    R["idg"] = inp("Inside diameter – if given (overrides NPS / schedule)", "len", "mm")
    R["len"] = inp("Total line length (from isometric)", "len", "m", example=120 if ex else None)
    g = lambda key, col: f"{col}{R[key]}"
    si = lambda key, kind, col: f"{col}{R[key]}*{fac(kind, R[key])}"

    def idu(c):
        wt = f"INDEX(PipeWT,MATCH({g('nps',c)},NPSU,0),MATCH({g('sch',c)},SchU,0))"
        return (f'IF({num(c,R["idg"])},{si("idg","len",c)}*1000,IF(OR({g("nps",c)}="",{g("sch",c)}=""),"",'
                f'IFERROR(IF({wt}="","Sch n/a",INDEX(PipeOD,MATCH({g("nps",c)},NPSU,0))-2*{wt}),"")))')
    R["idu"] = calc("Inside diameter used", idu, nf=NF_MM, total=False, unit="mm", note="Given ID, else OD − 2 × wall (Lookup sheet)")
    R["VOL"] = calc("LINE VOLUME  (π/4 × ID² × length)",
                    lambda c: f'IF(AND({num(c,R["idu"])},{num(c,R["len"])}),PI()/4*({g("idu",c)}/1000)^2*{si("len","len",c)},"")',
                    key="P_VOL", strong=True)
    R["pct"] = calc("% of volume considered",
                    lambda c: f'IF({g("VOL",c)}="","",IF(JobType="Chemical Cleaning",CC_PIPE,INDEX(DecFill,MATCH("Piping",DecTypeU,0))))',
                    nf=NF_PCT, total=False, unit="%")
    R["CONS"] = calc("CONSIDERED VOLUME for job", lambda c: f'IF(OR({g("VOL",c)}="",{g("pct",c)}=""),"",{g("VOL",c)}*{g("pct",c)})',
                     key="P_CONS", strong=True)
    R["CHEM"] = calc("Decontamination chemical required",
                     lambda c: f'IF(OR(JobType<>"Decontamination",{g("CONS",c)}=""),"",{g("CONS",c)}*INDEX(DecChem,MATCH("Piping",DecTypeU,0))*1000)',
                     nf=NF_L, key="P_CHEM", unit="L")
    ITEMS.append(("Piping", k, R["tag"], R["CONS"]))
    row += 1
LAST = row
ws.print_title_rows = f"{HR}:{HR}"
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = f"A1:J{LAST}"

# ================================================================ SUMMARY sheet
S = ws_sum
S.sheet_view.showGridLines = False
for c, w in {"A": 2, "B": 52, "C": 9, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16, "I": 18}.items():
    S.column_dimensions[c].width = w
S["B1"] = '="VOLUME & CHEMICAL SUMMARY – "&UPPER(JobType)'; S["B1"].font = f_title
S["B2"] = '="Client: "&ProjClient&"   ·   Plant: "&ProjPlant&"   ·   Proposal: "&ProjNo'; S["B2"].font = f_sub
ex_rows = [TAG_ROWS[0], TAG_ROWS[N_HX], TAG_ROWS[N_HX + N_VES]]
S["B3"] = ("=IF(" + "+".join(f'COUNTIF({SVC}!D{r}:H{r},"EXAMPLE*")' for r in ex_rows) +
           '>0,"⚠ Example data is still present (Tag starting with EXAMPLE) – delete it before issuing the proposal.","")')
S["B3"].font = f_red
VC = SVC
srow = 5


def s_head(title):
    global srow
    S.merge_cells(f"B{srow}:I{srow}")
    c = S[f"B{srow}"]; c.value = title; c.font = f_sect; c.fill = FILL_SECT
    srow += 1
    for j, h in enumerate(["Item", "Unit"] + [f"=TrainName{i+1}" for i in range(N_TRAINS)] + ["Total"]):
        style(S.cell(row=srow, column=2 + j, value=h), f_hdr, PatternFill("solid", fgColor="44546A"), al=AL_C)
    srow += 1


def s_row(label, fn, unit="m³", nf=NF_M3, strong=False, total=True):
    global srow
    r = srow
    style(S[f"B{r}"], f_bold if strong else f_body, al=AL_LN); S[f"B{r}"].value = label
    style(S[f"C{r}"], f_grey, al=AL_C); S[f"C{r}"].value = unit
    for i, col in enumerate(TR):
        c = S[f"{col}{r}"]; c.value = "=" + fn(i, col)
        style(c, f_res if strong else f_body, FILL_KEY if strong else FILL_RES, nf=nf, al=AL_C)
    S[f"I{r}"] = f"=SUM(D{r}:H{r})" if total else None
    style(S[f"I{r}"], f_res, FILL_KEY if strong else FILL_RES, nf=nf, al=AL_C)
    srow += 1
    return r


def sumkey(key, col):
    return f'SUMIFS({VC}!{col}:{col},{VC}!$K:$K,"{key}")'


def sumtype(key, typ, i, col):
    return f'SUMIFS({VC}!{col}:{col},{VC}!$K:$K,"{key}",{VC}!${HELP[i]}:${HELP[i]},"{typ}")'


TYPES = ["Vessel – Vertical", "Vessel – Horizontal", "Column", "Tank"]
TLAB = {"Vessel – Vertical": "Vessels – vertical", "Vessel – Horizontal": "Vessels – horizontal", "Column": "Columns", "Tank": "Tanks"}
s_head("1.  EQUIPMENT VOLUMES (full volume, selected basis)")
a = s_row("Heat exchangers – tube side", lambda i, c: sumkey("HX_TS", c))
b = s_row("Heat exchangers – shell side", lambda i, c: sumkey("HX_SS", c))
ty = [s_row(TLAB[t], lambda i, c, t=t: sumtype("V_SEL", t, i, c)) for t in TYPES]
ty_none = s_row("Vessels with no equipment type selected",
                lambda i, c: f"{sumkey('V_SEL', c)}-SUM({','.join(f'{c}{x}' for x in ty)})")
p = s_row("Piping", lambda i, c: sumkey("P_VOL", c))
tot = s_row("TOTAL SYSTEM VOLUME", lambda i, c: f"SUM({c}{a}:{c}{p})", strong=True)
s_row("TOTAL SYSTEM VOLUME", lambda i, c: f"{c}{tot}*1000", unit="L", nf=NF_L)
s_row("TOTAL SYSTEM VOLUME", lambda i, c: f"{c}{tot}/0.003785411784", unit="US gal", nf=NF_L)
srow += 1
s_head("2.  CONSIDERED VOLUME FOR THE SELECTED JOB TYPE  (% per Settings / per-item method)")
c1 = s_row("Heat exchangers (tube + shell side)", lambda i, c: sumkey("HX_CONS", c))
cty = [s_row(TLAB[t], lambda i, c, t=t: sumtype("V_CONS", t, i, c)) for t in TYPES]
cnone = s_row("Vessels with no equipment type selected",
              lambda i, c: f"{sumkey('V_CONS', c)}-SUM({','.join(f'{c}{x}' for x in cty)})")
cp = s_row("Piping", lambda i, c: sumkey("P_CONS", c))
ctot = s_row("TOTAL CONSIDERED VOLUME (water / solution)", lambda i, c: f"SUM({c}{c1}:{c}{cp})", strong=True)
s_row("TOTAL CONSIDERED VOLUME", lambda i, c: f"{c}{ctot}*1000", unit="L", nf=NF_L)
srow += 1

s_head("3A.  CHEMICAL CLEANING – chemical per step  (= total considered volume × concentration × fills)")
S[f"B{srow}"] = '=IF(JobType<>"Chemical Cleaning","(Not applicable – job type is "&JobType&")","")'
S[f"B{srow}"].font = f_grey; srow += 1
step_rows = []
for k, sr in enumerate(STEP_ROWS):
    lab = f'=IF({SS}!$B${sr}="","Step {k+1} (not used)",{SS}!$B${sr}&"  @ "&TEXT({SS}!$C${sr},"0.0%")&" × "&N({SS}!$D${sr})&" fill(s)")'
    r = s_row("", lambda i, c, sr=sr: f'IF(OR(JobType<>"Chemical Cleaning",{SS}!$C${sr}=""),0,{c}{ctot}*1000*{SS}!$C${sr}*IF(N({SS}!$D${sr})=0,1,{SS}!$D${sr}))',
              unit="L", nf=NF_L)
    S[f"B{r}"] = lab
    step_rows.append(r)
cc_tot = s_row("TOTAL CHEMICAL – chemical cleaning", lambda i, c: f"SUM({c}{step_rows[0]}:{c}{step_rows[-1]})", unit="L", nf=NF_L, strong=True)
srow += 1
s_head("3B.  DECONTAMINATION – chemical  (= considered volume × chemical % per equipment type)")
S[f"B{srow}"] = '=IF(JobType<>"Decontamination","(Not applicable – job type is "&JobType&")","")'
S[f"B{srow}"].font = f_grey; srow += 1
d1 = s_row("Heat exchangers", lambda i, c: sumkey("HX_CHEM", c), unit="L", nf=NF_L)
dty = [s_row(TLAB[t], lambda i, c, t=t: sumtype("V_CHEM", t, i, c), unit="L", nf=NF_L) for t in TYPES]
dp = s_row("Piping", lambda i, c: sumkey("P_CHEM", c), unit="L", nf=NF_L)
dc_tot = s_row("TOTAL CHEMICAL – decontamination", lambda i, c: f"SUM({c}{d1}:{c}{dp})", unit="L", nf=NF_L, strong=True)
srow += 1
s_head("4.  CHEMICAL PACKAGING  (selected job type)")
ch = s_row("Total chemical required", lambda i, c: f'IF(JobType="Chemical Cleaning",{c}{cc_tot},{c}{dc_tot})', unit="L", nf=NF_L, strong=True)
s_row('="Drums ("&DrumL&" L each)"', lambda i, c: f"IF({c}{ch}=0,0,ROUNDUP({c}{ch}/DrumL,0))", unit="nos", nf=NF_L)
S[f"B{srow-1}"].value = '="Drums ("&DrumL&" L each) – rounded up per train"'
s_row("", lambda i, c: f"IF({c}{ch}=0,0,ROUNDUP({c}{ch}/IBCL,0))", unit="nos", nf=NF_L)
S[f"B{srow-1}"].value = '="IBCs ("&IBCL&" L each) – rounded up per train"'
srow += 1
for t in ["Notes:",
          f"• Volumes per item come from the basis chosen on '{SHEET_NAMES['vc']}' (dimensions or water weight).",
          "• Chemical cleaning: exchangers and piping full volume; vessels / columns / tanks per item (Full fill or 20% + gamma jet circulation).",
          "• Decontamination: columns, exchangers, piping – vapour phase, complete volume, 1% chemical; vessels – boil-out 30% water + 2% chemical with steam injection; "
          "tanks – gamma-jet circulation (heating up to 80 °C if required) – % per Settings (to confirm).",
          "• Internals, nozzles, hoses and temporary circuit volumes are not included – add allowances separately if required."]:
    S[f"B{srow}"] = t; S[f"B{srow}"].font = f_grey if t != "Notes:" else f_bold
    srow += 1
S.page_setup.orientation = "landscape"; S.sheet_properties.pageSetUpPr.fitToPage = True
S.page_setup.fitToWidth = 1; S.page_setup.fitToHeight = 0
S.freeze_panes = "D5"

if STANDALONE and __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Cleaning_Decon_Volume_Template.xlsx"
    wb.save(out)
    print("saved", out, "calc rows:", LAST)

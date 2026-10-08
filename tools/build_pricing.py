"""Delight International – service pricing workbooks (master + single-service files).

Pilot: framework (Quote Info, Rates, markups, client quote) + Chemical Cleaning, with the row-based
volume sheets (tools/volume_rows.py). Rates are the values from ASAB CC Pricing R2 and DUQM Pricing
Rev 1.0 (highest where both list the same item).

Usage:  python3 tools/build_pricing.py pricing/
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import volume_rows as vr  # noqa: E402

SERVICES = [("CC", "Chemical Cleaning"), ("DC", "Decontamination"), ("LOF", "Lube Oil Flushing"), ("AB", "Air Blowing"),
            ("PIG", "Pigging"), ("TC", "Tank Cleaning"), ("N2L", "Nitrogen / Helium Leak Test"), ("N2F", "Nitrogen Flushing / Purging")]

FN = "Arial"
BLUE, GREEN, NAVY = "1E73BE", "3AAA35", "1F3864"
f_title = Font(name=FN, size=16, bold=True, color=NAVY)
f_sub = Font(name=FN, size=10, italic=True, color="595959")
f_sect = Font(name=FN, size=11, bold=True, color="FFFFFF")
f_hdr = Font(name=FN, size=9, bold=True, color="FFFFFF")
f_body = Font(name=FN, size=10)
f_bold = Font(name=FN, size=10, bold=True)
f_in = Font(name=FN, size=10, color="0000FF")
f_link = Font(name=FN, size=10, color="008000")
f_grey = Font(name=FN, size=9, color="808080")
f_red = Font(name=FN, size=10, bold=True, color="C00000")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_RES = PatternFill("solid", fgColor="E2EFDA")
FILL_KEY = PatternFill("solid", fgColor="C6E0B4")
FILL_SECT = PatternFill("solid", fgColor=BLUE)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_GREEN = PatternFill("solid", fgColor=GREEN)
FILL_MISS = PatternFill("solid", fgColor="F4B084")
FILL_UNIT = PatternFill("solid", fgColor="D9E1F2")
FILL_TOT = PatternFill("solid", fgColor="BDD7EE")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
AL_LN = Alignment(horizontal="left", vertical="center")
AL_R = Alignment(horizontal="right", vertical="center")
NF_AED = '#,##0.00;-#,##0.00;"-"'
NF_NUM = '#,##0.##;-#,##0.##;"-"'
NF_PCT = '0.0%;-0.0%;"-"'
NF_M3 = '#,##0.000;-#,##0.000;"-"'
q = vr.q


def style(c, font=f_body, fill=None, nf=None, al=None, border=True):
    c.font = font
    if fill: c.fill = fill
    if nf: c.number_format = nf
    if al: c.alignment = al
    if border: c.border = BORDER


class Book:
    def __init__(self):
        self.wb = Workbook()
        self.dvs = {}

    def name(self, n, sheet, ref):
        self.wb.defined_names[n] = DefinedName(n, attr_text=f"{q(sheet)}!{ref}")

    def dv(self, ws, formula, cells, num=False):
        key = (ws.title, formula)
        if key not in self.dvs:
            if num:
                d = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                   showErrorMessage=True, errorTitle="Number", error="Enter a positive number.")
            else:
                d = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                                   errorTitle="Invalid entry", error="Pick a value from the drop-down list.")
            ws.add_data_validation(d)
            self.dvs[key] = d
        self.dvs[key].add(cells)


def label(ws, cell, text, font=f_bold):
    style(ws[cell], font, al=AL_LN); ws[cell].value = text


def inp(ws, cell, value=None, nf=None, al=AL_C):
    style(ws[cell], f_in, FILL_IN, nf=nf, al=al)
    if value is not None:
        ws[cell].value = value


def bar(ws, row, text, c1="A", c2="L", fill=FILL_SECT):
    ws.merge_cells(f"{c1}{row}:{c2}{row}")
    c = ws[f"{c1}{row}"]; c.value = text; c.font = f_sect; c.fill = fill; c.alignment = AL_LN
    ws.row_dimensions[row].height = 20


# ================================================================ QUOTE INFO
def build_quote_info(B):
    ws = B.wb.active
    ws.title = "Quote Info"
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 44, "C": 22, "D": 16, "E": 16, "F": 16, "G": 50}.items():
        ws.column_dimensions[c].width = w
    ws["B1"] = "DELIGHT INTERNATIONAL – PRICING: QUOTE INFORMATION"; ws["B1"].font = f_title
    ws["B2"] = ("Fill in this sheet first. Yellow = input (blue text) · green = calculated. Markups, contingency, discount and VAT "
                "apply to every service in this workbook."); ws["B2"].font = f_sub
    r = 4
    bar(ws, r, "QUOTATION DETAILS", "B", "G")
    fields = [("Quotation No.", "QNo"), ("Revision", "QRev"), ("Date", "QDate"), ("Client", "QClient"),
              ("Attention (client contact)", "QAttn"), ("Project / Plant", "QProject"), ("Location / site", "QLoc"),
              ("Client enquiry / RFQ reference", "QRef"), ("Prepared by", "QPrep")]
    for i, (t, n) in enumerate(fields):
        rr = r + 1 + i
        label(ws, f"B{rr}", t)
        inp(ws, f"C{rr}", al=AL_LN, nf="dd-mmm-yyyy" if n == "QDate" else None)
        ws.merge_cells(f"C{rr}:E{rr}")
        B.name(n, ws.title, f"$C${rr}")
    r = 15
    bar(ws, r, "CURRENCY  (rates are entered in AED; the client quote is converted)", "B", "G")
    label(ws, f"B{r+1}", "Quote currency"); inp(ws, f"C{r+1}", "AED"); B.name("QCurr", ws.title, f"$C${r+1}")
    for j, h in enumerate(["Currency", "AED per 1 unit", "", "", "Source"]):
        if h:
            style(ws.cell(row=r + 2, column=2 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
    cur = [("AED", 1, "Base currency"), ("USD", 3.6725, "UAE central bank peg (AED 3.6725 = USD 1)"),
           ("SAR", "=3.6725/3.75", "Via USD pegs (SAR 3.75 = USD 1)"), ("OMR", "=3.6725/0.3845", "Via USD pegs (OMR 0.3845 = USD 1)"),
           ("GBP", 4.92, "DUQM Rev 1.0 market rate 17-Jun-2026 – update")]
    for i, (c_, v, s_) in enumerate(cur):
        rr = r + 3 + i
        label(ws, f"B{rr}", c_); inp(ws, f"C{rr}", v, nf="0.0000")
        ws[f"F{rr}"] = s_; ws[f"F{rr}"].font = f_grey
    B.name("CurrU", ws.title, f"$B${r+3}:$B${r+2+len(cur)}")
    B.name("CurrRate", ws.title, f"$C${r+3}:$C${r+2+len(cur)}")
    B.dv(ws, "=CurrU", f"C{r+1}")
    rr = r + 3 + len(cur)
    label(ws, f"B{rr}", "Exchange rate used (AED per 1 quote-currency unit)")
    ws[f"C{rr}"] = "=INDEX(CurrRate,MATCH(QCurr,CurrU,0))"; style(ws[f"C{rr}"], f_bold, FILL_RES, nf="0.0000", al=AL_C)
    B.name("QRate", ws.title, f"$C${rr}")
    r = rr + 2
    bar(ws, r, "WORKING PATTERN", "B", "G")
    label(ws, f"B{r+1}", "Shift length (hours)"); inp(ws, f"C{r+1}", 12); B.name("QShiftH", ws.title, f"$C${r+1}")
    label(ws, f"B{r+2}", "Shifts per day (1 = day only, 2 = 24-hour)"); inp(ws, f"C{r+2}", 2); B.name("QShiftsDay", ws.title, f"$C${r+2}")
    ws[f"G{r+1}"] = "Manpower rates are per person per shift (12-hour day)"; ws[f"G{r+1}"].font = f_grey
    r = r + 4
    bar(ws, r, "MARKUP CONTROLS  (selling = cost × (1 + markup))", "B", "G")
    label(ws, f"B{r+1}", "Markup mode"); inp(ws, f"C{r+1}", "Simple"); B.name("QMkMode", ws.title, f"$C${r+1}")
    ws[f"G{r+1}"] = "Simple = 2 markups (in-house / rented).  Detailed = 5 markups."; ws[f"G{r+1}"].font = f_grey
    ws[f"E{r+1}"] = "Simple"; ws[f"F{r+1}"] = "Detailed"
    for c in "EF":
        ws[f"{c}{r+1}"].font = Font(name=FN, size=8, color="FFFFFF")
    B.name("MkModeU", ws.title, f"$E${r+1}:$F${r+1}"); B.dv(ws, "=MkModeU", f"C{r+1}")
    mk = [("SIMPLE – In-house crew & equipment", "MkIn"), ("SIMPLE – Rented crew & equipment, bought-in items", "MkRent"),
          ("DETAILED – In-house manpower", "MkInMP"), ("DETAILED – Rented manpower", "MkRentMP"),
          ("DETAILED – In-house equipment", "MkInEq"), ("DETAILED – Rented equipment", "MkRentEq"),
          ("DETAILED – Chemicals, consumables & third-party services", "MkCons")]
    for i, (t, n) in enumerate(mk):
        rr = r + 2 + i
        label(ws, f"B{rr}", t, f_body); inp(ws, f"C{rr}", 0.3, nf=NF_PCT); B.name(n, ws.title, f"$C${rr}")
    ws[f"G{r+2}"] = "30% = highest of ASAB R2 (30%) and DUQM Rev 1.0 (10–30%) – edit per job"; ws[f"G{r+2}"].font = f_grey
    m0 = r + 10
    label(ws, f"B{m0}", "Effective markup applied (calculated)")
    for j, h in enumerate(["In-house", "Rented", "Third-party"]):
        style(ws.cell(row=m0, column=3 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
    B.name("SrcU", ws.title, f"$C${m0}:$E${m0}")
    simple = [["MkIn", "MkRent", "MkRent"]] * 3
    detail = [["MkInMP", "MkRentMP", "MkCons"], ["MkInEq", "MkRentEq", "MkCons"], ["MkCons", "MkCons", "MkCons"]]
    for i, cname in enumerate(["Manpower", "Equipment", "Chemicals, consumables & services"]):
        rr = m0 + 1 + i
        label(ws, f"B{rr}", cname, f_body)
        for j in range(3):
            c = ws.cell(row=rr, column=3 + j, value=f'=IF(QMkMode="Detailed",N({detail[i][j]}),N({simple[i][j]}))')
            style(c, f_body, FILL_RES, nf=NF_PCT, al=AL_C)
    B.name("MkMatrix", ws.title, f"$C${m0+1}:$E${m0+3}")
    r = m0 + 5
    bar(ws, r, "CONTINGENCY, DISCOUNT & VAT", "B", "G")
    for i, (t, n, v, note) in enumerate([
            ("Contingency (% of total cost, internal risk allowance)", "QCont", None, "Added to the price, not marked up"),
            ("Commercial discount (% of price)", "QDisc", None, "Applied to the price before VAT"),
            ("VAT (%)", "QVAT", 0.05, "UAE 5% · KSA 15% · Oman 5% · Qatar 0% – set per quotation")]):
        rr = r + 1 + i
        label(ws, f"B{rr}", t, f_body); inp(ws, f"C{rr}", v, nf=NF_PCT); B.name(n, ws.title, f"$C${rr}")
        ws[f"G{rr}"] = note; ws[f"G{rr}"].font = f_grey
    r = r + 5
    bar(ws, r, "COMMERCIAL TERMS  (printed on the client quotation)", "B", "G")
    terms = [("Validity", "QTValidity"), ("Payment terms", "QTPayment"), ("Delivery / mobilisation time", "QTMob"),
             ("Exclusions", "QTExcl"), ("Client scope / supplied by client", "QTClient"), ("Other notes", "QTNotes")]
    for i, (t, n) in enumerate(terms):
        rr = r + 1 + i
        label(ws, f"B{rr}", t); inp(ws, f"C{rr}", al=AL_L); ws.merge_cells(f"C{rr}:G{rr}"); ws.row_dimensions[rr].height = 30
        B.name(n, ws.title, f"$C${rr}")
    r = r + len(terms) + 2
    bar(ws, r, "COMPANY DETAILS  (letterhead on the client quotation)", "B", "G")
    co = [("Company name", "CoName", "Delight Equipment International L.L.C."),
          ("Address", "CoAddr", "P.O. Box 2932, Abu Dhabi, United Arab Emirates"),
          ("Telephone / fax", "CoTel", "Tel: +971 2 5515641  |  Fax: +971 2 5515643"),
          ("E-mail / website", "CoWeb", "info@delightintl.ae  |  www.delightintl.ae"),
          ("Tagline", "CoTag", "MAINTENANCE SIMPLIFIED"), ("Certifications", "CoCert", "ISO 9001 · ISO 14001 · ISO 45001")]
    for i, (t, n, v) in enumerate(co):
        rr = r + 1 + i
        label(ws, f"B{rr}", t); inp(ws, f"C{rr}", v, al=AL_LN); ws.merge_cells(f"C{rr}:G{rr}")
        B.name(n, ws.title, f"$C${rr}")


# ================================================================ RATES (values from ASAB R2 + DUQM Rev 1.0)
MANPOWER = [("Project Manager", 1000, "ASAB 1,000 · DUQM 1,000"), ("Project Engineer", 700, "ASAB 500 · DUQM 700 → highest"),
            ("Supervisor", 500, "ASAB 400 · DUQM 500 → highest"), ("Job Performer / Permit Receiver", None, "not in either sheet"),
            ("Safety Officer", 300, "DUQM"), ("Technician / Operator", 400, "ASAB 300 · DUQM 400 → highest"),
            ("Pipe Fitter", None, "not in either sheet"), ("Helper", 300, "ASAB 200 · DUQM 300 → highest"),
            ("Others", 400, "DUQM"), ("Decontamination Specialist", 5400, "DUQM (per day)"),
            ("Specialised Services", 2400, "DUQM (per day)"), ("Remote Technical Support", 2250, "DUQM (per day)")]
EQUIPMENT = [("Chemical circulation pump skid (SS316, VFD)", "day", 800, "ASAB"),
             ("Vessel circulation pump (gamma jet feed)", "day", 800, "ASAB"),
             ("Centrifugal pump", "day", 1000, "DUQM"), ("Drum pump", "day", 500, "DUQM"),
             ("Diaphragm pump & manifolds (transfer set)", "day", 300, "ASAB"), ("Diaphragm pump", "day", 100, "DUQM"),
             ("Manifold", "day", 250, "DUQM"), ("Chemical mixing / dosing tank 10 m³", "day", 150, "ASAB"),
             ("DM water buffer tank", "day", 150, "ASAB"), ("Waste holding tank 500 bbl (~80 m³)", "day", 400, "ASAB rental"),
             ("Effluent tank 60 m³", "day", 500, "DUQM"), ("Diesel boiler 2 TPH", "day", 3000, "ASAB rental"),
             ("Heat exchanger 1000 kW", "day", 1000, "ASAB rental"), ("Gamma jetting nozzle", "day", 500, "ASAB 500 · DUQM 300 → highest"),
             ("Hot-air dryer unit c/w dew point meter", "day", 500, "ASAB"),
             ("Hoses & fittings package", "LS", 10000, "ASAB lump sum (DUQM 6,000 per day – see next line)"),
             ("Hoses & fittings package (per day)", "day", 6000, "DUQM"), ("Multi gas meter", "day", 300, "DUQM"),
             ("Benzene meter", "day", 300, "DUQM"), ("Benzene SEP tubes (set)", "LS", 600, "DUQM"), ("Pickup", "day", 500, "DUQM"),
             ("Material container", "day", 250, "DUQM"), ("Lab container", "day", 250, "DUQM"),
             ("20' Circulation unit (CORE)", "day", 820, "DUQM"), ("Mobile heat exchanger (CORE)", "day", 700, "DUQM"),
             ("Support container (CORE)", "day", 320, "DUQM"), ("10-point VP injection module (CORE)", "day", 890, "DUQM")]
SERVICES_RATES = [("Waste disposal", "m³", 290, "ASAB GMET/SQ/03002 290 · DUQM 50 (TBC) → highest"),
                  ("Disposables / PPE", "person-day", 150, "ASAB 100 · DUQM 150 → highest"),
                  ("Pre-engineering works", "person-day", 1500, "ASAB 500–1,500 → highest"),
                  ("Kurita litmus paper", "roll", 150, "DUQM"), ("Chemical shipping", "trip", 18350, "ASAB (USD 5,000 × 3.67)"),
                  ("Sea freight 40' HC (CFR Sohar, incl. VGM)", "container", 40453.37, "DUQM (£8,166.67 + £55.56) × 4.92"),
                  ("Local co-operative fee (Al Dhafra)", "month", 3000, "ASAB (GMET quote, billed as actual)")]
VEHICLES = [("40 ft trailer", 15000, "ASAB 2,500 · DUQM 15,000 → highest"), ("Low-bed trailer", None, ""),
            ("Truck with crane (HIAB)", None, ""), ("Crew bus", None, "")]
EXTRAS = [("FAT – food, accommodation & transport (per person per day)", "XFAT", 250, "ASAB 150 · DUQM 250 → highest"),
          ("Accommodation – specialists (per person per day)", "XAcc", 700, "DUQM"),
          ("Visa (per person)", "XVisa", 2300, "DUQM"), ("Flight – return (per person)", "XFlight", 3300, "DUQM"),
          ("Crew visas & tickets – India (per person)", "XCrewVT", 4500, "DUQM"),
          ("Crew transport – pickup (per vehicle per day)", "XTrans", 500, "DUQM pickup")]
CHEM_IBC = [("CORE CLEAN VP1", 17463.24, "DUQM – vapour phase"), ("CORE CLEAN 2600", 35625, "DUQM – steam-out / boil-out"),
            ("CORE CLEAN 3010", 29411.76, "DUQM (ref)"), ("CORE CLEAN 5421", 33823.53, "DUQM (ref)")]


def build_rates(B):
    ws = B.wb.create_sheet("Rates")
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 50, "C": 14, "D": 16, "E": 16, "F": 16, "G": 16, "H": 50}.items():
        ws.column_dimensions[c].width = w
    ws["B1"] = "RATES (COST, AED, excl. markup)"; ws["B1"].font = f_title
    ws["B2"] = ("Values taken from ASAB CC Pricing R2 and DUQM Pricing Rev 1.0 – the higher one where both list the same item. "
                "Rented column blank = same as cost. Add items in the blank yellow rows."); ws["B2"].font = f_sub
    r = 4

    def header(title, hdrs):
        nonlocal r
        bar(ws, r, title, "B", "H")
        for j, h in enumerate(hdrs):
            if h:
                style(ws.cell(row=r + 1, column=2 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
        ws.row_dimensions[r + 1].height = 28
        r += 2

    # manpower
    header("MANPOWER – cost per person per shift / day", ["Grade", "", "Cost", "Rented cost (if different)", "", "", "Source / remarks"])
    first = r
    for nm, cost, rem in MANPOWER + [(None, None, None)] * 6:
        inp(ws, f"B{r}", nm, al=AL_LN); style(ws[f"C{r}"], f_grey)
        inp(ws, f"D{r}", cost, nf=NF_AED); inp(ws, f"E{r}", nf=NF_AED)
        ws[f"H{r}"] = rem; style(ws[f"H{r}"], f_grey, al=AL_L)
        r += 1
    B.name("MPU", "Rates", f"$B${first}:$B${r-1}"); B.name("MPCost", "Rates", f"$D${first}:$D${r-1}"); B.name("MPRent", "Rates", f"$E${first}:$E${r-1}")
    r += 1
    header("EQUIPMENT – cost per unit per day (LS = lump sum, counted once)", ["Equipment", "Unit", "Cost", "Rented cost (if different)", "", "", "Source / remarks"])
    first = r
    for nm, unit, cost, rem in EQUIPMENT + [(None, None, None, None)] * 10:
        inp(ws, f"B{r}", nm, al=AL_LN); inp(ws, f"C{r}", unit); B.dv(ws, '"day,LS"', f"C{r}")
        inp(ws, f"D{r}", cost, nf=NF_AED); inp(ws, f"E{r}", nf=NF_AED)
        ws[f"H{r}"] = rem; style(ws[f"H{r}"], f_grey, al=AL_L)
        r += 1
    B.name("EQU", "Rates", f"$B${first}:$B${r-1}"); B.name("EQUnit", "Rates", f"$C${first}:$C${r-1}")
    B.name("EQCost", "Rates", f"$D${first}:$D${r-1}"); B.name("EQRent", "Rates", f"$E${first}:$E${r-1}")
    r += 1
    header("CHEMICALS – price per kg (or per IBC)", ["Chemical", "Purity", "Price per kg", "Price per IBC", "IBC net content (kg)", "Price per kg used", "Source / remarks"])
    first = r
    rows = [(n, p, pk, None, None, "ASAB R2") for n, p, pk, _a, _b in vr.CHEMS] + \
           [(n, 1, None, pi, None, rem + " – enter IBC net kg") for n, pi, rem in CHEM_IBC] + [(None,) * 6] * 6
    for nm, pur, pk, pi, kg, rem in rows:
        inp(ws, f"B{r}", nm, al=AL_LN); inp(ws, f"C{r}", pur, nf="0%"); inp(ws, f"D{r}", pk, nf=NF_AED)
        inp(ws, f"E{r}", pi, nf=NF_AED); inp(ws, f"F{r}", kg, nf=NF_NUM)
        ws[f"G{r}"] = f'=IF(ISNUMBER(D{r}),D{r},IF(AND(ISNUMBER(E{r}),N(F{r})>0),E{r}/F{r},""))'
        style(ws[f"G{r}"], f_bold, FILL_RES, nf=NF_AED, al=AL_C)
        ws[f"H{r}"] = rem; style(ws[f"H{r}"], f_grey, al=AL_L)
        r += 1
    B.name("CHU", "Rates", f"$B${first}:$B${r-1}"); B.name("CHPurity", "Rates", f"$C${first}:$C${r-1}")
    B.name("CHPriceKg", "Rates", f"$G${first}:$G${r-1}")
    r += 1
    header("CONSUMABLES & THIRD-PARTY SERVICES – cost per unit", ["Item", "Unit", "Cost", "", "", "", "Source / remarks"])
    first = r
    for nm, unit, cost, rem in SERVICES_RATES + [(None, None, None, None)] * 8:
        inp(ws, f"B{r}", nm, al=AL_LN); inp(ws, f"C{r}", unit); inp(ws, f"D{r}", cost, nf=NF_AED)
        ws[f"H{r}"] = rem; style(ws[f"H{r}"], f_grey, al=AL_L)
        r += 1
    B.name("CNU", "Rates", f"$B${first}:$B${r-1}"); B.name("CNUnit", "Rates", f"$C${first}:$C${r-1}"); B.name("CNCost", "Rates", f"$D${first}:$D${r-1}")
    B.name("XWaste", "Rates", f"$D${first}")
    r += 1
    header("LOGISTICS – mobilisation / demobilisation (cost per trip)", ["Vehicle", "", "Cost per trip", "", "", "", "Source / remarks"])
    first = r
    for nm, cost, rem in VEHICLES + [(None, None, None)] * 3:
        inp(ws, f"B{r}", nm, al=AL_LN); inp(ws, f"D{r}", cost, nf=NF_AED)
        ws[f"H{r}"] = rem; style(ws[f"H{r}"], f_grey, al=AL_L)
        r += 1
    B.name("VHU", "Rates", f"$B${first}:$B${r-1}"); B.name("VHCost", "Rates", f"$D${first}:$D${r-1}")
    r += 1
    header("PERSONNEL EXTRAS", ["Item", "", "Cost", "", "", "", "Source / remarks"])
    for t, n, v, rem in EXTRAS:
        label(ws, f"B{r}", t, f_body); inp(ws, f"D{r}", v, nf=NF_AED); B.name(n, "Rates", f"$D${r}")
        ws[f"H{r}"] = rem; style(ws[f"H{r}"], f_grey, al=AL_L)
        r += 1
    ws.freeze_panes = "C4"


# ================================================================ CHEMICAL CLEANING
def build_cc(B):
    P, code = "CC_", "CC"
    wb = B.wb
    vr.build_lookup(wb)
    vr.build_settings(wb, P, "CC Settings", job_fixed="Chemical Cleaning", project_links=["=QClient", "=QProject", "=QNo"])
    vol = vr.build_volume(wb, P, "CC Volume", "CC Settings")
    vr.build_chemicals(wb, P, "CC Chemicals", rates=True)
    pricing = build_item_pricing(B, P, vol)
    costing = build_costing(B, P, pricing)
    summ = build_cc_summary(B, P, pricing)
    alloc = build_alloc(B, P, pricing, summ)
    build_client_quote(B, code, "Chemical Cleaning Services", alloc,
                       "Chemical cleaning of equipment and piping as per scope, including chemicals, manpower, equipment, "
                       "waste handling, mobilisation and demobilisation.")


def build_item_pricing(B, P, vol):
    """One pricing row per volume row (same row numbers as the volume sheet)."""
    ws = B.wb.create_sheet("CC Item Pricing")
    ws.sheet_view.showGridLines = False
    VS = q(vol["ws"].title)
    VC = vr.COLS
    cols = [("A", "#", 4), ("B", "Loop", 11), ("C", "Tag No.", 16), ("D", "Description", 26), ("E", "Type", 15), ("F", "Regime", 9),
            ("G", "Equipment vol. (m³)", 11), ("H", "Solution vol. (m³)", 11), ("I", "Waste vol. (m³)", 11),
            ("J", "Chem. rate (AED/m³)", 10), ("K", "Chemical cost", 12), ("L", "Waste disposal cost", 12),
            ("M", "Allocation basis", 10), ("N", "Share of shared costs", 9), ("O", "Shared cost allocated", 12),
            ("P", "TOTAL COST", 13), ("Q", "Chemical selling", 12), ("R", "Waste selling", 12), ("S", "Shared selling allocated", 12),
            ("T", "Contingency", 11), ("U", "Price before discount", 13), ("V", "Discount", 11), ("W", "NET PRICE (AED)", 14),
            ("X", "Net price (quote cur.)", 14), ("Y", "unit", 5)]
    for col, t, w in cols:
        ws.column_dimensions[col].width = w
        c = ws[f"{col}{vr.HROW}"]; c.value = t
        style(c, f_hdr, FILL_HDR if col not in ("P", "W", "X") else PatternFill("solid", fgColor="375623"), al=AL_C)
    ws.column_dimensions["Y"].hidden = True
    ws["X7"] = '="("&QCurr&")"'; ws["X7"].font = f_grey; ws["X7"].alignment = AL_C
    ws.row_dimensions[vr.HROW].height = 42
    ws["A1"] = "CHEMICAL CLEANING – PRICE PER EQUIPMENT  (internal)"; ws["A1"].font = f_title
    ws["A2"] = '="Quotation "&QNo&"   ·   "&QClient&"   ·   "&QProject'; ws["A2"].font = f_sub
    ws["A3"] = ("Each row = the same row on 'CC Volume'. Direct: chemicals (solution × regime rate) and waste disposal. "
                "Shared costs from 'CC Costing' are allocated by the basis chosen there."); ws["A3"].font = f_grey
    ws.freeze_panes = f"E{vr.FIRST_DATA}"
    mk_chem = "INDEX(MkMatrix,3,3)"
    data_rows = []
    for u, title, first, last, tot in vol["units"]:
        ws.merge_cells(f"A{title}:D{title}")
        ws[f"A{title}"] = f'="UNIT {u}:  "&INDEX({P}UnitNames,{u})'; ws[f"A{title}"].font = Font(name=FN, size=10, bold=True, color=NAVY)
        for col, *_ in cols[:-1]:
            ws[f"{col}{title}"].fill = FILL_UNIT
        for r in range(first, last + 1):
            data_rows.append(r)
            ws[f"A{r}"] = f"={VS}!{VC['no']}{r}"
            for col, key in (("B", "loop"), ("C", "tag"), ("D", "desc"), ("E", "type"), ("F", "reg")):
                ws[f"{col}{r}"] = f'=IF({VS}!{VC[key]}{r}="","",{VS}!{VC[key]}{r})'
            for col, key in (("G", "vol"), ("H", "sol"), ("I", "waste")):
                ws[f"{col}{r}"] = f"={VS}!{VC[key]}{r}"
            ws[f"J{r}"] = f'=IF(N(H{r})=0,"",IFERROR(INDEX({P}RegRate,MATCH(F{r},{P}RegU,0)),0))'
            ws[f"K{r}"] = f"=N(H{r})*N(J{r})"
            ws[f"L{r}"] = f'=IF({P}IncWaste="Yes",N(I{r})*N(XWaste),0)'
            ws[f"M{r}"] = (f'=IF(N(G{r})=0,0,IF({P}AllocBasis="Equipment count",1,'
                           f'IF({P}AllocBasis="Equipment volume",N(G{r}),N(H{r}))))')
            ws[f"N{r}"] = f"=IF(SUM({P}P_M)=0,0,M{r}/SUM({P}P_M))"
            ws[f"O{r}"] = f"=N{r}*{P}SharedCost"
            ws[f"P{r}"] = f"=K{r}+L{r}+O{r}"
            ws[f"Q{r}"] = f"=K{r}*(1+{mk_chem})"
            ws[f"R{r}"] = f"=L{r}*(1+{mk_chem})"
            ws[f"S{r}"] = f"=N{r}*{P}SharedSell"
            ws[f"T{r}"] = f"=P{r}*N(QCont)"
            ws[f"U{r}"] = f"=Q{r}+R{r}+S{r}+T{r}"
            ws[f"V{r}"] = f"=-U{r}*N(QDisc)"
            ws[f"W{r}"] = f"=U{r}+V{r}"
            ws[f"X{r}"] = f"=W{r}/QRate"
            ws[f"Y{r}"] = u
            for col, *_ in cols[:-1]:
                c = ws[f"{col}{r}"]
                nf = NF_M3 if col in "GHI" else NF_PCT if col == "N" else NF_NUM if col == "M" else NF_AED
                if col in "ABCDEF":
                    style(c, f_body, al=AL_C if col in "AEF" else AL_LN)
                else:
                    strong = col in ("P", "W", "X")
                    style(c, f_bold if strong else f_body, FILL_KEY if strong else FILL_RES, nf=nf, al=AL_C)
            ws.conditional_formatting.add(f"J{r}", FormulaRule(formula=[f'AND(N(H{r})>0,N(J{r})=0)'], fill=FILL_MISS))
        ws.merge_cells(f"A{tot}:F{tot}")
        ws[f"A{tot}"] = f'="TOTAL – "&INDEX({P}UnitNames,{u})'; ws[f"A{tot}"].font = f_bold; ws[f"A{tot}"].alignment = AL_R
        for col, *_ in cols[:-1]:
            cc = ws[f"{col}{tot}"]; cc.fill = FILL_TOT; cc.border = BORDER
            if col in "GHIKLOPQRSTUVWX":
                cc.value = f"=SUM({col}{first}:{col}{last})"; cc.number_format = NF_M3 if col in "GHI" else NF_AED
                cc.font = f_bold; cc.alignment = AL_C
    g = vol["grand"]
    ws.merge_cells(f"A{g}:F{g}")
    ws[f"A{g}"] = "GRAND TOTAL – ALL UNITS"; ws[f"A{g}"].font = Font(name=FN, size=11, bold=True, color="FFFFFF"); ws[f"A{g}"].alignment = AL_R
    for col, *_ in cols[:-1]:
        cc = ws[f"{col}{g}"]; cc.fill = FILL_HDR
        if col in "GHIKLOPQRSTUVWX":
            cc.value = "=" + "+".join(f"{col}{t[4]}" for t in vol["units"])
            cc.number_format = NF_M3 if col in "GHI" else NF_AED; cc.font = Font(name=FN, size=10, bold=True, color="FFFFFF"); cc.alignment = AL_C
    f, l = vr.FIRST_DATA, g
    B.name(f"{P}P_M", ws.title, f"$M${f}:$M${l}")
    for col, k in (("B", "loop"), ("Y", "unit"), ("G", "vol"), ("H", "sol"), ("I", "waste"), ("K", "chem"), ("L", "wastec"),
                   ("P", "cost"), ("W", "net"), ("C", "tag"), ("D", "desc")):
        B.name(f"{P}P_{k}", ws.title, f"${col}${f}:${col}${l}")
    for col, k in (("K", "ChemCost"), ("L", "WasteCost"), ("Q", "ChemSell"), ("R", "WasteSell"), ("W", "NetCheck")):
        B.name(f"{P}{k}", ws.title, f"${col}${g}")
    ws.print_title_rows = f"{vr.HROW}:{vr.UROW}"
    ws.print_area = f"A1:X{g}"
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    return {"ws": ws, "units": vol["units"], "grand": g, "data_rows": data_rows, "vol": vol}


def build_costing(B, P, pricing):
    ws = B.wb.create_sheet("CC Costing")
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 4, "B": 44, "C": 13, "D": 11, "E": 16, "F": 11, "G": 12, "H": 14, "I": 15, "J": 9, "K": 15, "L": 44}.items():
        ws.column_dimensions[c].width = w
    ws["B1"] = "CHEMICAL CLEANING – INTERNAL COSTING  (do not send to client)"; ws["B1"].font = f_title
    ws["B2"] = '="Quotation "&QNo&" Rev "&QRev&"   ·   "&QClient&"   ·   "&QProject'; ws["B2"].font = f_sub
    ws["B3"] = "Shared costs (crew, equipment, services, mob/demob, extras) are entered here; chemicals & waste are priced per equipment."
    ws["B3"].font = f_grey
    bar(ws, 5, "SCOPE, DURATION & PRICING OPTIONS")
    scope = [("Total equipment volume (m³)", f"={P}V_GrandVol", NF_M3), ("Total solution volume (m³)", f"={P}V_GrandSol", NF_M3),
             ("Total waste volume (m³)", f"={P}V_GrandWaste", NF_M3), ("Equipment items with volume", f"=COUNT({P}V_vol)-{len(pricing['units'])+1}", "0")]
    g = pricing["vol"]["grand"]; VS = q(pricing["vol"]["ws"].title); VC = vr.COLS
    B.name(f"{P}V_GrandVol", "CC Volume", f"${VC['vol']}${g}"); B.name(f"{P}V_GrandSol", "CC Volume", f"${VC['sol']}${g}")
    B.name(f"{P}V_GrandWaste", "CC Volume", f"${VC['waste']}${g}")
    scope[3] = ("Equipment items with volume", f'=SUMPRODUCT(({P}V_x_unit>0)*ISNUMBER({P}V_vol))', "0")
    for i, (t, f, nf) in enumerate(scope):
        rr = 6 + i
        label(ws, f"B{rr}", t, f_body); ws[f"C{rr}"] = f; style(ws[f"C{rr}"], f_link, FILL_RES, nf=nf, al=AL_C)
    opts = [("Job duration on site (days)", "Days", None, None, "Total days (rig-up to rig-down); lines can override days"),
            ("Shifts per day (Quote Info)", None, "=QShiftsDay", None, ""),
            ("Shared-cost allocation basis", "AllocBasis", "Solution volume", '"Solution volume,Equipment volume,Equipment count"',
             "How crew / equipment / mob costs are split over the equipment"),
            ("Include waste disposal in price", "IncWaste", "Yes", '"Yes,No"', "Waste volume × waste disposal rate (Rates sheet)"),
            ("Client price breakdown", "Mode", "Per equipment", '"Lump sum,Per unit,Per loop,Per equipment"',
             "Per loop = each loop as one line + equipment not in a loop listed individually")]
    for i, (t, n, v, lst, note) in enumerate(opts):
        rr = 10 + i
        label(ws, f"B{rr}", t)
        if n:
            inp(ws, f"C{rr}", v, nf=NF_NUM if n == "Days" else None); ws.merge_cells(f"C{rr}:D{rr}")
            B.name(f"{P}{n}", ws.title, f"$C${rr}")
            if lst:
                B.dv(ws, lst, f"C{rr}")
            else:
                B.dv(ws, "num", f"C{rr}", num=True)
        else:
            ws[f"C{rr}"] = v; style(ws[f"C{rr}"], f_link, FILL_RES, al=AL_C)
        ws[f"E{rr}"] = note; ws[f"E{rr}"].font = f_grey
    SUM_TOP = 16
    row = 36
    subtotals = []
    HEAD = ["#", "Item", "Source", "Qty", "Unit", "Days / qty", "Total units", "Unit cost (AED)", "Cost (AED)", "Markup", "Selling (AED)", "Notes"]

    def head(title, hdr=HEAD):
        nonlocal row
        bar(ws, row, title); row += 1
        for j, h in enumerate(hdr):
            style(ws.cell(row=row, column=1 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
        ws.row_dimensions[row].height = 26
        row += 1

    def money(r, cat):
        ws[f"I{r}"] = f"=N(G{r})*N(H{r})"
        ws[f"J{r}"] = f'=IF(I{r}=0,0,IFERROR(INDEX(MkMatrix,{cat},MATCH(C{r},SrcU,0)),0))'
        ws[f"K{r}"] = f"=I{r}*(1+J{r})"
        for c, nf in (("I", NF_AED), ("J", NF_PCT), ("K", NF_AED)):
            style(ws[f"{c}{r}"], f_body, FILL_RES, nf=nf, al=AL_C)
        style(ws[f"L{r}"], f_grey, al=AL_L)
        ws.conditional_formatting.add(f"H{r}", FormulaRule(formula=[f'AND(N(G{r})>0,N(H{r})=0)'], fill=FILL_MISS))

    def subtotal(lab, r1, r2):
        nonlocal row
        label(ws, f"B{row}", f"Subtotal – {lab}")
        for c in "IK":
            ws[f"{c}{row}"] = f"=SUM({c}{r1}:{c}{r2})"; style(ws[f"{c}{row}"], f_bold, FILL_KEY, nf=NF_AED, al=AL_C)
        subtotals.append((lab, f"I{row}", f"K{row}"))
        row += 2

    def line_basic(r, i, name, lst, src, unit_text):
        ws[f"A{r}"] = i; style(ws[f"A{r}"], f_grey, al=AL_C)
        inp(ws, f"B{r}", name, al=AL_LN)
        if lst:
            B.dv(ws, lst, f"B{r}")
        inp(ws, f"C{r}", src); B.dv(ws, "=SrcU", f"C{r}")
        inp(ws, f"D{r}", nf=NF_NUM); B.dv(ws, "num", f"D{r}", num=True)
        if unit_text is not None:
            ws[f"E{r}"] = unit_text; style(ws[f"E{r}"], f_grey, al=AL_C)
        inp(ws, f"F{r}", nf=NF_NUM); B.dv(ws, "num", f"F{r}", num=True)
        style(ws[f"G{r}"], f_body, FILL_RES, nf=NF_NUM, al=AL_C)
        style(ws[f"H{r}"], f_link, FILL_RES, nf=NF_AED, al=AL_C)

    crew = f"SUM({P}Crew)*N(QShiftsDay)"
    # 1 manpower
    head("1.  MANPOWER  (persons per shift × days × shifts per day × cost per shift)")
    m1 = row
    for i, g_ in enumerate(["Project Manager", "Project Engineer", "Supervisor", "Job Performer / Permit Receiver", "Safety Officer",
                            "Technician / Operator", "Pipe Fitter", "Helper", None, None, None, None]):
        r = row
        line_basic(r, i + 1, g_, "=MPU", "In-house", "persons / shift")
        ws[f"G{r}"] = f'=IF(N(D{r})=0,"",D{r}*IF(ISNUMBER(F{r}),F{r},N({P}Days))*N(QShiftsDay))'
        ws[f"H{r}"] = (f'=IF(B{r}="","",IFERROR(IF(AND(C{r}="Rented",ISNUMBER(INDEX(MPRent,MATCH(B{r},MPU,0)))),'
                       f'INDEX(MPRent,MATCH(B{r},MPU,0)),INDEX(MPCost,MATCH(B{r},MPU,0))),0))')
        money(r, 1)
        row += 1
    ws[f"L{m1}"] = "Days blank = job duration. Total units = person-shifts."
    B.name(f"{P}Crew", ws.title, f"$D${m1}:$D${row-1}")
    subtotal("Manpower", m1, row - 1)
    # 2 equipment
    head("2.  EQUIPMENT  (qty × days × cost per day; LS items counted once)")
    e1 = row
    eqs = ["Chemical circulation pump skid (SS316, VFD)", "Vessel circulation pump (gamma jet feed)", "Chemical mixing / dosing tank 10 m³",
           "DM water buffer tank", "Waste holding tank 500 bbl (~80 m³)", "Diesel boiler 2 TPH", "Heat exchanger 1000 kW",
           "Gamma jetting nozzle", "Hot-air dryer unit c/w dew point meter", "Hoses & fittings package",
           "Diaphragm pump & manifolds (transfer set)", "Multi gas meter", None, None, None, None]
    for i, e in enumerate(eqs):
        r = row
        line_basic(r, i + 1, e, "=EQU", "In-house", None)
        ws[f"E{r}"] = f'=IF(B{r}="","",IFERROR("per "&INDEX(EQUnit,MATCH(B{r},EQU,0)),""))'; style(ws[f"E{r}"], f_grey, al=AL_C)
        ws[f"G{r}"] = (f'=IF(N(D{r})=0,"",D{r}*IF(IFERROR(INDEX(EQUnit,MATCH(B{r},EQU,0)),"")="LS",1,'
                       f'IF(ISNUMBER(F{r}),F{r},N({P}Days))))')
        ws[f"H{r}"] = (f'=IF(B{r}="","",IFERROR(IF(AND(C{r}="Rented",ISNUMBER(INDEX(EQRent,MATCH(B{r},EQU,0)))),'
                       f'INDEX(EQRent,MATCH(B{r},EQU,0)),INDEX(EQCost,MATCH(B{r},EQU,0))),0))')
        money(r, 2)
        row += 1
    ws[f"L{e1}"] = "Days blank = job duration. LS (lump sum) items: qty × rate."
    subtotal("Equipment", e1, row - 1)
    # 3 services
    head("3.  CONSUMABLES & THIRD-PARTY SERVICES (shared)")
    s1 = row
    for i, (it, un) in enumerate([("Pre-engineering works", "person-day"), ("Kurita litmus paper", "roll"), ("Chemical shipping", "trip"),
                                  ("Local co-operative fee (Al Dhafra)", "month"), (None, None), (None, None), (None, None), (None, None)]):
        r = row
        line_basic(r, i + 1, it, "=CNU", "Third-party", None)
        ws[f"E{r}"] = f'=IF(B{r}="","",IFERROR(INDEX(CNUnit,MATCH(B{r},CNU,0)),""))'; style(ws[f"E{r}"], f_grey, al=AL_C)
        ws[f"G{r}"] = f'=IF(N(D{r})=0,"",D{r}*IF(ISNUMBER(F{r}),F{r},1))'
        ws[f"H{r}"] = f'=IF(B{r}="","",IFERROR(INDEX(CNCost,MATCH(B{r},CNU,0)),0))'
        money(r, 3)
        row += 1
    ws[f"L{s1}"] = "Days / qty column = multiplier (e.g. person-days = persons × days); blank = 1. Waste disposal is priced per equipment."
    subtotal("Consumables & services", s1, row - 1)
    # 4 mob/demob
    head("4.  MOBILISATION / DEMOBILISATION  (vehicles × trips × cost per trip)")
    v1 = row
    for i, vh in enumerate(["40 ft trailer", None, None, None]):
        r = row
        line_basic(r, i + 1, vh, "=VHU", "Rented", "vehicles")
        ws[f"F{r}"] = 2
        ws[f"G{r}"] = f'=IF(N(D{r})=0,"",D{r}*N(F{r}))'
        ws[f"H{r}"] = f'=IF(B{r}="","",IFERROR(INDEX(VHCost,MATCH(B{r},VHU,0)),0))'
        money(r, 2)
        row += 1
    ws[f"L{v1}"] = "Days / qty = no. of trips (2 = mob + demob)"
    subtotal("Mobilisation / demobilisation", v1, row - 1)
    # 5 extras
    head("5.  PERSONNEL EXTRAS  (persons blank = total crew = persons per shift × shifts per day)")
    x1 = row
    extras = [("FAT – food, accommodation & transport", "XFAT", "Third-party", True), ("Disposables / PPE", None, "Third-party", True),
              ("Accommodation – specialists", "XAcc", "Third-party", False), ("Visa", "XVisa", "Third-party", False),
              ("Flight – return", "XFlight", "Third-party", False), ("Crew visas & tickets – India", "XCrewVT", "Third-party", False),
              ("Crew transport – pickup (vehicles × days)", "XTrans", "Rented", None)]
    for i, (t, n, src, per_day) in enumerate(extras):
        r = row
        ws[f"A{r}"] = i + 1; style(ws[f"A{r}"], f_grey, al=AL_C)
        label(ws, f"B{r}", t, f_body)
        inp(ws, f"C{r}", src); B.dv(ws, "=SrcU", f"C{r}")
        inp(ws, f"D{r}", nf=NF_NUM)
        inp(ws, f"F{r}", nf=NF_NUM)
        if per_day is True:
            ws[f"E{r}"] = "persons × days"
            ws[f"G{r}"] = f'=IF(ISNUMBER(D{r}),D{r},{crew})*IF(ISNUMBER(F{r}),F{r},N({P}Days))'
        elif per_day is False:
            ws[f"E{r}"] = "persons"
            ws[f"G{r}"] = f'=IF(ISNUMBER(D{r}),D{r},0)*IF(ISNUMBER(F{r}),F{r},1)'
        else:
            ws[f"E{r}"] = "vehicles × days"
            ws[f"G{r}"] = f'=N(D{r})*IF(ISNUMBER(F{r}),F{r},N({P}Days))'
        style(ws[f"E{r}"], f_grey, al=AL_C); style(ws[f"G{r}"], f_body, FILL_RES, nf=NF_NUM, al=AL_C)
        ws[f"H{r}"] = f"=N({n})" if n else '=IFERROR(INDEX(CNCost,MATCH("Disposables / PPE",CNU,0)),0)'
        style(ws[f"H{r}"], f_link, FILL_RES, nf=NF_AED, al=AL_C)
        money(r, 1 if per_day is not None else 2)
        row += 1
    ws[f"L{x1}"] = "FAT & disposables default to whole crew × job days; enter 0 persons to exclude"
    ws[f"L{x1+2}"] = "Specialists / visas / flights: enter persons"
    subtotal("Personnel extras", x1, row - 1)
    LAST = row
    # summary
    r = SUM_TOP
    bar(ws, r, "COST & PRICE SUMMARY (AED)")
    for j, h in enumerate(["", "Section", "", "", "", "", "", "", "Cost", "", "Selling", ""]):
        if h:
            style(ws.cell(row=r + 1, column=1 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
    lines = [(lab, cc_, sc) for lab, cc_, sc in subtotals]
    for i, (lab, cc_, sc) in enumerate(lines):
        rr = r + 2 + i
        label(ws, f"B{rr}", lab, f_body); ws[f"I{rr}"] = f"={cc_}"; ws[f"K{rr}"] = f"={sc}"
        for c in "IK":
            style(ws[f"{c}{rr}"], f_body, FILL_RES, nf=NF_AED, al=AL_C)
    rr = r + 2 + len(lines)
    label(ws, f"B{rr}", "SHARED COSTS (allocated to equipment)")
    ws[f"I{rr}"] = f"=SUM(I{r+2}:I{rr-1})"; ws[f"K{rr}"] = f"=SUM(K{r+2}:K{rr-1})"
    for c in "IK":
        style(ws[f"{c}{rr}"], f_bold, FILL_KEY, nf=NF_AED, al=AL_C)
    B.name(f"{P}SharedCost", ws.title, f"$I${rr}"); B.name(f"{P}SharedSell", ws.title, f"$K${rr}")
    sh = rr
    direct = [("Chemicals (per equipment: solution × regime rate)", f"={P}ChemCost", f"={P}ChemSell"),
              ("Waste disposal (per equipment: waste vol. × rate)", f"={P}WasteCost", f"={P}WasteSell")]
    for i, (lab, fi, fk) in enumerate(direct):
        x = sh + 1 + i
        label(ws, f"B{x}", lab, f_body); ws[f"I{x}"] = fi; ws[f"K{x}"] = fk
        for c in "IK":
            style(ws[f"{c}{x}"], f_body, FILL_RES, nf=NF_AED, al=AL_C)
    t = sh + 3
    rows_ = [("TOTAL COST / SELLING BEFORE CONTINGENCY", f"=I{sh}+I{sh+1}+I{sh+2}", f"=K{sh}+K{sh+1}+K{sh+2}", True),
             ("Contingency (on cost, Quote Info)", None, f"=I{t}*N(QCont)", False),
             ("Price before discount", None, f"=K{t}+K{t+1}", False),
             ("Discount (Quote Info)", None, f"=-K{t+2}*N(QDisc)", False),
             ("NET PRICE EXCL. VAT (AED)", None, f"=K{t+2}+K{t+3}", True),
             ("Gross margin on net price", None, f'=IF(K{t+4}=0,0,(K{t+4}-I{t})/K{t+4})', False),
             ('="NET PRICE EXCL. VAT ("&QCurr&")"', None, f"=K{t+4}/QRate", True),
             ("Check: sum of equipment prices (should equal net price)", None, f"={P}NetCheck", False)]
    for i, (lab, fi, fk, strong) in enumerate(rows_):
        x = t + i
        label(ws, f"B{x}", lab, f_bold if strong else f_body)
        if fi:
            ws[f"I{x}"] = fi
        ws[f"K{x}"] = fk
        for c in "IK":
            style(ws[f"{c}{x}"], f_bold if strong else f_body, FILL_KEY if strong else FILL_RES,
                  nf=NF_PCT if "margin" in lab else NF_AED, al=AL_C)
    ws.conditional_formatting.add(f"K{t+7}", FormulaRule(formula=[f"ABS(K{t+7}-K{t+4})>0.01"], fill=FILL_MISS))
    ws[f"L{t+7}"] = "Orange = shared costs entered but no equipment volume to allocate them to"; ws[f"L{t+7}"].font = f_grey
    B.name(f"{P}Cost", ws.title, f"$I${t}"); B.name(f"{P}PriceGross", ws.title, f"$K${t+2}")
    B.name(f"{P}Discount", ws.title, f"$K${t+3}"); B.name(f"{P}Net", ws.title, f"$K${t+4}")
    assert t + len(rows_) < 36, t
    ws.freeze_panes = "C5"
    ws.print_area = f"A1:L{LAST}"
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    return ws


def build_cc_summary(B, P, pricing):
    ws = B.wb.create_sheet("CC Summary")
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 5, "C": 32, "D": 8, "E": 14, "F": 14, "G": 14, "H": 15, "I": 15, "J": 16, "K": 16}.items():
        ws.column_dimensions[c].width = w
    ws["B1"] = "CHEMICAL CLEANING – SUMMARY BY UNIT AND BY LOOP"; ws["B1"].font = f_title
    hdr = ["#", "", "Items", "Equipment vol. (m³)", "Solution vol. (m³)", "Waste vol. (m³)", "Chemical cost", "Total cost", "NET PRICE (AED)", '="Net ("&QCurr&")"']
    keys = [("E", "vol"), ("F", "sol"), ("G", "waste"), ("H", "chem"), ("I", "cost"), ("J", "net")]

    def table(r, title, n, namef, crit_range, critf, extra=None):
        bar(ws, r, title, "B", "K")
        for j, h in enumerate(hdr):
            style(ws.cell(row=r + 1, column=2 + j, value=h if j != 1 else ("Unit / train" if crit_range.endswith("unit") else "Loop")),
                  f_hdr, FILL_HDR, al=AL_C)
        first = r + 2
        for i in range(n):
            rr = first + i
            ws[f"B{rr}"] = i + 1 if i < n - (1 if extra else 0) else "–"
            ws[f"C{rr}"] = namef(i)
            crit = critf(i, rr)
            ws[f"D{rr}"] = f'=SUMPRODUCT(({crit_range}={crit})*({P}P_unit>0)*ISNUMBER({P}P_vol))'
            for col, k in keys:
                ws[f"{col}{rr}"] = f'=SUMIFS({P}P_{k},{crit_range},{crit},{P}P_unit,">0")'
            ws[f"K{rr}"] = f"=J{rr}/QRate"
            for col in "BCDEFGHIJK":
                style(ws[f"{col}{rr}"], f_bold if col in "JK" else f_body, FILL_KEY if col in "JK" else (FILL_RES if col >= "E" else None),
                      nf=NF_M3 if col in "EFG" else NF_AED if col in "HIJK" else None, al=AL_LN if col == "C" else AL_C)
        last = first + n - 1
        t = last + 1
        ws[f"C{t}"] = "TOTAL"
        for col in "DEFGHIJK":
            ws[f"{col}{t}"] = f"=SUM({col}{first}:{col}{last})"
        for col in "BCDEFGHIJK":
            style(ws[f"{col}{t}"], f_bold, FILL_TOT, nf=NF_M3 if col in "EFG" else NF_AED if col in "HIJK" else None, al=AL_C)
        return first, last, t

    u1, u2, ut = table(3, "BY UNIT / TRAIN", vr.N_UNITS, lambda i: f"=INDEX({P}UnitNames,{i+1})", f"{P}P_unit", lambda i, rr: str(i + 1))
    L0 = ut + 3
    l1, l2, lt = table(L0, "BY LOOP  (each loop = one line item; 'Not in a loop' = equipment without a loop)", vr.N_LOOPS + 1,
                       lambda i: f"=INDEX({P}LoopU,{i+1})" if i < vr.N_LOOPS else "Not in a loop", f"{P}P_loop",
                       lambda i, rr: f"$C${rr}" if i < vr.N_LOOPS else '""', extra=True)
    return {"ws": ws, "units": (u1, u2), "loops": (l1, l2)}


def build_alloc(B, P, pricing, summ):
    """Hidden: client breakdown list for Lump sum / Per unit / Per loop / Per equipment."""
    ws = B.wb.create_sheet("CC Alloc")
    SM = q(summ["ws"].title); PR = q(pricing["ws"].title)
    ws["A1"] = "Hidden – client price breakdown list"; ws["A1"].font = f_bold
    for j, h in enumerate(["Kind", "Description", "Unit / train", "Volume (m³)", "Price (AED)", "Loop mode flag", "Equip mode flag",
                           "Run loop", "Run equip"]):
        ws.cell(row=3, column=1 + j, value=h).font = f_bold
    r = 4
    l1, l2 = summ["loops"]
    for i in range(vr.N_LOOPS):
        sr = l1 + i
        ws[f"A{r}"] = "Loop"
        ws[f"B{r}"] = (f'={SM}!C{sr}&IF(INDEX({P}LoopD,{i+1})="",""," – "&INDEX({P}LoopD,{i+1}))'
                       f'&" ("&{SM}!D{sr}&" item"&IF({SM}!D{sr}=1,"","s")&")"')
        ws[f"C{r}"] = "–"; ws[f"D{r}"] = f"={SM}!E{sr}"; ws[f"E{r}"] = f"={SM}!J{sr}"
        ws[f"F{r}"] = f"=IF({SM}!D{sr}>0,1,0)"; ws[f"G{r}"] = 0
        r += 1
    for dr in pricing["data_rows"]:
        ws[f"A{r}"] = "Item"
        ws[f"B{r}"] = f'={PR}!C{dr}&IF({PR}!D{dr}="",""," – "&{PR}!D{dr})'
        ws[f"C{r}"] = f"=INDEX({P}UnitNames,{PR}!Y{dr})"
        ws[f"D{r}"] = f"=N({PR}!G{dr})"; ws[f"E{r}"] = f"={PR}!W{dr}"
        ws[f"G{r}"] = f"=IF(D{r}>0,1,0)"
        ws[f"F{r}"] = f'=IF(AND(G{r}=1,{PR}!B{dr}=""),1,0)'
        r += 1
    last = r - 1
    for rr in range(4, last + 1):
        ws[f"H{rr}"] = f'=IF(F{rr}=1,COUNTIF($F$4:F{rr},1),"")'
        ws[f"I{rr}"] = f'=IF(G{rr}=1,COUNTIF($G$4:G{rr},1),"")'
    u1, u2 = summ["units"]
    for j, h in enumerate(["Unit", "Net", "Vol", "Run unit"]):
        ws.cell(row=3, column=11 + j, value=h).font = f_bold
    for i in range(vr.N_UNITS):
        rr = 4 + i; sr = u1 + i
        ws[f"K{rr}"] = f"={SM}!C{sr}"; ws[f"L{rr}"] = f"={SM}!J{sr}"; ws[f"M{rr}"] = f"={SM}!E{sr}"
        ws[f"N{rr}"] = f'=IF(L{rr}>0,COUNTIF($L$4:L{rr},">0"),"")'
    ue = 3 + vr.N_UNITS
    ws["P3"] = "Rows needed"
    ws["P4"] = (f'=IF({P}Mode="Per equipment",MAX(0,MAX(I4:I{last})),IF({P}Mode="Per loop",MAX(0,MAX(H4:H{last})),'
                f'IF({P}Mode="Per unit",MAX(0,MAX(N4:N{ue})),1)))')
    B.name(f"{P}NItems", ws.title, "$P$4")
    nd = last - 3
    for k in range(1, nd + 1):
        rr = 3 + k
        me = f"MATCH({k},$I$4:$I${last},0)"; ml = f"MATCH({k},$H$4:$H${last},0)"; mu = f"MATCH({k},$N$4:$N${ue},0)"

        def pick(colrng, unitcol, lump):
            return (f'=IF({P}Mode="Per equipment",IFERROR(INDEX(${colrng}$4:${colrng}${last},{me}),""),'
                    f'IF({P}Mode="Per loop",IFERROR(INDEX(${colrng}$4:${colrng}${last},{ml}),""),'
                    f'IF({P}Mode="Per unit",IFERROR(INDEX(${unitcol}$4:${unitcol}${ue},{mu}),""),IF({k}=1,{lump},""))))')
        ws[f"R{rr}"] = pick("B", "K", '"LUMP SUM – "&QProject')
        ws[f"S{rr}"] = pick("C", "K", '"All units"')
        ws[f"T{rr}"] = pick("D", "M", f"{P}V_GrandVol")
        ws[f"U{rr}"] = pick("E", "L", f"{P}Net")
    # per-unit description should read "Chemical cleaning – <unit>"
    for k in range(1, nd + 1):
        rr = 3 + k
        ws[f"R{rr}"] = ws[f"R{rr}"].value.replace(f'IFERROR(INDEX($K$4:$K${ue},', f'IFERROR("Chemical cleaning – "&INDEX($K$4:$K${ue},', 1)
    ws.sheet_state = "hidden"
    ws.ndisp = nd
    return ws


def build_client_quote(B, code, title, alloc, scope_text):
    P = f"{code}_"
    ws = B.wb.create_sheet(f"{code} Client Quote")
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 17, "C": 44, "D": 18, "E": 14, "F": 20, "G": 2}.items():
        ws.column_dimensions[c].width = w
    AQ = q(alloc.title)
    ws.merge_cells("B1:F1"); ws["B1"] = "=UPPER(CoName)"; ws["B1"].font = Font(name=FN, size=18, bold=True, color=BLUE)
    ws.merge_cells("B2:F2"); ws["B2"] = "=CoTag"; ws["B2"].font = Font(name=FN, size=10, bold=True, color="FFFFFF")
    ws["B2"].fill = FILL_GREEN
    for i, n in enumerate(["CoAddr", "CoTel", "CoWeb", "CoCert"]):
        ws.merge_cells(f"B{3+i}:F{3+i}"); ws[f"B{3+i}"] = f"={n}"; ws[f"B{3+i}"].font = Font(name=FN, size=9, color="404040")
    for c in "BCDEF":
        ws[f"{c}7"].border = Border(bottom=Side(style="thick", color=BLUE))
    ws.merge_cells("B9:F9"); ws["B9"] = f"QUOTATION – {title.upper()}"
    ws["B9"].font = Font(name=FN, size=14, bold=True, color=NAVY); ws["B9"].alignment = AL_C
    blank = lambda n: f'=IF({n}="","",{n})'
    info = [("Quotation No.", '=QNo&IF(QRev="","","  Rev "&QRev)', "Date", blank("QDate")),
            ("Client", blank("QClient"), "Currency", "=QCurr"), ("Attention", blank("QAttn"), "Your ref.", blank("QRef")),
            ("Project / Plant", blank("QProject"), "Location", blank("QLoc"))]
    for i, (a, fa, b, fb) in enumerate(info):
        r = 11 + i
        ws[f"B{r}"] = a; ws[f"B{r}"].font = f_bold
        ws[f"C{r}"] = fa; ws[f"C{r}"].font = f_body; ws[f"C{r}"].alignment = AL_LN
        ws[f"E{r}"] = b; ws[f"E{r}"].font = f_bold; ws[f"E{r}"].alignment = AL_R
        ws[f"F{r}"] = fb; ws[f"F{r}"].font = f_body; ws[f"F{r}"].alignment = AL_LN
    ws["F11"].number_format = "dd-mmm-yyyy"
    ws.merge_cells("B16:F16"); ws["B16"] = scope_text; ws["B16"].font = f_body; ws["B16"].alignment = AL_L
    ws.row_dimensions[16].height = 30
    bar(ws, 18, "PRICE SUMMARY", "B", "F", fill=FILL_HDR)
    cur = '"("&QCurr&")"'
    summ = [("Price before discount", f"={P}PriceGross/QRate"), ("Discount", f"={P}Discount/QRate"),
            ("TOTAL PRICE EXCL. VAT", f"={P}Net/QRate"), ('="VAT @ "&TEXT(QVAT,"0%")', f"={P}Net/QRate*N(QVAT)"),
            ("TOTAL PRICE INCL. VAT", f"={P}Net/QRate*(1+N(QVAT))")]
    for i, (lab, f) in enumerate(summ):
        r = 19 + i
        strong = lab.startswith("TOTAL")
        ws.merge_cells(f"B{r}:D{r}")
        ws[f"B{r}"] = lab if lab.startswith("=") else f'="{lab} "&{cur}'
        ws[f"B{r}"].font = f_bold if strong else f_body; ws[f"B{r}"].alignment = AL_R
        ws.merge_cells(f"E{r}:F{r}")
        ws[f"E{r}"] = f; style(ws[f"E{r}"], f_bold if strong else f_body, FILL_KEY if strong else None, nf=NF_AED, al=AL_R)
        if strong:
            for c in "BCD":
                ws[f"{c}{r}"].fill = FILL_KEY
    bar(ws, 25, "COMMERCIAL TERMS", "B", "F", fill=FILL_HDR)
    terms = [("Validity", "QTValidity"), ("Payment terms", "QTPayment"), ("Mobilisation", "QTMob"),
             ("Exclusions", "QTExcl"), ("By client", "QTClient"), ("Notes", "QTNotes")]
    for i, (t, n) in enumerate(terms):
        r = 26 + i
        ws[f"B{r}"] = t; ws[f"B{r}"].font = f_bold; ws[f"B{r}"].alignment = AL_LN
        ws.merge_cells(f"C{r}:F{r}"); ws[f"C{r}"] = f'=IF({n}="","–",{n})'; ws[f"C{r}"].alignment = AL_L
        ws[f"C{r}"].font = f_body; ws.row_dimensions[r].height = 28
    r = 26 + len(terms) + 1
    ws[f"B{r}"] = "For and on behalf of"; ws[f"B{r}"].font = f_body
    ws[f"B{r+1}"] = "=CoName"; ws[f"B{r+1}"].font = f_bold
    ws[f"B{r+3}"] = '=IF(QPrep="","",QPrep)'; ws[f"B{r+3}"].font = f_body
    r0 = r + 6
    bar(ws, r0, "PRICE BREAKDOWN", "B", "F", fill=FILL_HDR)
    for j, h in enumerate(["Item", "Description", "Unit / train", "Volume (m³)", '="Price ("&QCurr&")"']):
        style(ws.cell(row=r0 + 1, column=2 + j, value=h), f_hdr, FILL_SECT, al=AL_C)
    for k in range(1, alloc.ndisp + 1):
        r = r0 + 1 + k
        a = 3 + k
        ws[f"B{r}"] = f'=IF({AQ}!R{a}="","",{k})'
        ws[f"C{r}"] = f"={AQ}!R{a}"; ws[f"D{r}"] = f"={AQ}!S{a}"; ws[f"E{r}"] = f"={AQ}!T{a}"
        ws[f"F{r}"] = f'=IF({AQ}!U{a}="","",{AQ}!U{a}/QRate)'
        for c, nf in (("B", "0"), ("C", None), ("D", None), ("E", NF_M3), ("F", NF_AED)):
            ws[f"{c}{r}"].font = f_body; ws[f"{c}{r}"].alignment = AL_C if c != "C" else AL_L
            if nf: ws[f"{c}{r}"].number_format = nf
    last = r0 + 1 + alloc.ndisp
    ws.conditional_formatting.add(f"B{r0+2}:F{last}", FormulaRule(formula=[f'$C{r0+2}<>""'], border=BORDER))
    ws[f"H{r0}"] = "Rows used"; ws[f"I{r0}"] = f"={P}NItems"
    for c in (f"H{r0}", f"I{r0}"):
        ws[c].font = f_grey
    ws.page_setup.orientation = "portrait"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.defined_names["_xlnm.Print_Area"] = DefinedName("_xlnm.Print_Area",
                                                       attr_text=f"OFFSET({q(ws.title)}!$A$1,0,0,{r0+1}+MAX(1,{P}NItems),7)")
    return ws


def build_quotation(B, built):
    ws = B.wb.create_sheet("Quotation")
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 6, "C": 40, "D": 14, "E": 20, "F": 20, "G": 40}.items():
        ws.column_dimensions[c].width = w
    ws["B1"] = "COMBINED QUOTATION – multi-service jobs"; ws["B1"].font = f_title
    ws["B2"] = "Choose the services included. Prices come from each service's costing sheet."; ws["B2"].font = f_sub
    for j, h in enumerate(["#", "Service", "Include?", "Net price (AED)", '="Net price ("&QCurr&")"', "Status"]):
        style(ws.cell(row=4, column=2 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
    for i, (code, nm) in enumerate(SERVICES):
        r = 5 + i
        ws[f"B{r}"] = i + 1; style(ws[f"B{r}"], f_grey, al=AL_C)
        label(ws, f"C{r}", nm, f_body)
        inp(ws, f"D{r}", "Yes" if code in built else "No"); B.dv(ws, '"Yes,No"', f"D{r}")
        if code in built:
            ws[f"E{r}"] = f'=IF(D{r}="Yes",{code}_Net,0)'; ws[f"G{r}"] = "Template built"
        else:
            ws[f"E{r}"] = 0; ws[f"G{r}"] = "Template not built yet (pilot)"
        style(ws[f"E{r}"], f_body, FILL_RES, nf=NF_AED, al=AL_C)
        ws[f"F{r}"] = f"=E{r}/QRate"; style(ws[f"F{r}"], f_body, FILL_RES, nf=NF_AED, al=AL_C)
        style(ws[f"G{r}"], f_grey, al=AL_LN)
    r = 5 + len(SERVICES)
    for i, (lab, f) in enumerate([("TOTAL EXCL. VAT", f"=SUM(E5:E{r-1})"), ('="VAT @ "&TEXT(QVAT,"0%")', f"=E{r}*N(QVAT)"),
                                  ("TOTAL INCL. VAT", f"=E{r}+E{r+1}")]):
        x = r + i
        label(ws, f"C{x}", lab); ws[f"E{x}"] = f; ws[f"F{x}"] = f"=E{x}/QRate"
        for c in "EF":
            style(ws[f"{c}{x}"], f_bold, FILL_KEY, nf=NF_AED, al=AL_C)


def build(services, master, out):
    B = Book()
    build_quote_info(B)
    build_rates(B)
    built = []
    if "CC" in services:
        build_cc(B); built.append("CC")
    if master:
        build_quotation(B, built)
    order = ["Quote Info", "Rates", "Quotation", "CC Settings", "CC Volume", "CC Chemicals", "CC Costing", "CC Item Pricing",
             "CC Summary", "CC Client Quote"]
    sheets = B.wb._sheets
    B.wb._sheets = [s for n in order for s in sheets if s.title == n] + [s for s in sheets if s.title not in order]
    B.wb.active = 0
    B.wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    outdir = sys.argv[1] if len(sys.argv) > 1 else "pricing"
    os.makedirs(outdir, exist_ok=True)
    build(["CC"], True, os.path.join(outdir, "Delight_Pricing_Master.xlsx"))
    build(["CC"], False, os.path.join(outdir, "Delight_Pricing_Chemical_Cleaning.xlsx"))

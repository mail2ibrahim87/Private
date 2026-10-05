"""Row-based volume calculation (one row per equipment, units/trains stacked down the sheet).

Used by build_volume_template.py (standalone template) and build_pricing.py (pricing workbook).
Every defined name is prefixed (e.g. "CC_") so several modules can live in one workbook; the
Lookup sheet / names are shared and created once.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string as cidx
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

N_UNITS, N_ROWS, N_LOOPS, N_REG, N_CHEM_ROWS = 12, 30, 40, 4, 15

FN = "Arial"
NAVY, BLUE = "1F3864", "1E73BE"
f_title = Font(name=FN, size=16, bold=True, color=NAVY)
f_sub = Font(name=FN, size=10, italic=True, color="595959")
f_sect = Font(name=FN, size=11, bold=True, color="FFFFFF")
f_hdr = Font(name=FN, size=9, bold=True, color="FFFFFF")
f_body = Font(name=FN, size=9)
f_bold = Font(name=FN, size=9, bold=True)
f_in = Font(name=FN, size=9, color="0000FF")
f_link = Font(name=FN, size=9, color="008000")
f_grey = Font(name=FN, size=8, color="808080")
f_red = Font(name=FN, size=10, bold=True, color="C00000")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_RES = PatternFill("solid", fgColor="E2EFDA")
FILL_KEY = PatternFill("solid", fgColor="C6E0B4")
FILL_SECT = PatternFill("solid", fgColor=BLUE)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_UNIT = PatternFill("solid", fgColor="D9E1F2")
FILL_TOT = PatternFill("solid", fgColor="BDD7EE")
FILL_MISS = PatternFill("solid", fgColor="F4B084")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
AL_LN = Alignment(horizontal="left", vertical="center")
AL_R = Alignment(horizontal="right", vertical="center")
NF_M3 = '#,##0.000;-#,##0.000;"-"'
NF_NUM = '#,##0.##;-#,##0.##;"-"'
NF_PCT = '0%;-0%;"-"'
NF_PCT1 = '0.0%;-0.0%;"-"'
NF_AED = '#,##0.00;-#,##0.00;"-"'
NF_T = '#,##0.000;-#,##0.000;"-"'

TYPES = ["Exchanger", "Piping", "Column", "Vessel – Vertical", "Vessel – Horizontal", "Tank", "Filter / Other"]
# default % of equipment volume filled: (chemical cleaning, decontamination)
TYPE_FILL = {"Exchanger": (1, 1), "Piping": (1, 1), "Column": (0.2, 1), "Vessel – Vertical": (0.2, 0.3),
             "Vessel – Horizontal": (0.2, 0.3), "Tank": (0.2, 0.3), "Filter / Other": (0.2, 0.3)}
TYPE_NOTE = {"Exchanger": "CC full volume · Decon vapour phase", "Piping": "CC full volume · Decon vapour phase",
             "Column": "CC 20% + gamma jet · Decon vapour phase (complete volume)",
             "Vessel – Vertical": "CC 20% + gamma jet · Decon boil-out 30%",
             "Vessel – Horizontal": "CC 20% + gamma jet · Decon boil-out 30%",
             "Tank": "CC 20% + gamma jet · Decon gamma-jet circulation (TO CONFIRM)",
             "Filter / Other": "CC 20% + gamma jet · Decon boil-out 30%"}
BASIS = ["Dimensions", "Water weight", "Given volume"]
JOBS = ["Chemical Cleaning", "Decontamination"]

# ASAB chemical list (values only): name, purity, price AED/kg, SS conc, CS conc
CHEMS = [("SODIUM HYDROXIDE (CAUSTIC SODA)", 0.49, 3, None, 0.005), ("CITRIC ACID", 1, 4, 0.03, 0.04),
         ("AMMONIUM BIFLUORIDE", 1, 16, None, None), ("SODIUM META SILICATE", 1, 9, None, None),
         ("TRI SODIUM PHOSPHATE", 1, 6, None, None), ("INHIBITORS - RODINE & ARMOHIB", 1, 12, 0.0025, 0.0025),
         ("SURFACTANTS - NP-9 / MULTI CLEAN", 1, 12, 0.002, 0.002), ("HYDROGEN PEROXIDE (35%)", 0.35, 4, None, None),
         ("AMMONIA (25%)", 0.25, 6, None, None), ("SODA ASH (SODIUM CARBONATE)", 1, 10, 0.02, 0.02),
         ("HYDROCHLORIC ACID (HCl)", 0.33, 2, None, None), ("SULPHURIC ACID", 0.98, 8, None, None),
         ("SODIUM NITRITE", 1, 8, None, 0.01), ("SULPHAMIC ACID", 1, 6, None, None), ("ANTIFOAM", 1, 9, None, None)]

LEN_U = [("mm", 0.001), ("cm", 0.01), ("m", 1), ("in", 0.0254), ("ft", 0.3048)]
VOL_U = [("m³", 1), ("L", 0.001), ("US gal", 0.003785411784), ("Imp gal", 0.00454609), ("bbl", 0.158987294928), ("ft³", 0.028316846592)]
WT_U = [("kg", 1), ("t", 1000), ("lb", 0.45359237)]
HEADS = [("None / Open", "0"), ("Flat", "0"), ("2:1 Ellipsoidal", "=PI()/24"), ("Hemispherical", "=PI()/12"),
         ("Torispherical (F&D)", "0.0809"), ("Conical", "0")]
LAYOUTS = [("Triangular (30°)", "=SQRT(3)/2"), ("Rotated triangular (60°)", "=SQRT(3)/2"), ("Square (90°)", 1), ("Rotated square (45°)", 1)]
SCHEDS = ["Sch 5S", "Sch 10S", "Sch 10", "Sch 20", "Sch 30", "Sch 40", "STD", "Sch 40S", "Sch 60",
          "Sch 80", "XS", "Sch 80S", "Sch 100", "Sch 120", "Sch 140", "Sch 160", "XXS"]
PIPES = [
 ('1/2" (DN 15)', 21.3, {"Sch 5S": 1.65, "Sch 10S": 2.11, "Sch 40": 2.77, "STD": 2.77, "Sch 40S": 2.77, "Sch 80": 3.73, "XS": 3.73, "Sch 80S": 3.73, "Sch 160": 4.78, "XXS": 7.47}),
 ('3/4" (DN 20)', 26.7, {"Sch 5S": 1.65, "Sch 10S": 2.11, "Sch 40": 2.87, "STD": 2.87, "Sch 40S": 2.87, "Sch 80": 3.91, "XS": 3.91, "Sch 80S": 3.91, "Sch 160": 5.56, "XXS": 7.82}),
 ('1" (DN 25)', 33.4, {"Sch 5S": 1.65, "Sch 10S": 2.77, "Sch 40": 3.38, "STD": 3.38, "Sch 40S": 3.38, "Sch 80": 4.55, "XS": 4.55, "Sch 80S": 4.55, "Sch 160": 6.35, "XXS": 9.09}),
 ('1-1/4" (DN 32)', 42.2, {"Sch 5S": 1.65, "Sch 10S": 2.77, "Sch 40": 3.56, "STD": 3.56, "Sch 40S": 3.56, "Sch 80": 4.85, "XS": 4.85, "Sch 80S": 4.85, "Sch 160": 6.35, "XXS": 9.70}),
 ('1-1/2" (DN 40)', 48.3, {"Sch 5S": 1.65, "Sch 10S": 2.77, "Sch 40": 3.68, "STD": 3.68, "Sch 40S": 3.68, "Sch 80": 5.08, "XS": 5.08, "Sch 80S": 5.08, "Sch 160": 7.14, "XXS": 10.15}),
 ('2" (DN 50)', 60.3, {"Sch 5S": 1.65, "Sch 10S": 2.77, "Sch 40": 3.91, "STD": 3.91, "Sch 40S": 3.91, "Sch 80": 5.54, "XS": 5.54, "Sch 80S": 5.54, "Sch 160": 8.74, "XXS": 11.07}),
 ('2-1/2" (DN 65)', 73.0, {"Sch 5S": 2.11, "Sch 10S": 3.05, "Sch 40": 5.16, "STD": 5.16, "Sch 40S": 5.16, "Sch 80": 7.01, "XS": 7.01, "Sch 80S": 7.01, "Sch 160": 9.53, "XXS": 14.02}),
 ('3" (DN 80)', 88.9, {"Sch 5S": 2.11, "Sch 10S": 3.05, "Sch 40": 5.49, "STD": 5.49, "Sch 40S": 5.49, "Sch 80": 7.62, "XS": 7.62, "Sch 80S": 7.62, "Sch 160": 11.13, "XXS": 15.24}),
 ('4" (DN 100)', 114.3, {"Sch 5S": 2.11, "Sch 10S": 3.05, "Sch 40": 6.02, "STD": 6.02, "Sch 40S": 6.02, "Sch 80": 8.56, "XS": 8.56, "Sch 80S": 8.56, "Sch 120": 11.13, "Sch 160": 13.49, "XXS": 17.12}),
 ('5" (DN 125)', 141.3, {"Sch 5S": 2.77, "Sch 10S": 3.40, "Sch 40": 6.55, "STD": 6.55, "Sch 40S": 6.55, "Sch 80": 9.53, "XS": 9.53, "Sch 80S": 9.53, "Sch 120": 12.70, "Sch 160": 15.88, "XXS": 19.05}),
 ('6" (DN 150)', 168.3, {"Sch 5S": 2.77, "Sch 10S": 3.40, "Sch 40": 7.11, "STD": 7.11, "Sch 40S": 7.11, "Sch 80": 10.97, "XS": 10.97, "Sch 80S": 10.97, "Sch 120": 14.27, "Sch 160": 18.26, "XXS": 21.95}),
 ('8" (DN 200)', 219.1, {"Sch 5S": 2.77, "Sch 10S": 3.76, "Sch 20": 6.35, "Sch 30": 7.04, "Sch 40": 8.18, "STD": 8.18, "Sch 40S": 8.18, "Sch 60": 10.31, "Sch 80": 12.70, "XS": 12.70, "Sch 80S": 12.70, "Sch 100": 15.09, "Sch 120": 18.26, "Sch 140": 20.62, "Sch 160": 23.01, "XXS": 22.23}),
 ('10" (DN 250)', 273.0, {"Sch 5S": 3.40, "Sch 10S": 4.19, "Sch 20": 6.35, "Sch 30": 7.80, "Sch 40": 9.27, "STD": 9.27, "Sch 40S": 9.27, "Sch 60": 12.70, "XS": 12.70, "Sch 80S": 12.70, "Sch 80": 15.09, "Sch 100": 18.26, "Sch 120": 21.44, "Sch 140": 25.40, "Sch 160": 28.58, "XXS": 25.40}),
 ('12" (DN 300)', 323.8, {"Sch 5S": 3.96, "Sch 10S": 4.57, "Sch 20": 6.35, "Sch 30": 8.38, "STD": 9.53, "Sch 40S": 9.53, "Sch 40": 10.31, "XS": 12.70, "Sch 80S": 12.70, "Sch 60": 14.27, "Sch 80": 17.48, "Sch 100": 21.44, "Sch 120": 25.40, "Sch 140": 28.58, "Sch 160": 33.32, "XXS": 25.40}),
 ('14" (DN 350)', 355.6, {"Sch 5S": 3.96, "Sch 10S": 4.78, "Sch 10": 6.35, "Sch 20": 7.92, "Sch 30": 9.53, "STD": 9.53, "Sch 40": 11.13, "XS": 12.70, "Sch 60": 15.09, "Sch 80": 19.05, "Sch 100": 23.83, "Sch 120": 27.79, "Sch 140": 31.75, "Sch 160": 35.71}),
 ('16" (DN 400)', 406.4, {"Sch 5S": 4.19, "Sch 10S": 4.78, "Sch 10": 6.35, "Sch 20": 7.92, "Sch 30": 9.53, "STD": 9.53, "Sch 40": 12.70, "XS": 12.70, "Sch 60": 16.66, "Sch 80": 21.44, "Sch 100": 26.19, "Sch 120": 30.96, "Sch 140": 36.53, "Sch 160": 40.49}),
 ('18" (DN 450)', 457.0, {"Sch 5S": 4.19, "Sch 10S": 4.78, "Sch 10": 6.35, "Sch 20": 7.92, "STD": 9.53, "Sch 30": 11.13, "XS": 12.70, "Sch 40": 14.27, "Sch 60": 19.05, "Sch 80": 23.83, "Sch 100": 29.36, "Sch 120": 34.93, "Sch 140": 39.67, "Sch 160": 45.24}),
 ('20" (DN 500)', 508.0, {"Sch 5S": 4.78, "Sch 10S": 5.54, "Sch 10": 6.35, "Sch 20": 9.53, "STD": 9.53, "Sch 30": 12.70, "XS": 12.70, "Sch 40": 15.09, "Sch 60": 20.62, "Sch 80": 26.19, "Sch 100": 32.54, "Sch 120": 38.10, "Sch 140": 44.45, "Sch 160": 50.01}),
 ('24" (DN 600)', 610.0, {"Sch 5S": 5.54, "Sch 10S": 6.35, "Sch 10": 6.35, "Sch 20": 9.53, "STD": 9.53, "XS": 12.70, "Sch 30": 14.27, "Sch 40": 17.48, "Sch 60": 24.61, "Sch 80": 30.96, "Sch 100": 38.89, "Sch 120": 46.02, "Sch 140": 52.37, "Sch 160": 59.54}),
 ('30" (DN 750)', 762.0, {"Sch 5S": 6.35, "Sch 10S": 7.92, "Sch 10": 7.92, "STD": 9.53, "Sch 20": 12.70, "XS": 12.70, "Sch 30": 15.88}),
 ('36" (DN 900)', 914.0, {"Sch 10": 7.92, "STD": 9.53, "Sch 20": 12.70, "XS": 12.70, "Sch 30": 15.88, "Sch 40": 19.05}),
]


def q(n):
    return f"'{n}'" if (" " in n or "-" in n or "&" in n) else n


def style(c, font=f_body, fill=None, nf=None, al=None, border=True):
    c.font = font
    if fill: c.fill = fill
    if nf: c.number_format = nf
    if al: c.alignment = al
    if border: c.border = BORDER


def add_name(wb, n, sheet, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=f"{q(sheet)}!{ref}")


class DV:
    def __init__(self, ws):
        self.ws, self.d = ws, {}

    def add(self, formula, cells, num=False):
        key = formula
        if key not in self.d:
            if num:
                d = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                   showErrorMessage=True, errorTitle="Number", error="Enter a positive number.")
            else:
                d = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                                   errorTitle="Invalid entry", error="Pick a value from the drop-down list.")
            self.ws.add_data_validation(d)
            self.d[key] = d
        self.d[key].add(cells)


# ------------------------------------------------------------------ shared lookup sheet
def build_lookup(wb):
    if "Lookup" in wb.sheetnames:
        return wb["Lookup"]
    lk = wb.create_sheet("Lookup")
    lk["A1"] = "Lookup tables (shared by all volume sheets)"; lk["A1"].font = f_bold

    def table(col, title, rows, names):
        lk.cell(row=3, column=col, value=title).font = f_bold
        for i, row in enumerate(rows):
            for j, v in enumerate(row if isinstance(row, tuple) else (row,)):
                lk.cell(row=4 + i, column=col + j, value=v)
        for j, n in enumerate(names):
            L = get_column_letter(col + j)
            add_name(wb, n, "Lookup", f"${L}$4:${L}${3 + len(rows)}")

    table(1, "Length", LEN_U, ["LenU", "LenF"])
    table(4, "Volume", VOL_U, ["VolU", "VolF"])
    table(7, "Weight", WT_U, ["WtU", "WtF"])
    table(10, "Heads", HEADS, ["HeadU", "HeadK"])
    table(13, "Layout", LAYOUTS, ["LayoutU", "LayoutF"])
    table(16, "Basis", BASIS, ["BasisU"])
    table(18, "Types", TYPES, ["TypeU"])
    table(20, "Jobs", JOBS, ["JobU"])
    pr = 14
    lk.cell(row=pr, column=1, value="Pipe wall thickness (mm) – ASME B36.10M / B36.19M (verify against piping class)").font = f_bold
    lk.cell(row=pr + 1, column=1, value="NPS"); lk.cell(row=pr + 1, column=2, value="OD (mm)")
    for j, s in enumerate(SCHEDS):
        lk.cell(row=pr + 1, column=3 + j, value=s)
    for i, (nps, od, wt) in enumerate(PIPES):
        lk.cell(row=pr + 2 + i, column=1, value=nps); lk.cell(row=pr + 2 + i, column=2, value=od)
        for j, s in enumerate(SCHEDS):
            lk.cell(row=pr + 2 + i, column=3 + j, value=wt.get(s))
    p1, p2 = pr + 2, pr + 1 + len(PIPES)
    lc = get_column_letter(2 + len(SCHEDS))
    add_name(wb, "NPSU", "Lookup", f"$A${p1}:$A${p2}")
    add_name(wb, "PipeOD", "Lookup", f"$B${p1}:$B${p2}")
    add_name(wb, "SchU", "Lookup", f"$C${pr+1}:${lc}${pr+1}")
    add_name(wb, "PipeWT", "Lookup", f"$C${p1}:${lc}${p2}")
    lk.sheet_state = "hidden"
    return lk


# ------------------------------------------------------------------ settings
def build_settings(wb, P, sname, job_fixed=None, project_links=None):
    """Settings sheet. P = name prefix (e.g. 'CC_'). Returns dict of useful refs."""
    s = wb.create_sheet(sname)
    s.sheet_view.showGridLines = False
    dv = DV(s)
    for c, w in {"A": 2, "B": 36, "C": 22, "D": 22, "E": 16, "F": 46}.items():
        s.column_dimensions[c].width = w
    s["B1"] = "VOLUME SETTINGS"; s["B1"].font = f_title
    s["B2"] = "Yellow = input. Unit names, loop names and regime names feed the drop-down lists on the volume sheet."; s["B2"].font = f_sub
    r = 4
    s[f"B{r}"] = "PROJECT"; s[f"B{r}"].font = f_bold
    proj = [("Client", "Client"), ("Plant / Project", "Plant"), ("Proposal No.", "ProjNo")]
    for i, (t, n) in enumerate(proj):
        rr = r + 1 + i
        style(s[f"B{rr}"], f_bold, al=AL_LN); s[f"B{rr}"].value = t
        if project_links:
            s[f"C{rr}"] = project_links[i]; style(s[f"C{rr}"], f_link, FILL_RES, al=AL_LN)
        else:
            style(s[f"C{rr}"], f_in, FILL_IN, al=AL_LN)
        s.merge_cells(f"C{rr}:D{rr}")
        add_name(wb, f"{P}{n}", sname, f"$C${rr}")
    rr = r + 5
    style(s[f"B{rr}"], f_bold, al=AL_LN); s[f"B{rr}"].value = "Job type"
    if job_fixed:
        s[f"C{rr}"] = job_fixed; style(s[f"C{rr}"], Font(name=FN, size=10, bold=True), FILL_RES, al=AL_C)
        s[f"D{rr}"] = "Fixed for this service"; s[f"D{rr}"].font = f_grey
    else:
        s[f"C{rr}"] = "Chemical Cleaning"; style(s[f"C{rr}"], Font(name=FN, size=10, bold=True, color="0000FF"), PatternFill("solid", fgColor="FFFF00"), al=AL_C)
        dv.add("=JobU", f"C{rr}")
    add_name(wb, f"{P}JobType", sname, f"$C${rr}")

    r = 11
    s[f"B{r}"] = "GENERAL"; s[f"B{r}"].font = f_bold
    gen = [("Water density (water weight → volume)", 1000, "kg/m³", "WaterRho", "0"),
           ("Tube pitch ratio (pitch ÷ OD)", 1.25, "–", "PitchRatio", "0.00"),
           ("Default circulation / solution allowance", 1.15, "× fill volume", "DefCirc", NF_PCT),
           ("Default no. of fills (for waste volume)", 5, "fills", "DefFills", "0"),
           ("Default waste allowance", 0.25, "on top of fills", "DefAllow", NF_PCT)]
    for i, (t, v, u, n, nf) in enumerate(gen):
        rr = r + 1 + i
        style(s[f"B{rr}"], f_body, al=AL_LN); s[f"B{rr}"].value = t
        style(s[f"C{rr}"], f_in, FILL_IN, nf=nf, al=AL_C); s[f"C{rr}"].value = v
        s[f"D{rr}"] = u; s[f"D{rr}"].font = f_grey
        add_name(wb, f"{P}{n}", sname, f"$C${rr}")
    s[f"F{r+3}"] = "Solution = equipment volume × fill % × circulation (ASAB: 115%)"; s[f"F{r+3}"].font = f_grey
    s[f"F{r+5}"] = "Waste = solution × fills × (1 + allowance)  (ASAB: 5 fills + 25%)"; s[f"F{r+5}"].font = f_grey

    r = 18
    s[f"B{r}"] = "DEFAULT FILL % BY EQUIPMENT TYPE  (used when the Fill % cell on a row is blank)"; s[f"B{r}"].font = f_bold
    for j, h in enumerate(["Equipment type", "Chemical cleaning", "Decontamination", "", "Method note"]):
        if h:
            style(s.cell(row=r + 1, column=2 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
    for i, t in enumerate(TYPES):
        rr = r + 2 + i
        style(s[f"B{rr}"], f_bold, al=AL_LN); s[f"B{rr}"].value = t
        for j, v in enumerate(TYPE_FILL[t]):
            c = s.cell(row=rr, column=3 + j, value=v); style(c, f_in, FILL_IN, nf=NF_PCT, al=AL_C)
        s[f"F{rr}"] = TYPE_NOTE[t]; s[f"F{rr}"].font = f_grey
    t1, t2 = r + 2, r + 1 + len(TYPES)
    add_name(wb, f"{P}TypeFillCC", sname, f"$C${t1}:$C${t2}")
    add_name(wb, f"{P}TypeFillDC", sname, f"$D${t1}:$D${t2}")
    add_name(wb, f"{P}TypeL", sname, f"$B${t1}:$B${t2}")
    s[f"B{t2+1}"] = "ASAB R2 used 15% fill for vessels with gamma jetting – set per project."; s[f"B{t2+1}"].font = f_grey

    r = t2 + 3
    s[f"B{r}"] = f"UNITS / TRAINS  (up to {N_UNITS})"; s[f"B{r}"].font = f_bold
    for i in range(N_UNITS):
        rr = r + 1 + i
        style(s[f"B{rr}"], f_body, al=AL_LN); s[f"B{rr}"].value = f"Unit {i+1} name"
        style(s[f"C{rr}"], f_in, FILL_IN, al=AL_LN); s[f"C{rr}"].value = f"Unit {i+1}"
        s.merge_cells(f"C{rr}:D{rr}")
    u1 = r + 1
    add_name(wb, f"{P}UnitNames", sname, f"$C${u1}:$C${u1+N_UNITS-1}")

    r = u1 + N_UNITS + 1
    s[f"B{r}"] = f"CHEMICAL REGIMES  (up to {N_REG}; concentrations on the chemicals sheet)"; s[f"B{r}"].font = f_bold
    for i, nm in enumerate(["SS", "CS", "Regime 3", "Regime 4"]):
        rr = r + 1 + i
        style(s[f"B{rr}"], f_body, al=AL_LN); s[f"B{rr}"].value = f"Regime {i+1} name"
        style(s[f"C{rr}"], f_in, FILL_IN, al=AL_LN); s[f"C{rr}"].value = nm
    g1 = r + 1
    add_name(wb, f"{P}RegU", sname, f"$C${g1}:$C${g1+N_REG-1}")

    r = g1 + N_REG + 1
    s[f"B{r}"] = f"LOOPS  (up to {N_LOOPS} – rename freely, e.g. ML-1, MV-1)"; s[f"B{r}"].font = f_bold
    style(s[f"B{r+1}"], f_hdr, FILL_HDR, al=AL_C); s[f"B{r+1}"].value = "Loop name (drop-down)"
    style(s[f"C{r+1}"], f_hdr, FILL_HDR, al=AL_C); s[f"C{r+1}"].value = "Loop description"
    s.merge_cells(f"C{r+1}:F{r+1}")
    l1 = r + 2
    for i in range(N_LOOPS):
        rr = l1 + i
        style(s[f"B{rr}"], f_in, FILL_IN, al=AL_LN); s[f"B{rr}"].value = f"Loop {i+1}"
        style(s[f"C{rr}"], f_in, FILL_IN, al=AL_LN); s.merge_cells(f"C{rr}:F{rr}")
    add_name(wb, f"{P}LoopU", sname, f"$B${l1}:$B${l1+N_LOOPS-1}")
    add_name(wb, f"{P}LoopD", sname, f"$C${l1}:$C${l1+N_LOOPS-1}")
    return s


# ------------------------------------------------------------------ volume sheet
COLS = {}  # key -> column letter


def _cols():
    order = [
        ("no", "#", 4, None), ("loop", "Loop", 11, None), ("tag", "Tag No.", 14, None), ("desc", "Description / service", 26, None),
        ("type", "Equipment type", 16, None), ("reg", "Chemical regime", 10, None), ("basis", "Volume basis", 12, None),
        ("given", "Given volume", 10, "vol"), ("ww", "Water weight", 10, "wt"),
        ("nps", "NPS", 12, None), ("sch", "Schedule", 9, None), ("pid", "Pipe ID (if given)", 9, "len"), ("plen", "Pipe length", 9, "len"),
        ("D", "Shell / vessel ID", 9, "len"), ("L", "Length T/T", 9, "len"), ("h1", "Head 1 / shell cover", 13, None),
        ("h2", "Head 2", 13, None), ("cone", "Cone height", 8, "len"),
        ("nt", "No. of tubes", 8, None), ("tod", "Tube OD", 8, "len"), ("tid", "Tube ID", 8, "len"), ("tl", "Tube length", 9, "len"),
        ("tp", "Tube pitch", 8, "len"), ("lay", "Tube layout", 13, None),
        ("cid", "Channel / bonnet ID", 9, "len"), ("fl", "Front channel length", 9, "len"), ("fh", "Front head", 13, None),
        ("rl", "Rear channel length", 9, "len"), ("rh", "Rear head", 13, None),
        ("ts", "Tube side vol.", 10, None), ("ss", "Shell side / vessel vol.", 10, None), ("pv", "Piping vol.", 10, None),
        ("vd", "Volume from dimensions", 10, None), ("vw", "Volume from water weight", 10, None),
        ("fill", "Fill %", 7, None), ("circ", "Circulation %", 8, None), ("fills", "No. of fills", 7, None), ("allow", "Waste allowance %", 8, None),
        ("waste", "WASTE VOLUME", 11, None), ("vol", "EQUIPMENT VOLUME", 12, None), ("sol", "SOLUTION VOLUME", 12, None),
        # hidden helpers
        ("x_od", "OD used (m)", 9, None), ("x_area", "Shell area (m²)", 9, None), ("x_len", "Shell length (m)", 9, None),
        ("x_pid", "Pipe ID used (mm)", 9, None), ("x_fill", "Fill used", 7, None), ("x_unit", "Unit no.", 6, None),
    ]
    for i, (k, *_r) in enumerate(order):
        COLS[k] = get_column_letter(i + 1)
    return order


ORDER = _cols()
DEF_UNITS = {"given": "m³", "ww": "kg", "pid": "mm", "plen": "m", "D": "mm", "L": "mm", "cone": "mm", "tod": "mm",
             "tid": "mm", "tl": "mm", "tp": "mm", "cid": "mm", "fl": "mm", "rl": "mm"}
GROUPS = [("ITEM", "no", "ww"), ("PIPING", "nps", "plen"), ("SHELL / VESSEL", "D", "cone"), ("TUBES (exchangers)", "nt", "lay"),
          ("CHANNELS / BONNETS", "cid", "rh"), ("CALCULATED (m³)", "ts", "vw"), ("FACTORS  (blank = default)", "fill", "allow"),
          ("RESULT (m³)", "waste", "sol")]
UROW, HROW, GROW = 7, 6, 5
FIRST_DATA = 9


def C(k):
    return COLS[k]


def build_volume(wb, P, sname, settings_name, example=True):
    """Volume sheet. Returns layout info: list of units with (title_row, first, last, total_row), grand row."""
    ws = wb.create_sheet(sname)
    ws.sheet_view.showGridLines = False
    dv = DV(ws)
    for k, (key, title, width, kind) in zip(COLS, ORDER):
        ws.column_dimensions[C(key)].width = width
    for key in ("x_od", "x_area", "x_len", "x_pid", "x_fill", "x_unit"):
        ws.column_dimensions[C(key)].hidden = True
    last_vis = C("sol")
    ws["A1"] = "VOLUME CALCULATION – one row per equipment, units stacked"; ws["A1"].font = f_title
    ws["A2"] = f'="Job type: "&{P}JobType&"   ·   Client: "&{P}Client&"   ·   Plant: "&{P}Plant&"   ·   Proposal: "&{P}ProjNo'
    ws["A2"].font = Font(name=FN, size=10, bold=True, color="C00000")
    ws["A3"] = ("Yellow = input (only the columns that apply to the equipment type). Units of each dimension column are chosen "
                "in row 7. Factors left blank use the defaults on the settings sheet. Volumes in m³.")
    ws["A3"].font = f_grey
    for g, a, b in GROUPS:
        ws.merge_cells(f"{C(a)}{GROW}:{C(b)}{GROW}")
        c = ws[f"{C(a)}{GROW}"]; c.value = g; c.font = f_sect; c.fill = FILL_SECT; c.alignment = AL_C
    for key, title, width, kind in ORDER:
        c = ws[f"{C(key)}{HROW}"]; c.value = title
        style(c, f_hdr, FILL_HDR if key not in ("waste", "vol", "sol") else PatternFill("solid", fgColor="375623"), al=AL_C)
        u = ws[f"{C(key)}{UROW}"]
        if kind:
            style(u, Font(name=FN, size=9, bold=True, color="0000FF"), FILL_IN, al=AL_C); u.value = DEF_UNITS[key]
            dv.add({"len": "=LenU", "vol": "=VolU", "wt": "=WtU"}[kind], f"{C(key)}{UROW}")
        else:
            style(u, f_grey, al=AL_C)
            u.value = {"ts": "m³", "ss": "m³", "pv": "m³", "vd": "m³", "vw": "m³", "waste": "m³", "vol": "m³", "sol": "m³",
                       "fill": "%", "circ": "%", "allow": "%", "fills": "nos", "nt": "nos"}.get(key, "")
    ws.row_dimensions[HROW].height = 40
    ws.freeze_panes = f"E{FIRST_DATA}"

    def fL(key):
        return f"INDEX(LenF,MATCH(${C(key)}${UROW},LenU,0))"

    def si(key, r):
        return f"{C(key)}{r}*{fL(key)}"

    def num(key, r):
        return f"ISNUMBER({C(key)}{r})"

    def head(tcell, D, h="0"):
        return (f'IF({tcell}="",0,IF({tcell}="Conical",PI()/12*({D})^2*({h}),'
                f'INDEX(HeadK,MATCH({tcell},HeadU,0))*({D})^3))')

    units = []
    r = FIRST_DATA
    vol_cols = ["ts", "ss", "pv", "vd", "vw", "waste", "vol", "sol"]
    for u in range(1, N_UNITS + 1):
        title = r
        ws.merge_cells(f"A{r}:{C('desc')}{r}")
        c = ws[f"A{r}"]; c.value = f'="UNIT {u}:  "&INDEX({P}UnitNames,{u})'; c.font = Font(name=FN, size=10, bold=True, color=NAVY)
        for k in range(1, cidx(last_vis) + 1):
            ws.cell(row=r, column=k).fill = FILL_UNIT
        r += 1
        first = r
        for i in range(N_ROWS):
            row = r
            ws[f"{C('no')}{row}"] = i + 1; style(ws[f"{C('no')}{row}"], f_grey, al=AL_C)
            for key, title_, width, kind in ORDER:
                if key in ("no",) or key.startswith("x_") or key in vol_cols:
                    continue
                cell = ws[f"{C(key)}{row}"]
                style(cell, f_in, FILL_IN, al=AL_C if key not in ("desc", "tag") else AL_LN,
                      nf=NF_PCT if key in ("fill", "circ", "allow") else None)
            dv.add(f"={P}LoopU", f"{C('loop')}{row}")
            dv.add("=TypeU", f"{C('type')}{row}")
            dv.add(f"={P}RegU", f"{C('reg')}{row}")
            dv.add("=BasisU", f"{C('basis')}{row}")
            dv.add("=NPSU", f"{C('nps')}{row}")
            dv.add("=SchU", f"{C('sch')}{row}")
            for hk in ("h1", "h2", "fh", "rh"):
                dv.add("=HeadU", f"{C(hk)}{row}")
            dv.add("=LayoutU", f"{C('lay')}{row}")
            for k in ("given", "ww", "pid", "plen", "D", "L", "cone", "nt", "tod", "tid", "tl", "tp", "cid", "fl", "rl", "fill", "circ", "fills", "allow"):
                dv.add("num", f"{C(k)}{row}", num=True)
            t = f"{C('type')}{row}"
            R = row
            # helpers
            ws[f"{C('x_od')}{R}"] = f'=IF({num("tod",R)},{si("tod",R)},IF({num("tp",R)},{si("tp",R)}/{P}PitchRatio,0))'
            ws[f"{C('x_area')}{R}"] = (f'=IF({num("D",R)},PI()/4*({si("D",R)})^2,IF(AND({num("nt",R)},{num("tp",R)},{C("lay")}{R}<>""),'
                                       f'{C("nt")}{R}*({si("tp",R)})^2*INDEX(LayoutF,MATCH({C("lay")}{R},LayoutU,0)),0))')
            ws[f"{C('x_len')}{R}"] = f'=IF({num("L",R)},{si("L",R)},IF({num("tl",R)},{si("tl",R)},0))'
            wt = f"INDEX(PipeWT,MATCH({C('nps')}{R},NPSU,0),MATCH({C('sch')}{R},SchU,0))"
            ws[f"{C('x_pid')}{R}"] = (f'=IF({num("pid",R)},{si("pid",R)}*1000,IF(OR({C("nps")}{R}="",{C("sch")}{R}=""),"",'
                                      f'IFERROR(IF({wt}="","Sch n/a",INDEX(PipeOD,MATCH({C("nps")}{R},NPSU,0))-2*{wt}),"")))')
            jobcol = f'IF({P}JobType="Decontamination",{P}TypeFillDC,{P}TypeFillCC)'
            ws[f"{C('x_fill')}{R}"] = f'=IF({num("fill",R)},{C("fill")}{R},IFERROR(INDEX({jobcol},MATCH({t},{P}TypeL,0)),1))'
            ws[f"{C('x_unit')}{R}"] = u
            # calculated volumes
            tubes = f'IF(AND({num("nt",R)},{num("tid",R)},{num("tl",R)}),{C("nt")}{R}*PI()/4*({si("tid",R)})^2*{si("tl",R)},0)'
            front = f'IF({num("cid",R)},PI()/4*({si("cid",R)})^2*N({C("fl")}{R})*{fL("fl")}+{head(C("fh")+str(R), si("cid",R))},0)'
            rear = f'IF({num("cid",R)},PI()/4*({si("cid",R)})^2*N({C("rl")}{R})*{fL("rl")}+{head(C("rh")+str(R), si("cid",R))},0)'
            ws[f"{C('ts')}{R}"] = f'=IF({t}="Exchanger",{tubes}+{front}+{rear},"")'
            disp = f'IF(AND({num("nt",R)},{num("tl",R)}),{C("nt")}{R}*PI()/4*{C("x_od")}{R}^2*{si("tl",R)},0)'
            shellx = (f'{C("x_area")}{R}*{C("x_len")}{R}+' + head(C("h1") + str(R), f'SQRT(4*{C("x_area")}{R}/PI())') + f"-{disp}")
            vessel = (f'IF({num("D",R)},PI()/4*({si("D",R)})^2*N({C("L")}{R})*{fL("L")}+'
                      + head(C("h1") + str(R), si("D", R), f'N({C("cone")}{R})*{fL("cone")}') + "+"
                      + head(C("h2") + str(R), si("D", R), f'N({C("cone")}{R})*{fL("cone")}') + ",0)")
            ws[f"{C('ss')}{R}"] = f'=IF(OR({t}="",{t}="Piping"),"",IF({t}="Exchanger",{shellx},{vessel}))'
            ws[f"{C('pv')}{R}"] = (f'=IF({t}<>"Piping","",IF(AND(ISNUMBER({C("x_pid")}{R}),{num("plen",R)}),'
                                   f'PI()/4*({C("x_pid")}{R}/1000)^2*{si("plen",R)},0))')
            ws[f"{C('vd')}{R}"] = f'=IF({t}="","",N({C("ts")}{R})+N({C("ss")}{R})+N({C("pv")}{R}))'
            ws[f"{C('vw')}{R}"] = f'=IF({num("ww",R)},{C("ww")}{R}*INDEX(WtF,MATCH(${C("ww")}${UROW},WtU,0))/{P}WaterRho,"")'
            given = f'IF({num("given",R)},{C("given")}{R}*INDEX(VolF,MATCH(${C("given")}${UROW},VolU,0)),"")'
            ws[f"{C('vol')}{R}"] = (f'=IF({C("basis")}{R}="Given volume",{given},IF({C("basis")}{R}="Water weight",{C("vw")}{R},'
                                    f'IF(N({C("vd")}{R})=0,"",{C("vd")}{R})))')
            circ = f'IF({num("circ",R)},{C("circ")}{R},{P}DefCirc)'
            ws[f"{C('sol')}{R}"] = f'=IF(N({C("vol")}{R})=0,"",{C("vol")}{R}*{C("x_fill")}{R}*{circ})'
            fills = f'IF({num("fills",R)},{C("fills")}{R},{P}DefFills)'
            allow = f'IF({num("allow",R)},{C("allow")}{R},{P}DefAllow)'
            ws[f"{C('waste')}{R}"] = f'=IF({C("sol")}{R}="","",{C("sol")}{R}*{fills}*(1+{allow}))'
            for key in vol_cols:
                strong = key in ("vol", "sol")
                style(ws[f"{C(key)}{R}"], f_bold if strong else f_body, FILL_KEY if strong else FILL_RES, nf=NF_M3, al=AL_C)
            for key in ("x_od", "x_area", "x_len", "x_pid", "x_fill", "x_unit"):
                ws[f"{C(key)}{R}"].font = f_grey
            r += 1
        last = r - 1
        # unit total row
        tot = r
        ws.merge_cells(f"A{r}:{C('desc')}{r}")
        ws[f"A{r}"] = f'="TOTAL – "&INDEX({P}UnitNames,{u})'
        for k in range(1, cidx(last_vis) + 1):
            cell = ws.cell(row=r, column=k); cell.fill = FILL_TOT; cell.border = BORDER
        ws[f"A{r}"].font = Font(name=FN, size=10, bold=True); ws[f"A{r}"].alignment = AL_R
        for key in vol_cols:
            ws[f"{C(key)}{r}"] = f"=SUM({C(key)}{first}:{C(key)}{last})"
            ws[f"{C(key)}{r}"].font = Font(name=FN, size=10, bold=True); ws[f"{C(key)}{r}"].number_format = NF_M3
            ws[f"{C(key)}{r}"].alignment = AL_C
        ws[f"{C('nt')}{r}"] = f'=COUNT({C("vol")}{first}:{C("vol")}{last})&" items"'
        ws[f"{C('nt')}{r}"].font = f_bold
        units.append((u, title, first, last, tot))
        r += 2
    grand = r
    ws.merge_cells(f"A{r}:{C('desc')}{r}")
    ws[f"A{r}"] = "GRAND TOTAL – ALL UNITS"; ws[f"A{r}"].font = Font(name=FN, size=11, bold=True, color="FFFFFF")
    for k in range(1, cidx(last_vis) + 1):
        ws.cell(row=r, column=k).fill = FILL_HDR
    ws[f"A{r}"].alignment = AL_R
    for key in vol_cols:
        ws[f"{C(key)}{r}"] = "=" + "+".join(f"{C(key)}{u[4]}" for u in units)
        ws[f"{C(key)}{r}"].font = Font(name=FN, size=10, bold=True, color="FFFFFF"); ws[f"{C(key)}{r}"].number_format = NF_M3
        ws[f"{C(key)}{r}"].alignment = AL_C
    # conditional flags
    data = f"{C('vol')}{FIRST_DATA}:{C('vol')}{grand}"
    rng_reg = f"{C('reg')}{FIRST_DATA}:{C('reg')}{grand}"
    ws.conditional_formatting.add(rng_reg, FormulaRule(formula=[f'AND(N(${C("vol")}{FIRST_DATA})>0,${C("reg")}{FIRST_DATA}="",${C("x_unit")}{FIRST_DATA}<>"")'], fill=FILL_MISS))
    ws.conditional_formatting.add(f"{C('vol')}{FIRST_DATA}:{C('vol')}{grand}",
                                  FormulaRule(formula=[f'AND(${C("type")}{FIRST_DATA}<>"",N(${C("vol")}{FIRST_DATA})=0,${C("x_unit")}{FIRST_DATA}<>"")'], fill=FILL_MISS))
    ws.print_title_rows = f"{GROW}:{UROW}"
    ws.print_area = f"A1:{last_vis}{grand}"
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    if example:
        _example(ws, units)
    VR = f"{q(sname)}!"
    rngs = {k: f"{VR}${C(k)}${FIRST_DATA}:${C(k)}${grand}" for k in ("loop", "reg", "vol", "sol", "waste", "type", "x_unit", "tag")}
    for k, v in rngs.items():
        wb.defined_names[f"{P}V_{k}"] = DefinedName(f"{P}V_{k}", attr_text=v)
    return {"ws": ws, "units": units, "grand": grand}


def _example(ws, units):
    """One example row of each kind in Unit 1 (values from ASAB loop sheets / earlier template)."""
    u, title, first, last, tot = units[0]
    ex = [
        {"loop": "Loop 1", "tag": "EXAMPLE ML-1", "desc": "Piping loop – volume given (ASAB ML-1)", "type": "Piping", "reg": "SS",
         "basis": "Given volume", "given": 27.115},
        {"loop": "Loop 2", "tag": "EXAMPLE V-101", "desc": "Vertical vessel (dimensions)", "type": "Vessel – Vertical", "reg": "SS",
         "D": 2000, "L": 6000, "h1": "2:1 Ellipsoidal", "h2": "2:1 Ellipsoidal", "fill": 0.15},
        {"loop": "Loop 2", "tag": "EXAMPLE E-101", "desc": "Shell & tube exchanger", "type": "Exchanger", "reg": "SS",
         "D": 800, "L": 6000, "nt": 500, "tod": 19.05, "tid": 15.75, "tl": 6096, "tp": 23.81, "lay": "Triangular (30°)",
         "cid": 800, "fl": 600, "fh": "Flat", "rl": 300, "rh": "2:1 Ellipsoidal"},
        {"loop": "", "tag": 'EXAMPLE 6"-P-1001', "desc": "Line from isometric", "type": "Piping", "reg": "CS",
         "nps": '6" (DN 150)', "sch": "Sch 40", "plen": 120},
    ]
    for i, d in enumerate(ex):
        r = first + i
        for k, v in d.items():
            if v != "":
                ws[f"{C(k)}{r}"] = v


# ------------------------------------------------------------------ chemicals (regime tables)
def build_chemicals(wb, P, sname, rates=False):
    """Regime tables (ASAB format). rates=True: purity / price looked up from the Rates sheet."""
    ws = wb.create_sheet(sname)
    ws.sheet_view.showGridLines = False
    dv = DV(ws)
    for c, w in {"A": 2, "B": 5, "C": 38, "D": 14, "E": 14, "F": 12, "G": 14, "H": 14, "I": 16, "J": 30}.items():
        ws.column_dimensions[c].width = w
    ws["B1"] = "CHEMICALS % CALCULATION – one table per regime (rate per m³ applies to every equipment of that regime)"
    ws["B1"].font = f_title
    ws["B2"] = ("Enter the % required (by weight of solution). Volume = total solution volume of all equipment rows set to the regime. "
                "Quantity (t) = volume × % ÷ purity (solution 1 t/m³). Rate per m³ = Σ(% ÷ purity × price/kg × 1000).")
    ws["B2"].font = f_grey
    # summary
    r = 4
    for j, h in enumerate(["#", "Regime", "Solution volume (m³)", "Chemical quantity (t)", "Chemical cost (AED)", "RATE (AED per m³)"]):
        style(ws.cell(row=r, column=2 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
    sum_rows = []
    for k in range(N_REG):
        sum_rows.append(r + 1 + k)
    tables = []
    tr = r + N_REG + 3
    for k in range(N_REG):
        ws.merge_cells(f"B{tr}:J{tr}")
        c = ws[f"B{tr}"]; c.value = f'="REGIME {k+1}:  "&INDEX({P}RegU,{k+1})'; c.font = f_sect; c.fill = FILL_SECT
        vol_cell = f"$E${tr+1}"
        ws[f"C{tr+1}"] = "Solution volume of this regime (m³)"; ws[f"C{tr+1}"].font = f_bold
        ws[f"E{tr+1}"] = f'=SUMIFS({P}V_sol,{P}V_reg,INDEX({P}RegU,{k+1}),{P}V_x_unit,">0")'
        style(ws[f"E{tr+1}"], f_link, FILL_RES, nf=NF_M3, al=AL_C)
        hdr = ["S.No", "Chemical", "% required", "Purity %", "Price (AED/kg)", "Quantity (t)", "Cost (AED)", "Cost per m³ (AED)", "Remarks"]
        for j, h in enumerate(hdr):
            style(ws.cell(row=tr + 2, column=2 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
        first = tr + 3
        for i in range(N_CHEM_ROWS):
            rr = first + i
            ws[f"B{rr}"] = i + 1; style(ws[f"B{rr}"], f_grey, al=AL_C)
            name, pur, price, ss, cs = CHEMS[i] if i < len(CHEMS) else (None, None, None, None, None)
            conc = ss if k == 0 else cs if k == 1 else None
            style(ws[f"C{rr}"], f_in, FILL_IN, al=AL_LN); ws[f"C{rr}"].value = name
            style(ws[f"D{rr}"], f_in, FILL_IN, nf='0.00%', al=AL_C); ws[f"D{rr}"].value = conc
            dv.add("num", f"D{rr}", num=True)
            if rates:
                dv.add("=CHU", f"C{rr}")
                ws[f"E{rr}"] = f'=IF(C{rr}="","",IFERROR(INDEX(CHPurity,MATCH(C{rr},CHU,0)),""))'
                ws[f"F{rr}"] = f'=IF(C{rr}="","",IFERROR(INDEX(CHPriceKg,MATCH(C{rr},CHU,0)),""))'
                style(ws[f"E{rr}"], f_link, FILL_RES, nf="0%", al=AL_C); style(ws[f"F{rr}"], f_link, FILL_RES, nf=NF_AED, al=AL_C)
            else:
                style(ws[f"E{rr}"], f_in, FILL_IN, nf="0%", al=AL_C); ws[f"E{rr}"].value = pur
                style(ws[f"F{rr}"], f_in, FILL_IN, nf=NF_AED, al=AL_C); ws[f"F{rr}"].value = price
            ws[f"I{rr}"] = f'=IF(OR(N(D{rr})=0,N(E{rr})=0),0,D{rr}/E{rr}*N(F{rr})*1000)'
            ws[f"G{rr}"] = f'=IF(OR(N(D{rr})=0,N(E{rr})=0),0,{vol_cell}*D{rr}/E{rr})'
            ws[f"H{rr}"] = f"=G{rr}*1000*N(F{rr})"
            for cc, nf in (("G", NF_T), ("H", NF_AED), ("I", NF_AED)):
                style(ws[f"{cc}{rr}"], f_body, FILL_RES, nf=nf, al=AL_C)
            style(ws[f"J{rr}"], f_grey, al=AL_L)
        last = first + N_CHEM_ROWS - 1
        tot = last + 1
        ws[f"C{tot}"] = "TOTAL"; ws[f"C{tot}"].font = f_bold
        for cc, nf in (("G", NF_T), ("H", NF_AED), ("I", NF_AED)):
            ws[f"{cc}{tot}"] = f"=SUM({cc}{first}:{cc}{last})"; style(ws[f"{cc}{tot}"], f_bold, FILL_KEY, nf=nf, al=AL_C)
        ws[f"J{tot}"] = "← RATE AED per m³ of solution"; ws[f"J{tot}"].font = f_red
        tables.append((tr, first, last, tot, vol_cell))
        tr = tot + 3
    for k, sr in enumerate(sum_rows):
        tr_, first, last, tot, vol_cell = tables[k]
        ws[f"B{sr}"] = k + 1; style(ws[f"B{sr}"], f_grey, al=AL_C)
        ws[f"C{sr}"] = f"=INDEX({P}RegU,{k+1})"; style(ws[f"C{sr}"], f_bold, al=AL_LN)
        ws[f"D{sr}"] = f"={vol_cell}"; ws[f"E{sr}"] = f"=G{tot}"; ws[f"F{sr}"] = f"=H{tot}"; ws[f"G{sr}"] = f"=I{tot}"
        for cc, nf in (("D", NF_M3), ("E", NF_T), ("F", NF_AED), ("G", NF_AED)):
            style(ws[f"{cc}{sr}"], f_bold if cc == "G" else f_body, FILL_KEY if cc == "G" else FILL_RES, nf=nf, al=AL_C)
    add_name(wb, f"{P}RegRate", sname, f"$G${sum_rows[0]}:$G${sum_rows[-1]}")
    add_name(wb, f"{P}RegCost", sname, f"$F${sum_rows[0]}:$F${sum_rows[-1]}")
    add_name(wb, f"{P}RegTons", sname, f"$E${sum_rows[0]}:$E${sum_rows[-1]}")
    ws[f"B{sum_rows[-1]+1}"] = ("Equipment rows with volume but no regime get no chemical cost (highlighted orange on the volume sheet). "
                                "Chemical prices per kg from ASAB R2 (values).")
    ws[f"B{sum_rows[-1]+1}"].font = f_grey
    ws.freeze_panes = "B4"
    return ws


# ------------------------------------------------------------------ unit / loop summary (standalone)
def build_summary(wb, P, sname, vol):
    ws = wb.create_sheet(sname)
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 5, "C": 30, "D": 10, "E": 16, "F": 16, "G": 16}.items():
        ws.column_dimensions[c].width = w
    ws["B1"] = "VOLUME SUMMARY – by unit and by loop"; ws["B1"].font = f_title
    VS = q(vol["ws"].title)
    r = 3
    for j, h in enumerate(["#", "Unit / train", "Items", "Equipment vol. (m³)", "Solution vol. (m³)", "Waste vol. (m³)"]):
        style(ws.cell(row=r, column=2 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
    for u, title, first, last, tot in vol["units"]:
        rr = r + u
        ws[f"B{rr}"] = u; ws[f"C{rr}"] = f"=INDEX({P}UnitNames,{u})"
        ws[f"D{rr}"] = f"=COUNT({VS}!{C('vol')}{first}:{C('vol')}{last})"
        ws[f"E{rr}"] = f"={VS}!{C('vol')}{tot}"; ws[f"F{rr}"] = f"={VS}!{C('sol')}{tot}"; ws[f"G{rr}"] = f"={VS}!{C('waste')}{tot}"
        for cc in "BCDEFG":
            style(ws[f"{cc}{rr}"], f_body, FILL_RES if cc in "EFG" else None, nf=NF_M3 if cc in "EFG" else None, al=AL_C if cc != "C" else AL_LN)
    rr = r + N_UNITS + 1
    ws[f"C{rr}"] = "TOTAL"
    for cc in "DEFG":
        ws[f"{cc}{rr}"] = f"=SUM({cc}{r+1}:{cc}{rr-1})"
    for cc in "BCDEFG":
        style(ws[f"{cc}{rr}"], f_bold, FILL_KEY, nf=NF_M3 if cc in "EFG" else None, al=AL_C)
    r = rr + 3
    ws[f"B{r-1}"] = "LOOPS – each loop is one line item (sum of its equipment rows); rows with no loop are listed under 'Not in a loop'"
    ws[f"B{r-1}"].font = f_bold
    for j, h in enumerate(["#", "Loop", "Items", "Equipment vol. (m³)", "Solution vol. (m³)", "Waste vol. (m³)"]):
        style(ws.cell(row=r, column=2 + j, value=h), f_hdr, FILL_HDR, al=AL_C)
    for i in range(N_LOOPS + 1):
        rr = r + 1 + i
        crit = f"INDEX({P}LoopU,{i+1})" if i < N_LOOPS else '""'
        ws[f"B{rr}"] = i + 1 if i < N_LOOPS else "–"
        ws[f"C{rr}"] = f"=INDEX({P}LoopU,{i+1})" if i < N_LOOPS else "Not in a loop"
        ws[f"D{rr}"] = f'=SUMPRODUCT(({P}V_loop={crit})*({P}V_x_unit>0)*ISNUMBER({P}V_vol))'
        for cc, key in (("E", "vol"), ("F", "sol"), ("G", "waste")):
            ws[f"{cc}{rr}"] = f'=SUMIFS({P}V_{key},{P}V_loop,{crit},{P}V_x_unit,">0")'
        for cc in "BCDEFG":
            style(ws[f"{cc}{rr}"], f_body, FILL_RES if cc in "EFG" else None, nf=NF_M3 if cc in "EFG" else None, al=AL_C if cc != "C" else AL_LN)
    return ws
